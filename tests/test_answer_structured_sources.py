from io import BytesIO
from datetime import date, datetime
import sys

import pytest
from fastapi import HTTPException
from openpyxl import Workbook

_original_argv = sys.argv
sys.argv = [sys.argv[0]]
try:
    from lightrag.api.routers.answer_routes import (
        _candidate_query_terms,
        _compact_answer_metadata,
        _compact_connector_materialization_sample,
        _embedding_to_pgvector,
        _extract_excel_rows,
        _body_for_structured_row,
        _guidance_from_structured_profile,
        _parse_db_table_ref,
        _validate_id_lookup_rows,
        SourceConnectorSampleResponse,
        StructuredProfileResponse,
    )
finally:
    sys.argv = _original_argv


def _workbook_bytes(row_count: int) -> bytes:
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("FAQ")
    sheet.append(["code", "question", "answer"])
    for index in range(1, row_count + 1):
        sheet.append(
            [
                f"FAQ-{index:04d}",
                f"Question {index}",
                f"Answer {index}",
            ]
        )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_excel_preview_accepts_exact_row_limit_without_truncation():
    result = _extract_excel_rows(_workbook_bytes(1000), max_rows=1000)

    rows = result[4]
    warnings = result[6]
    truncated = result[7]

    assert result[0][0].max_row == 1001
    assert result[0][0].max_column == 3
    assert len(rows) == 1000
    assert truncated is False
    assert not any("Only the first 1000" in warning for warning in warnings)


def test_excel_preview_marks_rows_above_limit_as_truncated():
    result = _extract_excel_rows(_workbook_bytes(1001), max_rows=1000)

    rows = result[4]
    warnings = result[6]
    truncated = result[7]

    assert len(rows) == 1000
    assert truncated is True
    assert "Only the first 1000 Excel data rows were loaded" in warnings


def test_excel_date_only_cells_are_normalized_without_midnight_time():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["valid_on", "created_at"])
    sheet.append([date(2026, 2, 2), datetime(2026, 2, 2, 9, 30)])
    output = BytesIO()
    workbook.save(output)

    result = _extract_excel_rows(output.getvalue(), max_rows=1000)

    assert result[4][0]["valid_on"] == "2026-02-02"
    assert result[4][0]["created_at"] == "2026-02-02T09:30:00"


@pytest.mark.parametrize(
    ("source_uri", "expected"),
    [
        ("db://public.faq_source", ("public", "faq_source")),
        ("quality.faq_source", ("quality", "faq_source")),
        ("faq_source", ("public", "faq_source")),
    ],
)
def test_db_table_reference_parsing(source_uri, expected):
    assert _parse_db_table_ref(source_uri) == expected


def test_db_table_reference_rejects_sql_fragments():
    with pytest.raises(HTTPException):
        _parse_db_table_ref("db://public.faq_source;DROP TABLE users")


def test_search_metadata_excludes_large_source_profiles():
    metadata = {
        "source_profile": {"structured": {"sample_rows": [{"large": "payload"}]}},
        "raw_content": "large source",
        "source_snapshot_id": "src-test",
        "source_row_index": 999,
    }

    assert _compact_answer_metadata(metadata) == {
        "source_snapshot_id": "src-test",
        "source_row_index": 999,
    }


def test_structured_synonym_column_becomes_finding_guidance():
    guidance = _guidance_from_structured_profile(
        [
            {
                "title": "모바일 앱 실행 오류",
                "question": "모바일 앱이 열리지 않아요",
                "answer": "앱을 다시 설치해 주세요.",
                "synonyms": "휴대폰, 스마트폰, 앱 열기",
            }
        ],
        {
            "title": "title",
            "question": "question",
            "answer": "answer",
        },
        ["synonyms"],
    )

    synonym_guidance = [item for item in guidance if item.source == "synonyms"]
    assert [item.guidance_type for item in synonym_guidance] == [
        "synonym",
        "synonym",
        "synonym",
    ]
    assert [item.text for item in synonym_guidance] == [
        "휴대폰",
        "스마트폰",
        "앱 열기",
    ]
    assert all(item.weight == 1.2 for item in synonym_guidance)


def test_embedding_vector_literal_is_compact_json():
    assert _embedding_to_pgvector([0.25, -0.5, 1.0]) == "[0.25,-0.5,1.0]"


def test_structured_exclusion_column_becomes_negative_guidance():
    guidance = _guidance_from_structured_profile(
        [{"title": "계정 잠금", "exclude_keywords": "퇴사자; 외부 고객"}],
        {"title": "title"},
        ["exclude_keywords"],
    )

    exclusions = [item for item in guidance if item.source == "exclude_keywords"]
    assert [item.guidance_type for item in exclusions] == [
        "negative_keyword",
        "negative_keyword",
    ]
    assert [item.text for item in exclusions] == ["퇴사자", "외부 고객"]


def test_id_lookup_returns_business_id_and_combines_detail_guidance():
    row = {
        "product_id": "PRD-1042",
        "manufacturer": "삼성",
        "model_name": "AX100",
        "symptom": "전원이 켜지지 않음",
    }
    mapping = {
        "id": "product_id",
        "title": "model_name",
        "question": "symptom",
    }

    assert _body_for_structured_row(row, mapping, "id_lookup") == "PRD-1042"
    guidance = _guidance_from_structured_profile(
        [row],
        mapping,
        ["manufacturer"],
        "id_lookup",
    )

    combined = [item for item in guidance if item.source == "id_lookup_details"]
    assert len(combined) == 1
    assert combined[0].guidance_type == "question"
    assert "전원이 켜지지 않음" in combined[0].text
    assert "삼성" in combined[0].text


def test_id_lookup_validation_rejects_blank_ids_and_ambiguous_details():
    validation = _validate_id_lookup_rows(
        [
            {"product_id": "PRD-1", "model": "AX100", "symptom": "전원 불량"},
            {"product_id": "PRD-2", "model": "AX100", "symptom": "전원 불량"},
            {"product_id": "", "model": "BX200", "symptom": "소음"},
        ],
        {"id": "product_id", "title": "model"},
        ["symptom"],
        "id_lookup",
    )

    assert validation.enabled is True
    assert validation.ready is False
    assert validation.blank_id_rows == [3]
    assert validation.ambiguous_detail_groups[0]["ids"] == ["PRD-1", "PRD-2"]


def test_id_lookup_validation_allows_repeated_id_with_different_details():
    validation = _validate_id_lookup_rows(
        [
            {"product_id": "PRD-1", "symptom": "전원 불량"},
            {"product_id": "PRD-1", "symptom": "화면 깜빡임"},
        ],
        {"id": "product_id", "question": "symptom"},
        [],
        "id_lookup",
    )

    assert validation.ready is True
    assert validation.duplicate_ids == ["PRD-1"]
    assert validation.ambiguous_detail_groups == []


def test_candidate_terms_include_korean_particle_and_ending_variants():
    terms = _candidate_query_terms("컴퓨터가 인쇄해도 출력물이")

    assert "컴퓨터" in terms
    assert "인쇄" in terms
    assert "출력물" in terms


def test_materialization_response_limits_embedded_connector_rows():
    rows = [{"id": index, "answer": f"Answer {index}"} for index in range(1000)]
    sample = SourceConnectorSampleResponse(
        connector_id="conn-test",
        connector_type="db_table",
        source_type="json",
        raw_content="large payload",
        rows=rows,
        columns=["id", "answer"],
        row_count=1000,
        row_limit=1000,
    )
    profile = StructuredProfileResponse(
        source_type="json",
        kind="records",
        row_count=1000,
        columns=["id", "answer"],
        sample_rows=rows[:20],
    )

    compact = _compact_connector_materialization_sample(sample, profile)

    assert compact.row_count == 1000
    assert len(compact.rows) == 20
    assert len(compact.raw_content) < 1000
    assert "20 of 1000 sample rows" in compact.warnings[-1]
