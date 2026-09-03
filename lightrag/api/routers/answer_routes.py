"""Answer catalog routes for fixed-answer / FAQ workspaces."""

from __future__ import annotations

import asyncio
import csv
import hashlib
import io
import json
import mimetypes
import os
import re
import tempfile
import time
import traceback
import unicodedata
import uuid
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Literal, Optional
from urllib.parse import urlparse

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, Request, UploadFile
from fastapi.responses import FileResponse, RedirectResponse, Response
from pydantic import BaseModel, Field

from lightrag.kg.shared_storage import get_default_workspace
from lightrag.utils import logger

from ..utils_api import decode_workspace_header, get_combined_auth_dependency
from ..utils_s3 import get_s3_client

router = APIRouter(prefix="/api/answers", tags=["answers"])

_get_rag_for_workspace = None
_ensured_dbs: set[int] = set()

AnswerStatus = Literal["draft", "published", "archived", "expired"]
DisplayPolicy = Literal["summary", "full", "both"]
ContentFormat = Literal["plain", "markdown", "html"]
GuidanceType = Literal["keyword", "question", "synonym", "negative_keyword", "note"]
ResolveStrategy = Literal["fast", "balanced"]
RetrievalMode = Literal["keyword", "hybrid", "graph_hybrid", "llm_rerank"]
SelectionPolicy = Literal["workspace", "coverage", "precision"]
StructuredOperator = Literal["contains", "equals", "starts_with", "ends_with"]
StructuredSourceType = Literal["csv", "json"]
StructuredMaterializationMode = Literal["table_as_dataset", "row_per_answer"]
StructuredConversionPurpose = Literal["faq", "id_lookup"]
GuidanceEnrichmentScope = Literal["missing_or_weak", "coverage", "all"]
TermCandidateType = Literal["synonym", "abbreviation", "neologism"]
TermCandidateStatus = Literal["suggested", "approved", "rejected"]
SourceConnectorType = Literal["manual_table", "db_table", "multi_table", "nosql_collection", "web"]
SourceConnectorStatus = Literal["draft", "active", "paused", "error"]
AnswerAssetType = Literal["image", "video", "audio", "table", "file"]
AnswerAssetStorageType = Literal["external", "s3", "local", "inline"]

VALID_WORKSPACE_MODES = {"kms", "answer_catalog", "hybrid"}
ANSWER_WORKSPACE_MODES = {"answer_catalog", "hybrid"}
MAX_EXCEL_UPLOAD_BYTES = 200 * 1024 * 1024
MAX_ANSWER_ASSET_UPLOAD_BYTES = 100 * 1024 * 1024
SUPPORTED_ANSWER_ASSET_EXTENSIONS = {
    ".png",
    ".jpg",
    ".jpeg",
    ".gif",
    ".webp",
    ".mp4",
    ".webm",
    ".mov",
    ".mp3",
    ".wav",
    ".ogg",
    ".pdf",
    ".csv",
    ".xlsx",
    ".xls",
    ".docx",
    ".pptx",
    ".txt",
    ".md",
}
MAX_EXCEL_ROWS_PER_PREVIEW = 1000
MAX_EXCEL_COLUMNS_PER_PREVIEW = 300
MAX_STRUCTURED_ROWS_PER_ANSWER_BATCH = 1000
MAX_KEYWORD_SEARCH_CANDIDATES = 5000
ANSWER_VECTOR_PROJECTION_VERSION = 2
ANSWER_GRAPH_PROJECTION_VERSION = 1
ANSWER_GRAPH_NODE_PREFIX = "FAQ"
DEFAULT_ANSWER_GRAPH_ENTITY_TYPES = [
    "FAQAnswer",
    "Intent",
    "Product",
    "System",
    "Symptom",
    "ErrorCode",
    "Procedure",
    "Category",
    "Term",
]
DEFAULT_ANSWER_GRAPH_RELATION_TYPES = [
    "REPRESENTS_QUESTION",
    "HAS_TERM",
    "HAS_SYMPTOM",
    "APPLIES_TO",
    "BELONGS_TO",
    "RESOLVES",
    "ALIAS_OF",
    "RELATED_TO",
]
DEFAULT_ANSWER_GRAPH_PROMPT = """Extract a compact FAQ retrieval graph from the supplied answer.
Return JSON only with `entities` and `relations`.
Each entity must contain `type`, `label`, and `description`.
Each relation must contain `source_type`, `source_label`, `target_type`,
`target_label`, `type`, and `description`.
Use only the configured entity and relation types. Do not create facts that are
not supported by the FAQ title, body, tags, metadata, or finding hints.
Always connect the FAQAnswer to an Intent with REPRESENTS_QUESTION.
Use HAS_TERM for at most four important domain terms, HAS_SYMPTOM only for an
explicit symptom, APPLIES_TO only for a named product or system, BELONGS_TO for
categories, and RESOLVES when the answer contains a procedure or action.
Prefer three to eight useful relations spanning at least two relation types when
the supplied content supports them. Do not use APPLIES_TO for procedures."""
TOKEN_RE = re.compile(r"[0-9A-Za-z가-힣_]+")
GUIDANCE_TYPE_MULTIPLIER = {
    "question": 1.25,
    "keyword": 1.0,
    "synonym": 0.9,
    "note": 0.55,
    "negative_keyword": -1.4,
}
BUILTIN_ANSWER_ALIAS_GROUPS = [
    {
        "alias_id": "builtin-microsoft-teams",
        "canonical_term": "Microsoft Teams",
        "aliases": ["Teams", "MS Teams", "팀즈", "마이크로소프트 팀즈"],
    },
    {
        "alias_id": "builtin-microsoft-outlook",
        "canonical_term": "Microsoft Outlook",
        "aliases": ["Outlook", "MS Outlook", "아웃룩", "마이크로소프트 아웃룩"],
    },
    {
        "alias_id": "builtin-microsoft-onedrive",
        "canonical_term": "Microsoft OneDrive",
        "aliases": ["OneDrive", "원드라이브", "마이크로소프트 원드라이브"],
    },
    {
        "alias_id": "builtin-wifi",
        "canonical_term": "Wi-Fi",
        "aliases": ["WiFi", "와이파이", "무선랜", "무선 네트워크"],
    },
    {
        "alias_id": "builtin-bluetooth",
        "canonical_term": "Bluetooth",
        "aliases": ["블루투스", "BT"],
    },
    {
        "alias_id": "builtin-vpn",
        "canonical_term": "VPN",
        "aliases": ["가상 사설망", "가상사설망", "사내 VPN"],
    },
]


def set_rag_workspace_getter(getter_func):
    """Set the function to get a workspace-scoped RAG instance."""
    global _get_rag_for_workspace
    _get_rag_for_workspace = getter_func


async def get_workspace_rag(workspace: str):
    if _get_rag_for_workspace is not None:
        return await _get_rag_for_workspace(workspace)
    return None


def _get_workspace_from_request(request: Request) -> str:
    workspace = decode_workspace_header(request.headers.get("LIGHTRAG-WORKSPACE", ""))
    if workspace:
        return workspace
    return get_default_workspace() or "base"


async def _workspace_rag_or_default(http_request: Request, default_rag):
    workspace = _get_workspace_from_request(http_request)
    workspace_rag = await get_workspace_rag(workspace)
    if workspace_rag is None:
        workspace_rag = default_rag
        logger.warning("[Answers] Using default RAG instance for workspace: %s", workspace)
    return workspace, workspace_rag


def _db_from_rag(rag):
    if hasattr(rag, "llm_response_cache") and hasattr(rag.llm_response_cache, "db"):
        return rag.llm_response_cache.db
    if hasattr(rag, "text_chunks") and hasattr(rag.text_chunks, "db"):
        return rag.text_chunks.db
    return None


def _json(value: Any) -> str:
    if isinstance(value, str):
        parsed = _coerce_json(value, None)
        if isinstance(parsed, dict):
            value = parsed
    return json.dumps(value if value is not None else {}, ensure_ascii=False)


def _json_list(value: Any) -> str:
    if isinstance(value, str):
        parsed = _coerce_json(value, None)
        if isinstance(parsed, list):
            value = parsed
    return json.dumps(value if value is not None else [], ensure_ascii=False)


def _coerce_json(value: Any, default: Any) -> Any:
    if value is None:
        return default
    if isinstance(value, str):
        current: Any = value
        for _ in range(2):
            if not isinstance(current, str):
                return current
            try:
                current = json.loads(current)
            except json.JSONDecodeError:
                return default
        return current
    return value


async def _workspace_mode_from_db(db, workspace: str) -> str:
    workspace_row = None
    if hasattr(db, "get_workspace"):
        workspace_row = await db.get_workspace(workspace)
    else:
        workspace_row = await db.query(
            "SELECT workspace_mode, metadata FROM LIGHTRAG_WORKSPACES WHERE workspace_id = $1",
            [workspace],
        )

    if not workspace_row:
        raise HTTPException(status_code=404, detail=f"Workspace '{workspace}' not found")

    metadata = _coerce_json(workspace_row.get("metadata"), {})
    if not isinstance(metadata, dict):
        metadata = {}
    mode = workspace_row.get("workspace_mode") or metadata.get("workspace_mode") or "kms"
    return mode if mode in VALID_WORKSPACE_MODES else "kms"


def _iso(value: Any) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value)


def _now() -> datetime:
    return datetime.now(timezone.utc)


async def _ensure_tables(db) -> None:
    db_key = id(db)
    if db_key in _ensured_dbs:
        return

    statements = [
        "CREATE EXTENSION IF NOT EXISTS vector",
        """
        CREATE TABLE IF NOT EXISTS LIGHTRAG_ANSWER_ITEMS (
            workspace TEXT NOT NULL,
            answer_id TEXT NOT NULL,
            title TEXT NOT NULL,
            body TEXT NOT NULL,
            approved_summary TEXT,
            content_format TEXT NOT NULL DEFAULT 'markdown',
            display_policy TEXT NOT NULL DEFAULT 'both',
            status TEXT NOT NULL DEFAULT 'draft',
            version INTEGER NOT NULL DEFAULT 1,
            valid_from TIMESTAMPTZ,
            valid_until TIMESTAMPTZ,
            priority INTEGER NOT NULL DEFAULT 0,
            tags JSONB NOT NULL DEFAULT '[]'::jsonb,
            metadata JSONB NOT NULL DEFAULT '{}'::jsonb,
            publish_time TIMESTAMPTZ,
            create_time TIMESTAMPTZ NOT NULL DEFAULT NOW(),
            update_time TIMESTAMPTZ NOT NULL DEFAULT NOW(),
            PRIMARY KEY (workspace, answer_id)
        )
        """,
        """
        CREATE TABLE IF NOT EXISTS LIGHTRAG_ANSWER_REVISIONS (
            revision_id TEXT PRIMARY KEY,
            workspace TEXT NOT NULL,
            answer_id TEXT NOT NULL,
            version INTEGER NOT NULL,
            snapshot_json JSONB NOT NULL,
            created_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
        )
        """,
        """
        CREATE TABLE IF NOT EXISTS LIGHTRAG_ANSWER_ASSETS (
            asset_id TEXT PRIMARY KEY,
            workspace TEXT NOT NULL,
            answer_id TEXT NOT NULL,
            answer_version INTEGER NOT NULL DEFAULT 1,
            asset_type TEXT NOT NULL,
            storage_type TEXT NOT NULL,
            storage_uri TEXT,
            file_name TEXT,
            mime_type TEXT,
            file_size BIGINT NOT NULL DEFAULT 0,
            caption TEXT,
            alt_text TEXT,
            search_text TEXT,
            content_text TEXT,
            display_order INTEGER NOT NULL DEFAULT 0,
            is_active BOOLEAN NOT NULL DEFAULT TRUE,
            metadata JSONB NOT NULL DEFAULT '{}'::jsonb,
            create_time TIMESTAMPTZ NOT NULL DEFAULT NOW(),
            update_time TIMESTAMPTZ NOT NULL DEFAULT NOW()
        )
        """,
        """
        CREATE INDEX IF NOT EXISTS idx_answer_assets_workspace_answer
        ON LIGHTRAG_ANSWER_ASSETS (workspace, answer_id, is_active, display_order)
        """,
        """
        CREATE TABLE IF NOT EXISTS LIGHTRAG_ANSWER_GUIDANCE (
            guidance_id TEXT PRIMARY KEY,
            workspace TEXT NOT NULL,
            answer_id TEXT NOT NULL,
            guidance_type TEXT NOT NULL DEFAULT 'keyword',
            text TEXT NOT NULL,
            weight DOUBLE PRECISION NOT NULL DEFAULT 1.0,
            metadata JSONB NOT NULL DEFAULT '{}'::jsonb,
            create_time TIMESTAMPTZ NOT NULL DEFAULT NOW()
        )
        """,
        """
        CREATE TABLE IF NOT EXISTS LIGHTRAG_ANSWER_TERM_ALIASES (
            alias_id TEXT PRIMARY KEY,
            workspace TEXT NOT NULL,
            canonical_term TEXT NOT NULL,
            aliases JSONB NOT NULL DEFAULT '[]'::jsonb,
            enabled BOOLEAN NOT NULL DEFAULT TRUE,
            source TEXT NOT NULL DEFAULT 'workspace',
            metadata JSONB NOT NULL DEFAULT '{}'::jsonb,
            create_time TIMESTAMPTZ NOT NULL DEFAULT NOW(),
            update_time TIMESTAMPTZ NOT NULL DEFAULT NOW()
        )
        """,
        """
        CREATE TABLE IF NOT EXISTS LIGHTRAG_ANSWER_TERM_CANDIDATES (
            candidate_id TEXT PRIMARY KEY,
            workspace TEXT NOT NULL,
            candidate_key TEXT NOT NULL,
            canonical_term TEXT NOT NULL,
            aliases JSONB NOT NULL DEFAULT '[]'::jsonb,
            term_type TEXT NOT NULL DEFAULT 'synonym',
            status TEXT NOT NULL DEFAULT 'suggested',
            confidence DOUBLE PRECISION NOT NULL DEFAULT 0.0,
            rationale TEXT,
            evidence JSONB NOT NULL DEFAULT '[]'::jsonb,
            source TEXT NOT NULL DEFAULT 'llm',
            metadata JSONB NOT NULL DEFAULT '{}'::jsonb,
            create_time TIMESTAMPTZ NOT NULL DEFAULT NOW(),
            update_time TIMESTAMPTZ NOT NULL DEFAULT NOW()
        )
        """,
        """
        CREATE TABLE IF NOT EXISTS LIGHTRAG_ANSWER_EVENTS (
            event_id TEXT PRIMARY KEY,
            workspace TEXT NOT NULL,
            event_type TEXT NOT NULL,
            query TEXT,
            selected_answer_id TEXT,
            candidate_ids JSONB NOT NULL DEFAULT '[]'::jsonb,
            scores JSONB NOT NULL DEFAULT '{}'::jsonb,
            metadata JSONB NOT NULL DEFAULT '{}'::jsonb,
            create_time TIMESTAMPTZ NOT NULL DEFAULT NOW()
        )
        """,
        """
        CREATE TABLE IF NOT EXISTS LIGHTRAG_ANSWER_VECTORS (
            vector_id TEXT PRIMARY KEY,
            workspace TEXT NOT NULL,
            answer_id TEXT NOT NULL,
            answer_version INTEGER NOT NULL DEFAULT 1,
            vector_kind TEXT NOT NULL DEFAULT 'combined',
            content_hash TEXT NOT NULL,
            content TEXT NOT NULL,
            embedding JSONB NOT NULL,
            embedding_vector vector,
            metadata JSONB NOT NULL DEFAULT '{}'::jsonb,
            create_time TIMESTAMPTZ NOT NULL DEFAULT NOW(),
            update_time TIMESTAMPTZ NOT NULL DEFAULT NOW()
        )
        """,
        """
        CREATE TABLE IF NOT EXISTS LIGHTRAG_ANSWER_GRAPH_CONFIG (
            workspace TEXT PRIMARY KEY,
            enabled BOOLEAN NOT NULL DEFAULT FALSE,
            auto_sync BOOLEAN NOT NULL DEFAULT FALSE,
            graph_weight DOUBLE PRECISION NOT NULL DEFAULT 0.35,
            min_similarity DOUBLE PRECISION NOT NULL DEFAULT 0.60,
            max_hops INTEGER NOT NULL DEFAULT 2,
            precision_mode BOOLEAN NOT NULL DEFAULT FALSE,
            precision_min_score DOUBLE PRECISION NOT NULL DEFAULT 0.35,
            min_score_margin DOUBLE PRECISION NOT NULL DEFAULT 0.08,
            min_category_margin DOUBLE PRECISION NOT NULL DEFAULT 0.05,
            min_evidence_sources INTEGER NOT NULL DEFAULT 2,
            llm_min_confidence DOUBLE PRECISION NOT NULL DEFAULT 0.80,
            entity_types JSONB NOT NULL DEFAULT '[]'::jsonb,
            relation_types JSONB NOT NULL DEFAULT '[]'::jsonb,
            extraction_prompt TEXT NOT NULL DEFAULT '',
            ai_extraction_strategy TEXT NOT NULL DEFAULT 'fast',
            ai_retry_max_tokens INTEGER NOT NULL DEFAULT 8192,
            ai_min_relations INTEGER NOT NULL DEFAULT 4,
            ai_min_relation_types INTEGER NOT NULL DEFAULT 3,
            schema_version INTEGER NOT NULL DEFAULT 1,
            create_time TIMESTAMPTZ NOT NULL DEFAULT NOW(),
            update_time TIMESTAMPTZ NOT NULL DEFAULT NOW()
        )
        """,
        """
        ALTER TABLE LIGHTRAG_ANSWER_GRAPH_CONFIG
        ADD COLUMN IF NOT EXISTS min_similarity DOUBLE PRECISION NOT NULL DEFAULT 0.60
        """,
        """
        ALTER TABLE LIGHTRAG_ANSWER_GRAPH_CONFIG
        ALTER COLUMN min_similarity SET DEFAULT 0.60
        """,
        """
        ALTER TABLE LIGHTRAG_ANSWER_GRAPH_CONFIG
        ADD COLUMN IF NOT EXISTS precision_mode BOOLEAN NOT NULL DEFAULT FALSE
        """,
        """
        ALTER TABLE LIGHTRAG_ANSWER_GRAPH_CONFIG
        ADD COLUMN IF NOT EXISTS precision_min_score DOUBLE PRECISION NOT NULL DEFAULT 0.35
        """,
        """
        ALTER TABLE LIGHTRAG_ANSWER_GRAPH_CONFIG
        ADD COLUMN IF NOT EXISTS min_score_margin DOUBLE PRECISION NOT NULL DEFAULT 0.08
        """,
        """
        ALTER TABLE LIGHTRAG_ANSWER_GRAPH_CONFIG
        ADD COLUMN IF NOT EXISTS min_category_margin DOUBLE PRECISION NOT NULL DEFAULT 0.05
        """,
        """
        ALTER TABLE LIGHTRAG_ANSWER_GRAPH_CONFIG
        ADD COLUMN IF NOT EXISTS min_evidence_sources INTEGER NOT NULL DEFAULT 2
        """,
        """
        ALTER TABLE LIGHTRAG_ANSWER_GRAPH_CONFIG
        ADD COLUMN IF NOT EXISTS llm_min_confidence DOUBLE PRECISION NOT NULL DEFAULT 0.80
        """,
        """
        ALTER TABLE LIGHTRAG_ANSWER_GRAPH_CONFIG
        ADD COLUMN IF NOT EXISTS ai_extraction_strategy TEXT NOT NULL DEFAULT 'fast'
        """,
        """
        ALTER TABLE LIGHTRAG_ANSWER_GRAPH_CONFIG
        ADD COLUMN IF NOT EXISTS ai_retry_max_tokens INTEGER NOT NULL DEFAULT 8192
        """,
        """
        ALTER TABLE LIGHTRAG_ANSWER_GRAPH_CONFIG
        ADD COLUMN IF NOT EXISTS ai_min_relations INTEGER NOT NULL DEFAULT 4
        """,
        """
        ALTER TABLE LIGHTRAG_ANSWER_GRAPH_CONFIG
        ALTER COLUMN ai_min_relations SET DEFAULT 4
        """,
        """
        ALTER TABLE LIGHTRAG_ANSWER_GRAPH_CONFIG
        ADD COLUMN IF NOT EXISTS ai_min_relation_types INTEGER NOT NULL DEFAULT 3
        """,
        """
        ALTER TABLE LIGHTRAG_ANSWER_GRAPH_CONFIG
        ALTER COLUMN ai_min_relation_types SET DEFAULT 3
        """,
        """
        CREATE TABLE IF NOT EXISTS LIGHTRAG_ANSWER_GRAPH_PROJECTIONS (
            workspace TEXT NOT NULL,
            answer_id TEXT NOT NULL,
            answer_version INTEGER NOT NULL DEFAULT 1,
            schema_version INTEGER NOT NULL DEFAULT 1,
            status TEXT NOT NULL DEFAULT 'pending',
            content_hash TEXT NOT NULL DEFAULT '',
            nodes JSONB NOT NULL DEFAULT '[]'::jsonb,
            relations JSONB NOT NULL DEFAULT '[]'::jsonb,
            error TEXT,
            built_at TIMESTAMPTZ,
            create_time TIMESTAMPTZ NOT NULL DEFAULT NOW(),
            update_time TIMESTAMPTZ NOT NULL DEFAULT NOW(),
            PRIMARY KEY (workspace, answer_id)
        )
        """,
        "ALTER TABLE LIGHTRAG_ANSWER_VECTORS ADD COLUMN IF NOT EXISTS embedding_vector vector",
        """
        UPDATE LIGHTRAG_ANSWER_VECTORS
        SET embedding_vector = (embedding::text)::vector
        WHERE embedding_vector IS NULL
        """,
        """
        CREATE TABLE IF NOT EXISTS LIGHTRAG_ANSWER_SOURCE_SNAPSHOTS (
            snapshot_id TEXT PRIMARY KEY,
            workspace TEXT NOT NULL,
            source_type TEXT NOT NULL,
            source_uri TEXT,
            file_name TEXT,
            title TEXT,
            raw_content TEXT NOT NULL,
            content_hash TEXT NOT NULL,
            content_length INTEGER NOT NULL DEFAULT 0,
            profile JSONB NOT NULL DEFAULT '{}'::jsonb,
            metadata JSONB NOT NULL DEFAULT '{}'::jsonb,
            created_answer_ids JSONB NOT NULL DEFAULT '[]'::jsonb,
            status TEXT NOT NULL DEFAULT 'captured',
            task_id TEXT,
            create_time TIMESTAMPTZ NOT NULL DEFAULT NOW()
        )
        """,
        """
        CREATE TABLE IF NOT EXISTS LIGHTRAG_ANSWER_SOURCE_LINKS (
            link_id TEXT PRIMARY KEY,
            workspace TEXT NOT NULL,
            answer_id TEXT NOT NULL,
            answer_version INTEGER NOT NULL DEFAULT 1,
            snapshot_id TEXT NOT NULL,
            link_type TEXT NOT NULL DEFAULT 'created_from',
            metadata JSONB NOT NULL DEFAULT '{}'::jsonb,
            create_time TIMESTAMPTZ NOT NULL DEFAULT NOW()
        )
        """,
        """
        CREATE TABLE IF NOT EXISTS LIGHTRAG_STRUCTURED_LOOKUP_LOGS (
            log_id TEXT PRIMARY KEY,
            workspace TEXT NOT NULL,
            dataset_id TEXT NOT NULL,
            dataset_title TEXT,
            pseudo_sql TEXT NOT NULL,
            filters JSONB NOT NULL DEFAULT '[]'::jsonb,
            result_count INTEGER NOT NULL DEFAULT 0,
            preview_only BOOLEAN NOT NULL DEFAULT FALSE,
            latency_ms DOUBLE PRECISION NOT NULL DEFAULT 0,
            metadata JSONB NOT NULL DEFAULT '{}'::jsonb,
            create_time TIMESTAMPTZ NOT NULL DEFAULT NOW()
        )
        """,
        """
        CREATE TABLE IF NOT EXISTS LIGHTRAG_SOURCE_CONNECTORS (
            connector_id TEXT PRIMARY KEY,
            workspace TEXT NOT NULL,
            name TEXT NOT NULL,
            connector_type TEXT NOT NULL,
            status TEXT NOT NULL DEFAULT 'draft',
            config JSONB NOT NULL DEFAULT '{}'::jsonb,
            auth_ref TEXT,
            refresh_policy JSONB NOT NULL DEFAULT '{}'::jsonb,
            enabled BOOLEAN NOT NULL DEFAULT TRUE,
            metadata JSONB NOT NULL DEFAULT '{}'::jsonb,
            create_time TIMESTAMPTZ NOT NULL DEFAULT NOW(),
            update_time TIMESTAMPTZ NOT NULL DEFAULT NOW()
        )
        """,
        "CREATE INDEX IF NOT EXISTS IDX_ANSWERS_WORKSPACE_STATUS ON LIGHTRAG_ANSWER_ITEMS(workspace, status)",
        "CREATE INDEX IF NOT EXISTS IDX_ANSWERS_WORKSPACE_UPDATE ON LIGHTRAG_ANSWER_ITEMS(workspace, update_time DESC)",
        "CREATE INDEX IF NOT EXISTS IDX_ANSWER_GUIDANCE_WORKSPACE_ANSWER ON LIGHTRAG_ANSWER_GUIDANCE(workspace, answer_id)",
        "CREATE INDEX IF NOT EXISTS IDX_ANSWER_TERM_ALIASES_WORKSPACE ON LIGHTRAG_ANSWER_TERM_ALIASES(workspace, enabled)",
        "CREATE UNIQUE INDEX IF NOT EXISTS UIDX_ANSWER_TERM_ALIASES_WORKSPACE_CANONICAL ON LIGHTRAG_ANSWER_TERM_ALIASES(workspace, LOWER(canonical_term))",
        "CREATE INDEX IF NOT EXISTS IDX_ANSWER_TERM_CANDIDATES_WORKSPACE_STATUS ON LIGHTRAG_ANSWER_TERM_CANDIDATES(workspace, status, update_time DESC)",
        "CREATE UNIQUE INDEX IF NOT EXISTS UIDX_ANSWER_TERM_CANDIDATES_WORKSPACE_KEY ON LIGHTRAG_ANSWER_TERM_CANDIDATES(workspace, candidate_key)",
        "CREATE INDEX IF NOT EXISTS IDX_ANSWER_EVENTS_WORKSPACE_TIME ON LIGHTRAG_ANSWER_EVENTS(workspace, create_time DESC)",
        "CREATE UNIQUE INDEX IF NOT EXISTS UIDX_ANSWER_VECTORS_WORKSPACE_ANSWER_KIND ON LIGHTRAG_ANSWER_VECTORS(workspace, answer_id, vector_kind)",
        "CREATE INDEX IF NOT EXISTS IDX_ANSWER_VECTORS_WORKSPACE_UPDATE ON LIGHTRAG_ANSWER_VECTORS(workspace, update_time DESC)",
        "CREATE INDEX IF NOT EXISTS IDX_ANSWER_GRAPH_PROJECTIONS_WORKSPACE_STATUS ON LIGHTRAG_ANSWER_GRAPH_PROJECTIONS(workspace, status, update_time DESC)",
        "CREATE INDEX IF NOT EXISTS IDX_ANSWER_SOURCE_SNAPSHOTS_WORKSPACE_TIME ON LIGHTRAG_ANSWER_SOURCE_SNAPSHOTS(workspace, create_time DESC)",
        "CREATE INDEX IF NOT EXISTS IDX_ANSWER_SOURCE_SNAPSHOTS_WORKSPACE_HASH ON LIGHTRAG_ANSWER_SOURCE_SNAPSHOTS(workspace, content_hash)",
        "CREATE INDEX IF NOT EXISTS IDX_ANSWER_SOURCE_LINKS_WORKSPACE_ANSWER ON LIGHTRAG_ANSWER_SOURCE_LINKS(workspace, answer_id)",
        "CREATE INDEX IF NOT EXISTS IDX_ANSWER_SOURCE_LINKS_WORKSPACE_SNAPSHOT ON LIGHTRAG_ANSWER_SOURCE_LINKS(workspace, snapshot_id)",
        "CREATE INDEX IF NOT EXISTS IDX_STRUCTURED_LOOKUP_LOGS_WORKSPACE_TIME ON LIGHTRAG_STRUCTURED_LOOKUP_LOGS(workspace, create_time DESC)",
        "CREATE INDEX IF NOT EXISTS IDX_STRUCTURED_LOOKUP_LOGS_WORKSPACE_DATASET ON LIGHTRAG_STRUCTURED_LOOKUP_LOGS(workspace, dataset_id)",
        "CREATE INDEX IF NOT EXISTS IDX_SOURCE_CONNECTORS_WORKSPACE_TYPE ON LIGHTRAG_SOURCE_CONNECTORS(workspace, connector_type)",
        "CREATE INDEX IF NOT EXISTS IDX_SOURCE_CONNECTORS_WORKSPACE_UPDATE ON LIGHTRAG_SOURCE_CONNECTORS(workspace, update_time DESC)",
    ]
    for statement in statements:
        await db.execute(statement)
    _ensured_dbs.add(db_key)


class AnswerAsset(BaseModel):
    asset_id: str
    workspace: str
    answer_id: str
    answer_version: int = 1
    asset_type: AnswerAssetType
    storage_type: AnswerAssetStorageType
    storage_uri: Optional[str] = None
    content_url: Optional[str] = None
    file_name: Optional[str] = None
    mime_type: Optional[str] = None
    file_size: int = 0
    caption: Optional[str] = None
    alt_text: Optional[str] = None
    search_text: Optional[str] = None
    content_text: Optional[str] = None
    display_order: int = 0
    is_active: bool = True
    metadata: dict[str, Any] = Field(default_factory=dict)
    create_time: Optional[str] = None
    update_time: Optional[str] = None


class AnswerAssetCreateRequest(BaseModel):
    asset_type: AnswerAssetType
    external_url: Optional[str] = None
    file_name: Optional[str] = None
    mime_type: Optional[str] = None
    caption: Optional[str] = None
    alt_text: Optional[str] = None
    search_text: Optional[str] = None
    content_text: Optional[str] = None
    display_order: int = 0
    metadata: dict[str, Any] = Field(default_factory=dict)


class AnswerAssetUpdateRequest(BaseModel):
    caption: Optional[str] = None
    alt_text: Optional[str] = None
    search_text: Optional[str] = None
    content_text: Optional[str] = None
    display_order: Optional[int] = None
    metadata: Optional[dict[str, Any]] = None


class AnswerItem(BaseModel):
    answer_id: str
    workspace: str
    title: str
    body: str
    approved_summary: Optional[str] = None
    content_format: ContentFormat = "markdown"
    display_policy: DisplayPolicy = "both"
    status: AnswerStatus = "draft"
    version: int = 1
    valid_from: Optional[str] = None
    valid_until: Optional[str] = None
    priority: int = 0
    tags: list[str] = Field(default_factory=list)
    metadata: dict[str, Any] = Field(default_factory=dict)
    assets: list[AnswerAsset] = Field(default_factory=list)
    publish_time: Optional[str] = None
    create_time: Optional[str] = None
    update_time: Optional[str] = None


class AnswerGuidance(BaseModel):
    guidance_id: str
    answer_id: str
    workspace: str
    guidance_type: GuidanceType = "keyword"
    text: str
    weight: float = 1.0
    metadata: dict[str, Any] = Field(default_factory=dict)
    create_time: Optional[str] = None


class AnswerAliasGroup(BaseModel):
    alias_id: str
    workspace: str
    canonical_term: str
    aliases: list[str] = Field(default_factory=list)
    enabled: bool = True
    source: str = "workspace"
    metadata: dict[str, Any] = Field(default_factory=dict)
    create_time: Optional[str] = None
    update_time: Optional[str] = None


class AnswerGraphConfig(BaseModel):
    workspace: str
    enabled: bool = False
    auto_sync: bool = False
    graph_weight: float = Field(default=0.35, ge=0.0, le=1.0)
    min_similarity: float = Field(default=0.60, ge=0.0, le=1.0)
    max_hops: int = Field(default=2, ge=1, le=3)
    precision_mode: bool = False
    precision_min_score: float = Field(default=0.35, ge=0.0, le=1.0)
    min_score_margin: float = Field(default=0.08, ge=0.0, le=1.0)
    min_category_margin: float = Field(default=0.05, ge=0.0, le=1.0)
    min_evidence_sources: int = Field(default=2, ge=1, le=3)
    llm_min_confidence: float = Field(default=0.80, ge=0.0, le=1.0)
    entity_types: list[str] = Field(
        default_factory=lambda: list(DEFAULT_ANSWER_GRAPH_ENTITY_TYPES)
    )
    relation_types: list[str] = Field(
        default_factory=lambda: list(DEFAULT_ANSWER_GRAPH_RELATION_TYPES)
    )
    extraction_prompt: str = DEFAULT_ANSWER_GRAPH_PROMPT
    ai_extraction_strategy: Literal["fast", "adaptive", "deep"] = "fast"
    ai_retry_max_tokens: int = Field(default=8192, ge=1024, le=32768)
    ai_min_relations: int = Field(default=4, ge=1, le=12)
    ai_min_relation_types: int = Field(default=3, ge=1, le=8)
    schema_version: int = 1
    create_time: Optional[str] = None
    update_time: Optional[str] = None


class AnswerGraphConfigUpdate(BaseModel):
    enabled: Optional[bool] = Field(
        default=None,
        description="Allow graph_hybrid retrieval in this FAQ workspace.",
        examples=[True],
    )
    auto_sync: Optional[bool] = Field(
        default=None,
        description="Rebuild an affected FAQ projection after FAQ, hint, or alias changes.",
        examples=[False],
    )
    graph_weight: Optional[float] = Field(
        default=None,
        ge=0.0,
        le=1.0,
        description="Graph contribution to the graph_hybrid candidate score.",
        examples=[0.35],
    )
    min_similarity: Optional[float] = Field(
        default=None,
        ge=0.0,
        le=1.0,
        description=(
            "Minimum semantic similarity required before a graph node can "
            "contribute to an FAQ candidate."
        ),
        examples=[0.60],
    )
    max_hops: Optional[int] = Field(
        default=None,
        ge=1,
        le=3,
        description="Maximum graph relations traversed from a matched term.",
        examples=[2],
    )
    precision_mode: Optional[bool] = Field(
        default=None,
        description=(
            "Answer only when minimum score and evidence checks are strong enough. "
            "Close candidates are recorded for diagnostics, but the highest-ranked "
            "eligible candidate is still selected."
        ),
        examples=[True],
    )
    precision_min_score: Optional[float] = Field(
        default=None,
        ge=0.0,
        le=1.0,
        description="Minimum candidate score used by precision mode.",
        examples=[0.35],
    )
    min_score_margin: Optional[float] = Field(
        default=None,
        ge=0.0,
        le=1.0,
        description=(
            "Score-gap threshold used to flag close FAQ candidates in diagnostics. "
            "It does not suppress the highest-ranked eligible answer."
        ),
        examples=[0.08],
    )
    min_category_margin: Optional[float] = Field(
        default=None,
        ge=0.0,
        le=1.0,
        description=(
            "Score-gap threshold used to flag close FAQ categories in diagnostics. "
            "It does not suppress the highest-ranked eligible answer."
        ),
        examples=[0.05],
    )
    min_evidence_sources: Optional[int] = Field(
        default=None,
        ge=1,
        le=3,
        description="Required count of keyword, vector, and graph evidence sources.",
        examples=[2],
    )
    llm_min_confidence: Optional[float] = Field(
        default=None,
        ge=0.0,
        le=1.0,
        description="Minimum confidence accepted from the constrained LLM ID selector.",
        examples=[0.8],
    )
    entity_types: Optional[list[str]] = Field(
        default=None,
        description="Allowed FAQ graph entity types. FAQAnswer is always retained.",
        examples=[["FAQAnswer", "Intent", "Product", "Symptom", "Term"]],
    )
    relation_types: Optional[list[str]] = Field(
        default=None,
        description="Allowed FAQ graph relation types.",
        examples=[["HAS_TERM", "HAS_SYMPTOM", "APPLIES_TO", "ALIAS_OF"]],
    )
    extraction_prompt: Optional[str] = Field(
        default=None,
        description="Optional LLM extraction prompt. Deterministic projection runs first.",
    )
    ai_extraction_strategy: Optional[Literal["fast", "adaptive", "deep"]] = Field(
        default=None,
        description=(
            "fast disables thinking, adaptive retries weak projections with thinking, "
            "and deep enables thinking for every FAQ."
        ),
        examples=["fast"],
    )
    ai_retry_max_tokens: Optional[int] = Field(default=None, ge=1024, le=32768)
    ai_min_relations: Optional[int] = Field(default=None, ge=1, le=12)
    ai_min_relation_types: Optional[int] = Field(default=None, ge=1, le=8)


class AnswerGraphNode(BaseModel):
    node_id: str
    entity_type: str
    label: str
    description: str = ""
    source: str = "deterministic"


class AnswerGraphRelation(BaseModel):
    source_id: str
    target_id: str
    relation_type: str
    description: str = ""
    weight: float = 1.0
    source: str = "deterministic"


class AnswerGraphProjection(BaseModel):
    workspace: str
    answer_id: str
    answer_version: int
    schema_version: int
    status: str
    content_hash: str = ""
    nodes: list[AnswerGraphNode] = Field(default_factory=list)
    relations: list[AnswerGraphRelation] = Field(default_factory=list)
    error: Optional[str] = None
    built_at: Optional[str] = None
    update_time: Optional[str] = None


class AnswerGraphPreviewRequest(BaseModel):
    answer_id: str = Field(
        description="FAQ answer ID to project without saving.",
        examples=["ANS-IT-OFFICE-007"],
    )
    use_llm: bool = Field(
        default=True,
        description="Build the deterministic projection, then extend it with bounded AI extraction.",
        examples=[True],
    )


class AnswerGraphRebuildRequest(BaseModel):
    answer_ids: Optional[list[str]] = Field(
        default=None,
        description="Optional FAQ IDs. Omit to select from the whole workspace.",
        examples=[["ANS-IT-OFFICE-007", "ANS-IT-OFFICE-008"]],
    )
    include_drafts: bool = Field(
        default=False,
        description="Include draft FAQs in addition to published FAQs.",
    )
    only_stale: bool = Field(
        default=True,
        description="Build only missing, stale, pending, or failed projections.",
    )
    use_llm: bool = Field(
        default=True,
        description="Build deterministic projections, then extend them with bounded AI extraction.",
    )
    limit: int = Field(
        default=500,
        ge=1,
        le=1000,
        description="Maximum FAQs selected for this asynchronous rebuild.",
        examples=[500],
    )


class AnswerGraphRebuildResponse(BaseModel):
    task_id: str
    stream_url: str
    answer_count: int
    message: str
    reused: bool = False


class AnswerGraphStatusResponse(BaseModel):
    workspace: str
    enabled: bool
    total_answers: int
    ready: int
    stale: int
    pending: int
    failed: int
    missing: int
    last_built_at: Optional[str] = None


class AnswerGraphEvidence(BaseModel):
    answer_id: str
    score: float
    start_node: str
    path: list[str] = Field(default_factory=list)
    relation_types: list[str] = Field(default_factory=list)


class AnswerAliasCreateRequest(BaseModel):
    canonical_term: str = Field(..., min_length=2, max_length=200)
    aliases: list[str] = Field(default_factory=list, max_length=30)
    enabled: bool = True
    metadata: dict[str, Any] = Field(default_factory=dict)


class AnswerAliasExpansion(BaseModel):
    canonical_term: str
    matched_term: str
    expanded_terms: list[str] = Field(default_factory=list)
    source: str = "workspace"


class AnswerTermCandidate(BaseModel):
    candidate_id: str
    workspace: str
    canonical_term: str
    aliases: list[str] = Field(default_factory=list)
    term_type: TermCandidateType = "synonym"
    status: TermCandidateStatus = "suggested"
    confidence: float = 0.0
    rationale: Optional[str] = None
    evidence: list[dict[str, Any]] = Field(default_factory=list)
    source: str = "llm"
    metadata: dict[str, Any] = Field(default_factory=dict)
    create_time: Optional[str] = None
    update_time: Optional[str] = None


class AnswerTermDiscoveryRequest(BaseModel):
    include_drafts: bool = True
    include_no_match_queries: bool = True
    answer_limit: int = Field(default=1000, ge=1, le=5000)
    event_limit: int = Field(default=300, ge=0, le=2000)
    batch_size: int = Field(default=20, ge=5, le=40)


class AnswerTermDiscoveryResponse(BaseModel):
    task_id: str
    stream_url: str
    message: str


class AnswerTermCandidateActionResponse(BaseModel):
    candidate: AnswerTermCandidate
    alias_group: Optional[AnswerAliasGroup] = None


class AnswerRevision(BaseModel):
    revision_id: str
    answer_id: str
    workspace: str
    version: int
    snapshot: dict[str, Any]
    created_at: Optional[str] = None


class AnswerEvent(BaseModel):
    event_id: str
    workspace: str
    event_type: str
    query: Optional[str] = None
    selected_answer_id: Optional[str] = None
    candidate_ids: list[str] = Field(default_factory=list)
    scores: dict[str, float] = Field(default_factory=dict)
    metadata: dict[str, Any] = Field(default_factory=dict)
    create_time: Optional[str] = None


class AnswerCreateRequest(BaseModel):
    answer_id: Optional[str] = Field(default=None, description="Optional stable answer id. Defaults to ANS-<uuid>.")
    title: str = Field(..., min_length=1, max_length=500)
    body: str = Field(..., min_length=1)
    approved_summary: Optional[str] = None
    content_format: ContentFormat = "markdown"
    display_policy: DisplayPolicy = "both"
    status: AnswerStatus = "draft"
    valid_from: Optional[datetime] = None
    valid_until: Optional[datetime] = None
    priority: int = 0
    tags: list[str] = Field(default_factory=list)
    metadata: dict[str, Any] = Field(default_factory=dict)
    guidance: list[str] = Field(default_factory=list, description="Initial keyword/question guidance texts.")


class AnswerUpdateRequest(BaseModel):
    title: Optional[str] = Field(default=None, min_length=1, max_length=500)
    body: Optional[str] = Field(default=None, min_length=1)
    approved_summary: Optional[str] = None
    content_format: Optional[ContentFormat] = None
    display_policy: Optional[DisplayPolicy] = None
    status: Optional[AnswerStatus] = None
    valid_from: Optional[datetime] = None
    valid_until: Optional[datetime] = None
    priority: Optional[int] = None
    tags: Optional[list[str]] = None
    metadata: Optional[dict[str, Any]] = None


class GuidanceCreateRequest(BaseModel):
    guidance_type: GuidanceType = "keyword"
    text: str = Field(..., min_length=1)
    weight: float = Field(default=1.0, ge=0.0, le=10.0)
    metadata: dict[str, Any] = Field(default_factory=dict)


class AnswerListResponse(BaseModel):
    answers: list[AnswerItem]
    total: int
    page: int
    page_size: int


class AnswerEventListResponse(BaseModel):
    events: list[AnswerEvent]
    answers: list[AnswerItem] = Field(default_factory=list)
    total: int
    page: int
    page_size: int


class AnswerAnalyticsGroupRow(BaseModel):
    key: str
    label: str
    count: int


class AnswerEventStatsResponse(BaseModel):
    workspace: str
    total_events: int
    no_match: int
    avg_latency_ms: int
    timezone: str
    selected_answers: list[AnswerAnalyticsGroupRow] = Field(default_factory=list)
    queries: list[AnswerAnalyticsGroupRow] = Field(default_factory=list)
    sources: list[AnswerAnalyticsGroupRow] = Field(default_factory=list)
    modes: list[AnswerAnalyticsGroupRow] = Field(default_factory=list)
    dates: list[AnswerAnalyticsGroupRow] = Field(default_factory=list)
    hours: list[AnswerAnalyticsGroupRow] = Field(default_factory=list)


class ResolveRequest(BaseModel):
    query: str = Field(..., min_length=1)
    top_k: int = Field(default=5, ge=1, le=20)
    min_score: float = Field(default=0.18, ge=0.0, le=1.0)
    include_drafts: bool = False
    strategy: ResolveStrategy = "balanced"
    retrieval_mode: RetrievalMode = Field(
        default="hybrid",
        description="keyword keeps the deterministic scorer, hybrid adds answer vectors, graph_hybrid also uses FAQ graph paths, and llm_rerank lets the LLM choose only from candidate IDs.",
    )
    vector_top_k: int = Field(default=8, ge=1, le=50)
    llm_candidate_count: int = Field(default=5, ge=1, le=10)
    allowed_answer_ids: Optional[list[str]] = Field(
        default=None,
        description="Optional answer IDs allowed for candidate selection. Omitted means unrestricted.",
    )
    selection_policy: SelectionPolicy = Field(
        default="workspace",
        description=(
            "workspace follows the workspace precision setting, coverage keeps the "
            "legacy highest-score behavior, and precision may abstain on weak or "
            "ambiguous evidence."
        ),
    )


class ResolveCandidate(BaseModel):
    answer: AnswerItem
    score: float
    matched_guidance: list[str] = Field(default_factory=list)
    reason: str
    score_details: dict[str, float] = Field(default_factory=dict)
    selected_by: str = "keyword"
    graph_evidence: list[AnswerGraphEvidence] = Field(default_factory=list)


class ResolveResponse(BaseModel):
    selected_answer: Optional[AnswerItem]
    matched_id: Optional[str] = Field(
        default=None,
        description="Business ID returned by an ID-lookup FAQ. This is separate from the internal answer_id.",
    )
    confidence: float
    candidates: list[ResolveCandidate]
    trace_id: str
    rationale: str
    retrieval_mode: RetrievalMode = Field(
        default="hybrid",
        description="Compatibility alias for requested_retrieval_mode.",
    )
    requested_retrieval_mode: RetrievalMode = Field(
        default="hybrid",
        description="Retrieval mode requested by the caller.",
    )
    effective_retrieval_mode: RetrievalMode = Field(
        default="hybrid",
        description="Retrieval mode actually used after graph availability checks.",
    )
    graph_status: str = Field(
        default="not_requested",
        description="Graph lookup outcome such as graph_ready or graph_below_similarity.",
    )
    retrieval_fallback_reason: Optional[str] = Field(
        default=None,
        description="Reason graph_hybrid fell back to hybrid; null when no fallback occurred.",
    )
    selected_by: str = "keyword"
    alias_expansions: list[AnswerAliasExpansion] = Field(default_factory=list)
    selection_policy: Literal["coverage", "precision"] = "coverage"
    abstention_reason: Optional[str] = None
    clarification_question: Optional[str] = Field(
        default=None,
        description="A concise follow-up question when the FAQ intent is ambiguous.",
    )


class AnswerSearchRequest(BaseModel):
    query: str = Field(..., min_length=1)
    top_k: int = Field(default=5, ge=1, le=20)
    min_score: float = Field(default=0.18, ge=0.0, le=1.0)
    include_drafts: bool = False
    strategy: ResolveStrategy = "balanced"
    retrieval_mode: RetrievalMode = Field(
        default="hybrid",
        description="Direct API default remains hybrid; products can request graph_hybrid explicitly.",
    )
    vector_top_k: int = Field(default=8, ge=1, le=50)
    llm_candidate_count: int = Field(default=5, ge=1, le=10)
    allowed_answer_ids: Optional[list[str]] = None
    selection_policy: SelectionPolicy = "workspace"
    response_policy: Optional[DisplayPolicy] = None
    include_candidates: bool = True


class AnswerSearchResponse(BaseModel):
    matched: bool
    answer_id: Optional[str] = None
    matched_id: Optional[str] = Field(
        default=None,
        description="Business ID returned by an ID-lookup FAQ. This is separate from the internal answer_id.",
    )
    title: Optional[str] = None
    response: Optional[str] = None
    summary: Optional[str] = None
    full_content: Optional[str] = None
    display_policy: Optional[DisplayPolicy] = None
    content_format: Optional[ContentFormat] = None
    status: Optional[AnswerStatus] = None
    version: Optional[int] = None
    valid_from: Optional[str] = None
    valid_until: Optional[str] = None
    confidence: float = 0.0
    tags: list[str] = Field(default_factory=list)
    assets: list[AnswerAsset] = Field(default_factory=list)
    source_type: Optional[str] = None
    source_uri: Optional[str] = None
    candidates: list[ResolveCandidate] = Field(default_factory=list)
    trace_id: str
    rationale: str
    retrieval_mode: RetrievalMode = Field(
        default="hybrid",
        description="Compatibility alias for requested_retrieval_mode.",
    )
    requested_retrieval_mode: RetrievalMode = Field(
        default="hybrid",
        description="Retrieval mode requested by the caller.",
    )
    effective_retrieval_mode: RetrievalMode = Field(
        default="hybrid",
        description="Retrieval mode actually used after graph availability checks.",
    )
    graph_status: str = Field(
        default="not_requested",
        description="Graph lookup outcome such as graph_ready or graph_below_similarity.",
    )
    retrieval_fallback_reason: Optional[str] = Field(
        default=None,
        description="Reason graph_hybrid fell back to hybrid; null when no fallback occurred.",
    )
    selected_by: str = "keyword"
    alias_expansions: list[AnswerAliasExpansion] = Field(default_factory=list)
    selection_policy: Literal["coverage", "precision"] = "coverage"
    abstention_reason: Optional[str] = None
    clarification_question: Optional[str] = None


class AnswerViewRequest(BaseModel):
    query: Optional[str] = None
    metadata: dict[str, Any] = Field(default_factory=dict)


class AnswerFeedbackRequest(BaseModel):
    query: Optional[str] = None
    helpful: Optional[bool] = None
    note: Optional[str] = None
    selected_alternative_id: Optional[str] = None
    metadata: dict[str, Any] = Field(default_factory=dict)


class SourceGuidanceCandidate(BaseModel):
    guidance_type: GuidanceType = "keyword"
    text: str = Field(..., min_length=1)
    weight: float = Field(default=1.0, ge=0.0, le=10.0)
    source: Optional[str] = None
    metadata: dict[str, Any] = Field(default_factory=dict)


class GuidanceSuggestRequest(BaseModel):
    query_examples: list[str] = Field(default_factory=list)
    max_suggestions: int = Field(default=12, ge=1, le=30)
    use_llm: bool = False


class GuidanceSuggestResponse(BaseModel):
    suggestions: list[SourceGuidanceCandidate] = Field(default_factory=list)
    mode: str = "heuristic"


class AnswerSourceDraftRequest(BaseModel):
    answer_id: Optional[str] = Field(default=None, description="Optional stable answer id. Defaults to ANS-<uuid>.")
    source_type: Literal["plain", "markdown", "html", "url", "file", "excel", "structured"] = "plain"
    source_uri: Optional[str] = None
    file_name: Optional[str] = None
    title: str = Field(..., min_length=1, max_length=500)
    body: str = Field(..., min_length=1)
    approved_summary: Optional[str] = None
    content_format: ContentFormat = "markdown"
    display_policy: DisplayPolicy = "both"
    status: AnswerStatus = "draft"
    priority: int = 0
    tags: list[str] = Field(default_factory=list)
    source_profile: dict[str, Any] = Field(default_factory=dict)
    guidance: list[SourceGuidanceCandidate] = Field(default_factory=list)
    metadata: dict[str, Any] = Field(default_factory=dict)


class AnswerSourceDraftResponse(BaseModel):
    answer: AnswerItem
    guidance: list[AnswerGuidance] = Field(default_factory=list)
    snapshot: Optional["AnswerSourceSnapshot"] = None
    source_link: Optional["AnswerSourceLink"] = None


class AnswerSourceSnapshot(BaseModel):
    snapshot_id: str
    workspace: str
    source_type: str
    source_uri: Optional[str] = None
    file_name: Optional[str] = None
    title: Optional[str] = None
    content_hash: str
    content_length: int = 0
    content_preview: Optional[str] = None
    profile: dict[str, Any] = Field(default_factory=dict)
    metadata: dict[str, Any] = Field(default_factory=dict)
    created_answer_ids: list[str] = Field(default_factory=list)
    status: str = "captured"
    task_id: Optional[str] = None
    create_time: Optional[str] = None


class AnswerSourceLink(BaseModel):
    link_id: str
    workspace: str
    answer_id: str
    answer_version: int = 1
    snapshot_id: str
    link_type: str = "created_from"
    metadata: dict[str, Any] = Field(default_factory=dict)
    create_time: Optional[str] = None
    snapshot: Optional[AnswerSourceSnapshot] = None


class StructuredDataset(BaseModel):
    answer_id: str
    title: str
    status: AnswerStatus
    source_type: str = "-"
    source_uri: Optional[str] = None
    kind: str = "text"
    columns: list[str] = Field(default_factory=list)
    row_count: int = 0
    sample_rows: list[dict[str, Any]] = Field(default_factory=list)
    metadata: dict[str, Any] = Field(default_factory=dict)


class StructuredFilter(BaseModel):
    field: str = Field(..., min_length=1)
    operator: StructuredOperator = "contains"
    value: str = ""


class StructuredQueryRequest(BaseModel):
    answer_id: str = Field(..., min_length=1)
    filters: list[StructuredFilter] = Field(default_factory=list)
    limit: int = Field(default=20, ge=1, le=100)
    preview_only: bool = False


class StructuredQueryResponse(BaseModel):
    answer_id: str
    title: str
    pseudo_sql: str
    columns: list[str]
    rows: list[dict[str, Any]]
    row_count: int
    preview_only: bool = False


class StructuredFieldProfile(BaseModel):
    name: str
    inferred_type: str = "text"
    null_count: int = 0
    null_rate: float = 0.0
    distinct_count: int = 0
    sample_values: list[str] = Field(default_factory=list)
    semantic_role: str = "metadata"
    confidence: float = 0.0


class StructuredProfileRequest(BaseModel):
    source_type: StructuredSourceType = "csv"
    raw_content: str = Field(..., min_length=1)
    source_uri: Optional[str] = None
    sample_limit: int = Field(default=20, ge=1, le=100)


class StructuredProfileResponse(BaseModel):
    source_type: StructuredSourceType
    kind: str
    row_count: int
    columns: list[str] = Field(default_factory=list)
    fields: list[StructuredFieldProfile] = Field(default_factory=list)
    sample_rows: list[dict[str, Any]] = Field(default_factory=list)
    mapping_suggestions: dict[str, Optional[str]] = Field(default_factory=dict)
    warnings: list[str] = Field(default_factory=list)


class ExcelSheetInfo(BaseModel):
    name: str
    max_row: int = 0
    max_column: int = 0


class ExcelPreviewResponse(BaseModel):
    file_name: str
    file_size: int
    sheets: list[ExcelSheetInfo] = Field(default_factory=list)
    selected_sheet: str
    header_row: int
    data_start_row: int
    row_count: int
    row_limit: int = MAX_EXCEL_ROWS_PER_PREVIEW
    truncated: bool = False
    columns: list[str] = Field(default_factory=list)
    raw_content: str
    source_uri: str
    profile: StructuredProfileResponse
    warnings: list[str] = Field(default_factory=list)


class BatchGuidanceEnrichmentConfig(BaseModel):
    enabled: bool = Field(
        default=False,
        description=(
            "After FAQ drafts are created, add LLM-generated representative questions, "
            "keywords, and synonyms without changing answer content."
        ),
    )
    scope: GuidanceEnrichmentScope = Field(
        default="coverage",
        description=(
            "missing_or_weak enriches only answers missing a question or keyword/synonym; "
            "coverage also checks representative-question, synonym, colloquial, and bilingual "
            "search coverage; all enriches every generated answer."
        ),
    )
    batch_size: int = Field(
        default=10,
        ge=1,
        le=20,
        description="Number of FAQ answers included in each LLM request.",
    )
    max_suggestions: int = Field(
        default=5,
        ge=1,
        le=8,
        description="Maximum new finding hints accepted per FAQ answer.",
    )


class StructuredMaterializeRequest(BaseModel):
    source_type: StructuredSourceType = "csv"
    raw_content: str = Field(..., min_length=1)
    title: str = Field(..., min_length=1, max_length=500)
    approved_summary: Optional[str] = None
    source_uri: Optional[str] = None
    file_name: Optional[str] = None
    status: AnswerStatus = "draft"
    priority: int = 0
    tags: list[str] = Field(default_factory=list)
    mapping: dict[str, str] = Field(default_factory=dict)
    guidance_columns: list[str] = Field(default_factory=list)
    materialization_mode: StructuredMaterializationMode = "table_as_dataset"
    conversion_purpose: StructuredConversionPurpose = "faq"
    source_truncated: bool = False
    llm_guidance_enrichment: BatchGuidanceEnrichmentConfig = Field(
        default_factory=BatchGuidanceEnrichmentConfig
    )
    metadata: dict[str, Any] = Field(default_factory=dict)


class StructuredIdLookupValidation(BaseModel):
    enabled: bool = False
    ready: bool = True
    id_column: Optional[str] = None
    searchable_columns: list[str] = Field(default_factory=list)
    row_count: int = 0
    valid_id_count: int = 0
    blank_id_rows: list[int] = Field(default_factory=list)
    duplicate_ids: list[str] = Field(default_factory=list)
    ambiguous_detail_groups: list[dict[str, Any]] = Field(default_factory=list)
    warnings: list[str] = Field(default_factory=list)


class StructuredMaterializeResponse(BaseModel):
    answer: AnswerItem
    dataset: StructuredDataset
    answers: list[AnswerItem] = Field(default_factory=list)
    datasets: list[StructuredDataset] = Field(default_factory=list)
    answer_count: int = 0
    answers_truncated: bool = False
    profile: StructuredProfileResponse
    guidance: list[AnswerGuidance] = Field(default_factory=list)
    guidance_count: int = 0
    guidance_truncated: bool = False
    snapshot: Optional[AnswerSourceSnapshot] = None
    source_link: Optional[AnswerSourceLink] = None
    validation: StructuredIdLookupValidation = Field(default_factory=StructuredIdLookupValidation)
    guidance_enrichment_task_id: Optional[str] = None
    guidance_enrichment_stream_url: Optional[str] = None
    vector_rebuild_task_id: Optional[str] = None
    vector_rebuild_stream_url: Optional[str] = None


class StructuredLookupLog(BaseModel):
    log_id: str
    workspace: str
    dataset_id: str
    dataset_title: Optional[str] = None
    pseudo_sql: str
    filters: list[dict[str, Any]] = Field(default_factory=list)
    result_count: int = 0
    preview_only: bool = False
    latency_ms: float = 0.0
    metadata: dict[str, Any] = Field(default_factory=dict)
    create_time: Optional[str] = None


class SourceConnector(BaseModel):
    connector_id: str
    workspace: str
    name: str
    connector_type: SourceConnectorType
    status: SourceConnectorStatus = "draft"
    config: dict[str, Any] = Field(default_factory=dict)
    auth_ref: Optional[str] = None
    refresh_policy: dict[str, Any] = Field(default_factory=dict)
    enabled: bool = True
    metadata: dict[str, Any] = Field(default_factory=dict)
    create_time: Optional[str] = None
    update_time: Optional[str] = None


class SourceConnectorCreateRequest(BaseModel):
    connector_id: Optional[str] = Field(default=None, description="Optional stable connector id. Defaults to conn-<uuid>.")
    name: str = Field(..., min_length=1, max_length=300)
    connector_type: SourceConnectorType
    status: SourceConnectorStatus = "draft"
    config: dict[str, Any] = Field(default_factory=dict)
    auth_ref: Optional[str] = None
    refresh_policy: dict[str, Any] = Field(default_factory=dict)
    enabled: bool = True
    metadata: dict[str, Any] = Field(default_factory=dict)


class SourceConnectorUpdateRequest(BaseModel):
    name: Optional[str] = Field(default=None, min_length=1, max_length=300)
    status: Optional[SourceConnectorStatus] = None
    config: Optional[dict[str, Any]] = None
    auth_ref: Optional[str] = None
    refresh_policy: Optional[dict[str, Any]] = None
    enabled: Optional[bool] = None
    metadata: Optional[dict[str, Any]] = None


class SourceConnectorSampleRequest(BaseModel):
    limit: int = Field(default=20, ge=1, le=200)


class SourceConnectorSampleResponse(BaseModel):
    connector_id: str
    connector_type: SourceConnectorType
    source_type: StructuredSourceType
    source_uri: Optional[str] = None
    raw_content: str
    rows: list[dict[str, Any]] = Field(default_factory=list)
    columns: list[str] = Field(default_factory=list)
    row_count: int = 0
    row_limit: int = 1000
    truncated: bool = False
    warnings: list[str] = Field(default_factory=list)


class SourceConnectorMappingPreviewRequest(BaseModel):
    mapping: dict[str, str] = Field(default_factory=dict)
    guidance_columns: list[str] = Field(default_factory=list)
    materialization_mode: StructuredMaterializationMode = "table_as_dataset"
    conversion_purpose: StructuredConversionPurpose = "faq"


class SourceConnectorMappingPreviewResponse(BaseModel):
    connector: SourceConnector
    sample: SourceConnectorSampleResponse
    profile: StructuredProfileResponse
    mapping: dict[str, str] = Field(default_factory=dict)
    guidance_columns: list[str] = Field(default_factory=list)
    materialization_modes: list[StructuredMaterializationMode] = Field(default_factory=list)
    validation: StructuredIdLookupValidation = Field(default_factory=StructuredIdLookupValidation)


class SourceConnectorMaterializeRequest(BaseModel):
    title: Optional[str] = None
    mapping: dict[str, str] = Field(default_factory=dict)
    guidance_columns: list[str] = Field(default_factory=list)
    materialization_mode: StructuredMaterializationMode = "table_as_dataset"
    conversion_purpose: StructuredConversionPurpose = "faq"
    status: AnswerStatus = "draft"
    tags: list[str] = Field(default_factory=list)
    llm_guidance_enrichment: BatchGuidanceEnrichmentConfig = Field(
        default_factory=BatchGuidanceEnrichmentConfig
    )
    metadata: dict[str, Any] = Field(default_factory=dict)


class SourceConnectorMaterializeResponse(BaseModel):
    connector: SourceConnector
    sample: SourceConnectorSampleResponse
    materialized: StructuredMaterializeResponse


def _answer_asset_content_url(row: dict[str, Any]) -> Optional[str]:
    storage_type = str(row.get("storage_type") or "")
    storage_uri = str(row.get("storage_uri") or "").strip()
    if storage_type in {"external", "s3"}:
        return storage_uri or None
    if storage_type == "local":
        return (
            f"/api/answers/{str(row['answer_id'])}/assets/"
            f"{str(row['asset_id'])}/content"
        )
    return None


def _answer_asset_from_row(row: dict[str, Any]) -> AnswerAsset:
    metadata = _coerce_json(row.get("metadata"), {})
    if not isinstance(metadata, dict):
        metadata = {}
    return AnswerAsset(
        asset_id=str(row["asset_id"]),
        workspace=str(row["workspace"]),
        answer_id=str(row["answer_id"]),
        answer_version=int(row.get("answer_version") or 1),
        asset_type=row.get("asset_type") or "file",
        storage_type=row.get("storage_type") or "external",
        storage_uri=row.get("storage_uri"),
        content_url=_answer_asset_content_url(row),
        file_name=row.get("file_name"),
        mime_type=row.get("mime_type"),
        file_size=int(row.get("file_size") or 0),
        caption=row.get("caption"),
        alt_text=row.get("alt_text"),
        search_text=row.get("search_text"),
        content_text=row.get("content_text"),
        display_order=int(row.get("display_order") or 0),
        is_active=bool(row.get("is_active", True)),
        metadata=metadata,
        create_time=_iso(row.get("create_time")),
        update_time=_iso(row.get("update_time")),
    )


def _answer_asset_search_content(assets: list[AnswerAsset]) -> str:
    return "\n".join(
        " ".join(
            part.strip()
            for part in [
                asset.caption or "",
                asset.alt_text or "",
                asset.search_text or "",
                asset.content_text or "",
            ]
            if part and part.strip()
        )
        for asset in assets
        if asset.is_active
    ).strip()


async def _answer_assets_by_answer(
    db,
    workspace: str,
    answer_ids: list[str],
    *,
    include_inactive: bool = False,
) -> dict[str, list[AnswerAsset]]:
    if not answer_ids:
        return {}
    active_filter = "" if include_inactive else "AND is_active = TRUE"
    rows = await db.query(
        f"""
        SELECT asset_id, workspace, answer_id, answer_version, asset_type,
               storage_type, storage_uri, file_name, mime_type, file_size,
               caption, alt_text, search_text, content_text, display_order,
               is_active, metadata, create_time, update_time
        FROM LIGHTRAG_ANSWER_ASSETS
        WHERE workspace = $1 AND answer_id = ANY($2::text[])
          {active_filter}
        ORDER BY answer_id, display_order, create_time, asset_id
        """,
        [workspace, answer_ids],
        multirows=True,
    )
    result: dict[str, list[AnswerAsset]] = {}
    for row in rows or []:
        asset = _answer_asset_from_row(dict(row))
        result.setdefault(asset.answer_id, []).append(asset)
    return result


async def _hydrate_answer_assets(
    db,
    workspace: str,
    answers: list[AnswerItem],
) -> list[AnswerItem]:
    assets_by_answer = await _answer_assets_by_answer(
        db,
        workspace,
        [answer.answer_id for answer in answers],
    )
    for answer in answers:
        answer.assets = assets_by_answer.get(answer.answer_id, [])
    return answers


def _answer_from_row(row: dict[str, Any]) -> AnswerItem:
    tags = _coerce_json(row.get("tags"), [])
    if isinstance(tags, str):
        tags = [tags]
    if not isinstance(tags, list):
        tags = []
    metadata = _coerce_json(row.get("metadata"), {})
    if not isinstance(metadata, dict):
        metadata = {}
    return AnswerItem(
        answer_id=str(row["answer_id"]),
        workspace=str(row["workspace"]),
        title=str(row["title"]),
        body=str(row["body"]),
        approved_summary=row.get("approved_summary"),
        content_format=row.get("content_format") or "markdown",
        display_policy=row.get("display_policy") or "both",
        status=row.get("status") or "draft",
        version=int(row.get("version") or 1),
        valid_from=_iso(row.get("valid_from")),
        valid_until=_iso(row.get("valid_until")),
        priority=int(row.get("priority") or 0),
        tags=tags,
        metadata=metadata,
        assets=[
            item if isinstance(item, AnswerAsset) else AnswerAsset.model_validate(item)
            for item in (row.get("assets") or [])
        ],
        publish_time=_iso(row.get("publish_time")),
        create_time=_iso(row.get("create_time")),
        update_time=_iso(row.get("update_time")),
    )


def _guidance_from_row(row: dict[str, Any]) -> AnswerGuidance:
    metadata = _coerce_json(row.get("metadata"), {})
    if not isinstance(metadata, dict):
        metadata = {}
    return AnswerGuidance(
        guidance_id=str(row["guidance_id"]),
        workspace=str(row["workspace"]),
        answer_id=str(row["answer_id"]),
        guidance_type=row.get("guidance_type") or "keyword",
        text=str(row["text"]),
        weight=float(row.get("weight") or 1.0),
        metadata=metadata,
        create_time=_iso(row.get("create_time")),
    )


def _alias_group_from_row(row: dict[str, Any]) -> AnswerAliasGroup:
    aliases = _coerce_json(row.get("aliases"), [])
    if not isinstance(aliases, list):
        aliases = []
    metadata = _coerce_json(row.get("metadata"), {})
    if not isinstance(metadata, dict):
        metadata = {}
    return AnswerAliasGroup(
        alias_id=str(row["alias_id"]),
        workspace=str(row["workspace"]),
        canonical_term=str(row["canonical_term"]),
        aliases=[str(item) for item in aliases if str(item).strip()],
        enabled=bool(row.get("enabled", True)),
        source=str(row.get("source") or "workspace"),
        metadata=metadata,
        create_time=_iso(row.get("create_time")),
        update_time=_iso(row.get("update_time")),
    )


def _term_candidate_from_row(row: dict[str, Any]) -> AnswerTermCandidate:
    aliases = _coerce_json(row.get("aliases"), [])
    evidence = _coerce_json(row.get("evidence"), [])
    metadata = _coerce_json(row.get("metadata"), {})
    return AnswerTermCandidate(
        candidate_id=str(row["candidate_id"]),
        workspace=str(row["workspace"]),
        canonical_term=str(row["canonical_term"]),
        aliases=[str(item) for item in aliases if str(item).strip()]
        if isinstance(aliases, list)
        else [],
        term_type=str(row.get("term_type") or "synonym"),
        status=str(row.get("status") or "suggested"),
        confidence=_safe_float(row.get("confidence"), 0.0, 0.0, 1.0),
        rationale=str(row.get("rationale") or "") or None,
        evidence=[
            dict(item) for item in evidence if isinstance(item, dict)
        ]
        if isinstance(evidence, list)
        else [],
        source=str(row.get("source") or "llm"),
        metadata=metadata if isinstance(metadata, dict) else {},
        create_time=_iso(row.get("create_time")),
        update_time=_iso(row.get("update_time")),
    )


def _event_from_row(row: dict[str, Any]) -> AnswerEvent:
    candidate_ids = _coerce_json(row.get("candidate_ids"), [])
    if not isinstance(candidate_ids, list):
        candidate_ids = []
    scores = _coerce_json(row.get("scores"), {})
    if not isinstance(scores, dict):
        scores = {}
    metadata = _coerce_json(row.get("metadata"), {})
    if not isinstance(metadata, dict):
        metadata = {}
    return AnswerEvent(
        event_id=str(row["event_id"]),
        workspace=str(row["workspace"]),
        event_type=str(row["event_type"]),
        query=row.get("query"),
        selected_answer_id=row.get("selected_answer_id"),
        candidate_ids=[str(item) for item in candidate_ids],
        scores={str(key): float(value) for key, value in scores.items() if isinstance(value, (int, float))},
        metadata=metadata,
        create_time=_iso(row.get("create_time")),
    )


def _analytics_group_from_row(row: dict[str, Any]) -> AnswerAnalyticsGroupRow:
    return AnswerAnalyticsGroupRow(
        key=str(row.get("key") or "-"),
        label=str(row.get("label") or row.get("key") or "-"),
        count=int(row.get("count") or 0),
    )


def _safe_timezone(value: str | None) -> str:
    timezone_name = (value or "Asia/Seoul").strip()
    if not timezone_name or len(timezone_name) > 64:
        return "Asia/Seoul"
    if not re.fullmatch(r"[A-Za-z0-9_./+-]+", timezone_name):
        return "Asia/Seoul"
    return timezone_name


def _add_answer_event_filters(
    params: list[Any],
    where: list[str],
    *,
    event_type: Optional[str] = None,
    selected_answer_id: Optional[str] = None,
    query_text: Optional[str] = None,
    match_status: Optional[str] = None,
    source: Optional[str] = None,
    mode: Optional[str] = None,
    date_key: Optional[str] = None,
    hour_key: Optional[str] = None,
    date_from: Optional[datetime] = None,
    date_to: Optional[datetime] = None,
    min_confidence: Optional[float] = None,
    max_latency_ms: Optional[int] = None,
    search: Optional[str] = None,
    timezone_name: str = "Asia/Seoul",
) -> None:
    source_expr = (
        "CASE WHEN e.selected_answer_id IS NULL THEN '__none__' "
        "ELSE COALESCE(a.metadata->>'source_type', a.metadata->>'created_from', "
        "a.metadata->>'materialization_mode', 'manual') END"
    )
    mode_expr = "COALESCE(e.metadata->>'mode', e.metadata->>'retrieval_mode', '-')"

    if event_type and event_type != "all":
        params.append(event_type)
        where.append(f"e.event_type = ${len(params)}")

    if selected_answer_id and selected_answer_id != "all":
        if selected_answer_id == "__none__":
            where.append("e.selected_answer_id IS NULL")
        else:
            params.append(selected_answer_id)
            where.append(f"e.selected_answer_id = ${len(params)}")

    if query_text:
        params.append(query_text)
        where.append(f"COALESCE(e.query, '') = ${len(params)}")

    if match_status == "matched":
        where.append("e.selected_answer_id IS NOT NULL")
    elif match_status == "no_match":
        where.append("e.selected_answer_id IS NULL")

    if source and source != "all":
        params.append(source)
        where.append(f"{source_expr} = ${len(params)}")

    if mode and mode != "all":
        params.append(mode)
        where.append(f"{mode_expr} = ${len(params)}")

    timezone_placeholder: Optional[str] = None
    if date_key or hour_key:
        params.append(timezone_name)
        timezone_placeholder = f"${len(params)}"

    if date_key:
        params.append(date_key)
        where.append(
            f"TO_CHAR(e.create_time AT TIME ZONE {timezone_placeholder}, 'YYYY-MM-DD') = ${len(params)}"
        )

    if hour_key:
        params.append(hour_key)
        where.append(
            f"TO_CHAR(e.create_time AT TIME ZONE {timezone_placeholder}, 'HH24:00') = ${len(params)}"
        )

    if date_from:
        params.append(date_from)
        where.append(f"e.create_time >= ${len(params)}")

    if date_to:
        params.append(date_to)
        where.append(f"e.create_time <= ${len(params)}")

    if min_confidence is not None:
        params.append(min_confidence)
        where.append(
            "("
            "e.selected_answer_id IS NOT NULL AND "
            "COALESCE(NULLIF(e.scores->>e.selected_answer_id, ''), '0') ~ '^[0-9]+(\\.[0-9]+)?$' AND "
            f"(e.scores->>e.selected_answer_id)::DOUBLE PRECISION >= ${len(params)}"
            ")"
        )

    if max_latency_ms is not None:
        params.append(max_latency_ms)
        where.append(
            "("
            "COALESCE(NULLIF(e.metadata->>'latency_ms', ''), '0') ~ '^[0-9]+(\\.[0-9]+)?$' AND "
            f"(e.metadata->>'latency_ms')::DOUBLE PRECISION <= ${len(params)}"
            ")"
        )

    if search:
        term = f"%{search.strip().lower()}%"
        params.append(term)
        search_placeholder = f"${len(params)}"
        where.append(
            "("
            f"LOWER(COALESCE(e.query, '')) LIKE {search_placeholder} OR "
            f"LOWER(COALESCE(e.event_id, '')) LIKE {search_placeholder} OR "
            f"LOWER(COALESCE(e.selected_answer_id, '')) LIKE {search_placeholder} OR "
            f"LOWER(COALESCE(a.title, '')) LIKE {search_placeholder} OR "
            f"LOWER(COALESCE(a.metadata->>'source_type', '')) LIKE {search_placeholder} OR "
            f"LOWER(COALESCE(a.metadata->>'source_uri', '')) LIKE {search_placeholder} OR "
            f"LOWER(COALESCE(e.metadata->>'mode', '')) LIKE {search_placeholder} OR "
            f"LOWER(COALESCE(e.metadata->>'retrieval_mode', '')) LIKE {search_placeholder}"
            ")"
        )


def _structured_lookup_log_from_row(row: dict[str, Any]) -> StructuredLookupLog:
    filters = _coerce_json(row.get("filters"), [])
    if not isinstance(filters, list):
        filters = []
    metadata = _coerce_json(row.get("metadata"), {})
    if not isinstance(metadata, dict):
        metadata = {}
    return StructuredLookupLog(
        log_id=str(row["log_id"]),
        workspace=str(row["workspace"]),
        dataset_id=str(row["dataset_id"]),
        dataset_title=row.get("dataset_title"),
        pseudo_sql=str(row.get("pseudo_sql") or ""),
        filters=[item for item in filters if isinstance(item, dict)],
        result_count=int(row.get("result_count") or 0),
        preview_only=bool(row.get("preview_only")),
        latency_ms=float(row.get("latency_ms") or 0.0),
        metadata=metadata,
        create_time=_iso(row.get("create_time")),
    )


def _source_connector_from_row(row: dict[str, Any]) -> SourceConnector:
    config = _coerce_json(row.get("config"), {})
    if not isinstance(config, dict):
        config = {}
    refresh_policy = _coerce_json(row.get("refresh_policy"), {})
    if not isinstance(refresh_policy, dict):
        refresh_policy = {}
    metadata = _coerce_json(row.get("metadata"), {})
    if not isinstance(metadata, dict):
        metadata = {}
    return SourceConnector(
        connector_id=str(row["connector_id"]),
        workspace=str(row["workspace"]),
        name=str(row.get("name") or ""),
        connector_type=row.get("connector_type") or "manual_table",
        status=row.get("status") or "draft",
        config=config,
        auth_ref=row.get("auth_ref"),
        refresh_policy=refresh_policy,
        enabled=bool(row.get("enabled")),
        metadata=metadata,
        create_time=_iso(row.get("create_time")),
        update_time=_iso(row.get("update_time")),
    )


def _content_hash(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def _source_snapshot_from_row(row: dict[str, Any]) -> AnswerSourceSnapshot:
    profile = _coerce_json(row.get("profile"), {})
    if not isinstance(profile, dict):
        profile = {}
    metadata = _coerce_json(row.get("metadata"), {})
    if not isinstance(metadata, dict):
        metadata = {}
    created_answer_ids = _coerce_json(row.get("created_answer_ids"), [])
    if not isinstance(created_answer_ids, list):
        created_answer_ids = []
    preview = row.get("content_preview")
    if preview is None and row.get("raw_content") is not None:
        preview = str(row.get("raw_content"))[:600]
    return AnswerSourceSnapshot(
        snapshot_id=str(row["snapshot_id"]),
        workspace=str(row["workspace"]),
        source_type=str(row.get("source_type") or "plain"),
        source_uri=row.get("source_uri"),
        file_name=row.get("file_name"),
        title=row.get("title"),
        content_hash=str(row.get("content_hash") or ""),
        content_length=int(row.get("content_length") or 0),
        content_preview=preview,
        profile=profile,
        metadata=metadata,
        created_answer_ids=[str(item) for item in created_answer_ids],
        status=str(row.get("status") or "captured"),
        task_id=row.get("task_id"),
        create_time=_iso(row.get("create_time")),
    )


def _source_link_from_row(
    row: dict[str, Any],
    snapshot: Optional[AnswerSourceSnapshot] = None,
) -> AnswerSourceLink:
    metadata = _coerce_json(row.get("metadata"), {})
    if not isinstance(metadata, dict):
        metadata = {}
    return AnswerSourceLink(
        link_id=str(row["link_id"]),
        workspace=str(row["workspace"]),
        answer_id=str(row["answer_id"]),
        answer_version=int(row.get("answer_version") or 1),
        snapshot_id=str(row["snapshot_id"]),
        link_type=str(row.get("link_type") or "created_from"),
        metadata=metadata,
        create_time=_iso(row.get("create_time")),
        snapshot=snapshot,
    )


def _structured_profile(answer: AnswerItem) -> dict[str, Any]:
    if (
        isinstance(answer.metadata, dict)
        and answer.metadata.get("materialization_mode") == "row_per_answer"
        and isinstance(answer.metadata.get("source_row"), dict)
    ):
        source_row = answer.metadata["source_row"]
        return {
            "kind": "structured_row",
            "columns": [str(column) for column in source_row.keys()],
            "row_count": 1,
        }

    profile = answer.metadata.get("source_profile") if isinstance(answer.metadata, dict) else {}
    if not isinstance(profile, dict):
        profile = {}
    structured = profile.get("structured")
    if not isinstance(structured, dict):
        structured = profile
    columns = structured.get("columns") or profile.get("columns") or []
    if not isinstance(columns, list):
        columns = []
    row_count = structured.get("rowCount") or structured.get("row_count") or profile.get("rowCount") or profile.get("row_count") or 0
    try:
        row_count = int(row_count)
    except (TypeError, ValueError):
        row_count = 0
    return {
        "kind": structured.get("kind") or profile.get("kind") or answer.metadata.get("source_type") or "text",
        "columns": [str(column) for column in columns if str(column).strip()],
        "row_count": row_count,
    }


def _row_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return json.dumps(value, ensure_ascii=False)


def _excel_cell_value(cell: Any) -> Any:
    value = cell.value
    if isinstance(value, datetime) and getattr(cell, "is_date", False):
        number_format = str(getattr(cell, "number_format", "") or "").lower()
        has_time_format = (
            "am/pm" in number_format
            or "a/p" in number_format
            or bool(re.search(r"[hs]", number_format))
        )
        if not has_time_format:
            return value.date().isoformat()
    return _row_value(value)


async def _read_upload_limited(file: UploadFile, max_bytes: int) -> tuple[bytes, int]:
    total = 0
    chunks: list[bytes] = []
    while True:
        chunk = await file.read(1024 * 1024)
        if not chunk:
            break
        total += len(chunk)
        if total > max_bytes:
            raise HTTPException(
                status_code=413,
                detail=f"Excel upload exceeds the {max_bytes // 1024 // 1024}MB limit",
            )
        chunks.append(chunk)
    return b"".join(chunks), total


async def _read_answer_asset_upload(file: UploadFile) -> tuple[bytes, int]:
    total = 0
    chunks: list[bytes] = []
    while True:
        chunk = await file.read(1024 * 1024)
        if not chunk:
            break
        total += len(chunk)
        if total > MAX_ANSWER_ASSET_UPLOAD_BYTES:
            raise HTTPException(
                status_code=413,
                detail=(
                    "Answer asset upload exceeds the "
                    f"{MAX_ANSWER_ASSET_UPLOAD_BYTES // 1024 // 1024}MB limit"
                ),
            )
        chunks.append(chunk)
    return b"".join(chunks), total


def _answer_asset_root() -> Path:
    configured = Path(
        os.getenv("ANSWER_ASSET_DIR", "rag_storage/answer_assets")
    ).expanduser()
    root = configured if configured.is_absolute() else Path.cwd() / configured
    return root.resolve()


def _safe_asset_segment(value: Any, fallback: str) -> str:
    name = Path(str(value or "")).name.strip()
    safe = re.sub(r"[^\w.\-]+", "_", name, flags=re.UNICODE).strip("._")
    return safe[:180] or fallback


def _local_answer_asset_path(storage_uri: str) -> Path:
    root = _answer_asset_root()
    candidate = (root / storage_uri).resolve()
    try:
        candidate.relative_to(root)
    except ValueError as exc:
        raise HTTPException(status_code=409, detail="Stored asset path is invalid") from exc
    return candidate


def _validate_external_asset_url(value: Any) -> str:
    url = str(value or "").strip()
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        raise HTTPException(
            status_code=422,
            detail="external_url must be an absolute HTTP or HTTPS URL",
        )
    return url


def _infer_answer_asset_type(
    file_name: str,
    mime_type: Optional[str],
) -> AnswerAssetType:
    mime = str(mime_type or "").lower()
    suffix = Path(file_name).suffix.lower()
    if mime.startswith("image/"):
        return "image"
    if mime.startswith("video/"):
        return "video"
    if mime.startswith("audio/"):
        return "audio"
    if suffix in {".csv", ".xlsx", ".xls"}:
        return "table"
    return "file"


def _validate_answer_asset_file(file_name: str) -> str:
    safe_name = _safe_asset_segment(file_name, f"asset-{uuid.uuid4().hex}")
    suffix = Path(safe_name).suffix.lower()
    if suffix not in SUPPORTED_ANSWER_ASSET_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_ANSWER_ASSET_EXTENSIONS))
        raise HTTPException(
            status_code=415,
            detail=f"Unsupported answer asset type. Supported extensions: {supported}",
        )
    return safe_name


def _unique_excel_column_name(value: Any, index: int, seen: set[str]) -> str:
    base = str(_row_value(value) or "").strip() or f"Column {index}"
    if len(base) > 120:
        base = base[:120].strip()
    name = base
    suffix = 2
    while name in seen:
        name = f"{base}_{suffix}"
        suffix += 1
    seen.add(name)
    return name


def _extract_excel_rows(
    content: bytes,
    *,
    sheet_name: Optional[str] = None,
    header_row: int = 1,
    data_start_row: Optional[int] = None,
    max_rows: int = MAX_EXCEL_ROWS_PER_PREVIEW,
) -> tuple[list[ExcelSheetInfo], str, int, int, list[dict[str, Any]], list[str], list[str], bool]:
    try:
        from openpyxl import load_workbook  # type: ignore
        from openpyxl.utils.cell import range_boundaries  # type: ignore
    except ImportError as exc:
        raise HTTPException(
            status_code=500,
            detail="openpyxl is required to process Excel sources",
        ) from exc

    if header_row < 1:
        raise HTTPException(status_code=400, detail="header_row must be greater than or equal to 1")

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Invalid Excel workbook: {exc}") from exc

    try:
        sheets: list[ExcelSheetInfo] = []
        for worksheet in workbook.worksheets:
            max_row = worksheet.max_row or 0
            max_column = worksheet.max_column or 0
            if not max_row or not max_column:
                try:
                    _, _, max_column, max_row = range_boundaries(
                        worksheet.calculate_dimension(force=True)
                    )
                except (TypeError, ValueError):
                    max_row = max_row or 0
                    max_column = max_column or 0
            sheets.append(
                ExcelSheetInfo(
                    name=worksheet.title,
                    max_row=max_row,
                    max_column=max_column,
                )
            )
        if not sheets:
            raise HTTPException(status_code=400, detail="Excel workbook has no sheets")

        selected_sheet = sheet_name if sheet_name in workbook.sheetnames else workbook.sheetnames[0]
        if sheet_name and sheet_name not in workbook.sheetnames:
            raise HTTPException(status_code=400, detail=f"Sheet '{sheet_name}' was not found")

        worksheet = workbook[selected_sheet]
        start_row = data_start_row if data_start_row and data_start_row > header_row else header_row + 1
        warnings: list[str] = []

        header_values = next(
            worksheet.iter_rows(
                min_row=header_row,
                max_row=header_row,
                values_only=True,
            ),
            None,
        )
        if not header_values:
            raise HTTPException(status_code=400, detail="Excel sheet must include a header row")

        seen_columns: set[str] = set()
        columns = [
            _unique_excel_column_name(value, index + 1, seen_columns)
            for index, value in enumerate(header_values[:MAX_EXCEL_COLUMNS_PER_PREVIEW])
        ]
        while columns and columns[-1].startswith("Column ") and len(columns[-1].split()) == 2:
            columns.pop()
        if not columns:
            raise HTTPException(status_code=400, detail="Excel sheet header row is empty")
        if len(header_values) > len(columns):
            warnings.append(f"Only the first {len(columns)} Excel columns were loaded")

        rows: list[dict[str, Any]] = []
        skipped_empty = 0
        reached_limit = False
        for cells in worksheet.iter_rows(min_row=start_row, values_only=False):
            if len(rows) >= max_rows:
                reached_limit = True
                break
            sliced_cells = list(cells[: len(columns)])
            sliced_values = [cell.value for cell in sliced_cells]
            if all(_is_empty_cell(value) for value in sliced_values):
                skipped_empty += 1
                continue
            rows.append(
                {
                    column: _excel_cell_value(sliced_cells[index]) if index < len(sliced_cells) else ""
                    for index, column in enumerate(columns)
                }
            )

        if skipped_empty:
            warnings.append(f"{skipped_empty} empty Excel row(s) were ignored")
        if reached_limit:
            warnings.append(f"Only the first {max_rows} Excel data rows were loaded")
        if not rows:
            raise HTTPException(status_code=400, detail="Excel sheet must include at least one data row")

        return sheets, selected_sheet, header_row, start_row, rows, columns, warnings, reached_limit
    finally:
        workbook.close()


def _parse_structured_rows(
    source_type: str,
    raw_content: str,
    *,
    max_rows: int = 1000,
) -> tuple[list[dict[str, Any]], str, list[str]]:
    content = (raw_content or "").strip()
    warnings: list[str] = []
    if not content:
        raise HTTPException(status_code=400, detail="Structured source content is empty")

    if source_type == "json":
        try:
            parsed = json.loads(content)
        except json.JSONDecodeError as exc:
            raise HTTPException(status_code=400, detail=f"Invalid JSON source: {exc.msg}") from exc
        raw_rows = parsed if isinstance(parsed, list) else [parsed]
        object_rows = [row for row in raw_rows if isinstance(row, dict)]
        if not object_rows:
            raise HTTPException(status_code=400, detail="JSON source must be an object or an array of objects")
        if len(object_rows) < len(raw_rows):
            warnings.append("Non-object JSON rows were ignored")
        rows = [
            {str(key): _row_value(value) for key, value in row.items()}
            for row in object_rows[:max_rows]
        ]
        if len(object_rows) > max_rows:
            warnings.append(f"Only the first {max_rows} JSON rows were profiled")
        return rows, "json", warnings

    sample = content[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel_tab if "\t" in content.splitlines()[0] else csv.excel
    reader = csv.DictReader(io.StringIO(content), dialect=dialect)
    if not reader.fieldnames:
        raise HTTPException(status_code=400, detail="CSV source must include a header row")

    rows: list[dict[str, Any]] = []
    for row in reader:
        if len(rows) >= max_rows:
            warnings.append(f"Only the first {max_rows} CSV rows were profiled")
            break
        if row:
            rows.append({str(key): _row_value(value) for key, value in row.items() if key is not None})
    if not rows:
        raise HTTPException(status_code=400, detail="CSV source must include at least one data row")
    return rows, "table", warnings


def _rows_to_csv(rows: list[dict[str, Any]], columns: Optional[list[str]] = None) -> str:
    ordered_columns = columns or []
    if not ordered_columns:
        for row in rows:
            for column in row.keys():
                if column not in ordered_columns:
                    ordered_columns.append(column)
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=ordered_columns)
    writer.writeheader()
    for row in rows:
        writer.writerow({column: row.get(column, "") for column in ordered_columns})
    return output.getvalue().strip()


def _column_order_from_rows(rows: list[dict[str, Any]]) -> list[str]:
    columns: list[str] = []
    for row in rows:
        for column in row.keys():
            if column not in columns:
                columns.append(column)
    return columns


def _connector_has_inline_sample(config: dict[str, Any]) -> bool:
    raw_content = str(config.get("raw_content") or "").strip()
    if raw_content:
        return True
    for key in ("sample_rows", "rows", "documents", "tables"):
        value = config.get(key)
        if isinstance(value, list) and value:
            return True
        if isinstance(value, dict):
            return True
    return False


def _quote_pg_identifier(identifier: str) -> str:
    if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", identifier):
        raise HTTPException(
            status_code=400,
            detail=f"Invalid DB table identifier '{identifier}'. Use schema.table with letters, numbers, and underscores.",
        )
    return f'"{identifier}"'


def _parse_db_table_ref(source_uri: Any) -> tuple[str, str]:
    table_ref = str(source_uri or "").strip()
    if table_ref.startswith("db://"):
        table_ref = table_ref[len("db://") :]
    table_ref = table_ref.split("?", 1)[0].strip().strip("/")
    table_ref = table_ref.replace("/", ".")
    parts = [part for part in table_ref.split(".") if part]
    for part in parts:
        _quote_pg_identifier(part)
    if len(parts) == 1:
        return "public", parts[0]
    if len(parts) == 2:
        return parts[0], parts[1]
    raise HTTPException(
        status_code=400,
        detail="DB table source must use db://schema.table or schema.table.",
    )


async def _connector_raw_content_from_db(
    db: Any,
    connector: SourceConnector,
    limit: int,
) -> tuple[StructuredSourceType, str, Optional[str], list[str], bool]:
    config = connector.config
    if connector.connector_type != "db_table" or _connector_has_inline_sample(config):
        return _connector_raw_content(connector, limit)

    source_uri = config.get("source_uri") or config.get("table") or config.get("uri")
    if not source_uri:
        return _connector_raw_content(connector, limit)

    schema, table = _parse_db_table_ref(source_uri)
    row_limit = max(1, min(int(limit or 100), 1000))
    qualified_table = f"{_quote_pg_identifier(schema)}.{_quote_pg_identifier(table)}"
    if connector.auth_ref:
        auth_ref = connector.auth_ref.strip()
        if not re.fullmatch(r"[A-Z][A-Z0-9_]{1,127}", auth_ref):
            raise HTTPException(
                status_code=400,
                detail="DB connector auth_ref must be an environment variable name such as FAQ_SOURCE_DATABASE_URL.",
            )
        connection_url = os.getenv(auth_ref)
        if not connection_url:
            raise HTTPException(
                status_code=400,
                detail=f"DB connector environment variable '{auth_ref}' is not configured.",
            )
        try:
            import asyncpg  # type: ignore

            connection = await asyncpg.connect(
                dsn=connection_url,
                timeout=10,
                command_timeout=30,
            )
            try:
                async with connection.transaction(readonly=True):
                    rows = await connection.fetch(
                        f"SELECT * FROM {qualified_table} LIMIT $1",
                        row_limit + 1,
                    )
            finally:
                await connection.close()
        except HTTPException:
            raise
        except Exception as exc:
            logger.warning(
                "[Answers] Failed to read external DB connector %s using %s: %s",
                connector.connector_id,
                auth_ref,
                exc,
            )
            raise HTTPException(
                status_code=502,
                detail=f"Failed to read DB connector '{connector.name}'. Check its connection environment and table.",
            ) from exc
    else:
        rows = await db.query(
            f"SELECT * FROM {qualified_table} LIMIT {row_limit + 1}",
            multirows=True,
        )
    truncated = len(rows or []) > row_limit
    rows = list(rows or [])[:row_limit]
    object_rows = [
        {str(key): _row_value(value) for key, value in row.items()}
        for row in rows
    ]
    if not object_rows:
        raise HTTPException(
            status_code=400,
            detail=f"DB table '{schema}.{table}' returned no rows.",
        )
    warnings = [
        f"Loaded {len(object_rows)} row(s) from "
        f"{'external ' if connector.auth_ref else ''}DB table {schema}.{table}"
    ]
    if truncated:
        warnings.append(
            f"DB table contains more than {row_limit} rows; additional rows were not loaded"
        )
    return (
        "json",
        json.dumps(object_rows, ensure_ascii=False),
        str(source_uri),
        warnings,
        truncated,
    )


def _connector_raw_content(
    connector: SourceConnector,
    limit: int,
) -> tuple[StructuredSourceType, str, Optional[str], list[str], bool]:
    config = connector.config
    warnings: list[str] = []
    source_uri = config.get("source_uri") or config.get("uri") or config.get("url") or config.get("table") or config.get("collection")

    raw_content = str(config.get("raw_content") or "").strip()
    if raw_content:
        source_type = str(config.get("source_type") or "").lower()
        if source_type not in {"csv", "json"}:
            source_type = "json" if raw_content.startswith(("{", "[")) else "csv"
        return source_type, raw_content, source_uri, warnings, False

    rows = config.get("sample_rows") or config.get("rows") or config.get("documents")
    if isinstance(rows, dict):
        rows = [rows]
    if isinstance(rows, list) and rows:
        object_rows = [row for row in rows if isinstance(row, dict)]
        if object_rows:
            sliced_rows = [{str(key): _row_value(value) for key, value in row.items()} for row in object_rows[:limit]]
            if len(object_rows) > limit:
                warnings.append(f"Only the first {limit} connector rows were sampled")
            return "json", json.dumps(sliced_rows, ensure_ascii=False), source_uri, warnings, len(object_rows) > limit
        array_rows = [row for row in rows if isinstance(row, list)]
        columns = config.get("columns")
        if array_rows and isinstance(columns, list) and columns:
            dict_rows = [
                {str(column): _row_value(row[index]) if index < len(row) else "" for index, column in enumerate(columns)}
                for row in array_rows[:limit]
            ]
            return "csv", _rows_to_csv(dict_rows, [str(column) for column in columns]), source_uri, warnings, len(array_rows) > limit

    tables = config.get("tables")
    if isinstance(tables, list) and tables:
        flattened: list[dict[str, Any]] = []
        for table in tables:
            if not isinstance(table, dict):
                continue
            table_name = str(table.get("name") or table.get("table") or "table")
            table_rows = table.get("sample_rows") or table.get("rows") or []
            if isinstance(table_rows, dict):
                table_rows = [table_rows]
            for row in table_rows[: max(1, limit // max(len(tables), 1))]:
                if isinstance(row, dict):
                    flattened.append({"table": table_name, **{str(key): _row_value(value) for key, value in row.items()}})
        if flattened:
            return "json", json.dumps(flattened[:limit], ensure_ascii=False), source_uri, warnings, len(flattened) > limit

    if connector.connector_type == "web":
        row = {
            "title": config.get("title") or connector.name,
            "url": config.get("url") or source_uri or "",
            "content": config.get("content") or config.get("summary") or "",
        }
        return "json", json.dumps([row], ensure_ascii=False), source_uri, warnings, False

    row = {
        "name": connector.name,
        "connector_type": connector.connector_type,
        "source_uri": source_uri or "",
        "description": config.get("description") or connector.metadata.get("description") or "",
    }
    warnings.append("Connector has no raw_content or sample_rows; generated a metadata sample row")
    return "json", json.dumps([row], ensure_ascii=False), source_uri, warnings, False


def _connector_sample_response_from_raw(
    connector: SourceConnector,
    source_type: StructuredSourceType,
    raw_content: str,
    source_uri: Optional[str],
    warnings: list[str],
    truncated: bool,
    limit: int = 1000,
) -> SourceConnectorSampleResponse:
    rows, _, parse_warnings = _parse_structured_rows(source_type, raw_content, max_rows=limit)
    warnings.extend(parse_warnings)
    return SourceConnectorSampleResponse(
        connector_id=connector.connector_id,
        connector_type=connector.connector_type,
        source_type=source_type,
        source_uri=source_uri,
        raw_content=raw_content,
        rows=rows,
        columns=_column_order_from_rows(rows),
        row_count=len(rows),
        row_limit=limit,
        truncated=truncated,
        warnings=warnings,
    )


def _compact_connector_materialization_sample(
    sample: SourceConnectorSampleResponse,
    profile: StructuredProfileResponse,
    *,
    limit: int = 20,
) -> SourceConnectorSampleResponse:
    rows = profile.sample_rows[:limit]
    if sample.source_type == "csv":
        raw_content = _rows_to_csv(rows, profile.columns)
    else:
        raw_content = json.dumps(rows, ensure_ascii=False)
    warnings = list(sample.warnings)
    if sample.row_count > len(rows):
        warnings.append(
            f"Materialization response includes {len(rows)} of {sample.row_count} sample rows. "
            "Use the connector sample or mapping preview API to inspect source rows."
        )
    return sample.model_copy(
        update={
            "raw_content": raw_content,
            "rows": rows,
            "warnings": warnings,
        }
    )


def _connector_sample_response(connector: SourceConnector, limit: int) -> SourceConnectorSampleResponse:
    return _connector_sample_response_from_raw(
        connector,
        *_connector_raw_content(connector, limit),
        limit=limit,
    )


def _mapping_from_profile(
    profile: StructuredProfileResponse,
    override: Optional[dict[str, str]] = None,
) -> dict[str, str]:
    mapping: dict[str, str] = {}
    for role, column in profile.mapping_suggestions.items():
        if column and column in profile.columns:
            mapping[role] = column
    for role, column in (override or {}).items():
        if column in profile.columns:
            mapping = {mapped_role: mapped_column for mapped_role, mapped_column in mapping.items() if mapped_column != column}
            mapping[str(role)] = column
    return mapping


def _guidance_columns_from_profile(
    profile: StructuredProfileResponse,
    mapping: dict[str, str],
    override: Optional[list[str]] = None,
) -> list[str]:
    guidance_columns = [
        column for column in (override or [])
        if column in profile.columns and column not in mapping.values()
    ]
    if guidance_columns:
        return guidance_columns
    suggested: list[str] = []
    guidance_name_hints = {
        "keyword", "keywords", "tag", "tags", "intent", "synonym",
        "키워드", "태그", "의도", "동의어", "검색어",
    }
    mapped_columns = set(mapping.values())
    for field in profile.fields:
        normalized = re.sub(r"[^0-9a-z가-힣]+", "_", field.name.strip().lower()).strip("_")
        name_parts = set(normalized.split("_"))
        should_include = (
            field.semantic_role in {"category", "title", "id"}
            or normalized in guidance_name_hints
            or bool(name_parts & guidance_name_hints)
        )
        if should_include and field.name not in mapped_columns and field.name not in suggested:
            suggested.append(field.name)
    return suggested[:5]


def _is_empty_cell(value: Any) -> bool:
    if value is None:
        return True
    normalized = str(value).strip().lower()
    return normalized in {"", "null", "none", "n/a", "na"}


def _infer_cell_type(value: Any) -> str:
    if value is None:
        return "empty"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, int):
        return "integer"
    if isinstance(value, float):
        return "number"
    text = str(value).strip()
    if not text:
        return "empty"
    if text.lower() in {"true", "false", "yes", "no", "y", "n"}:
        return "boolean"
    if re.fullmatch(r"[-+]?\d+", text):
        return "integer"
    if re.fullmatch(r"[-+]?\d+\.\d+", text):
        return "number"
    if re.fullmatch(r"\d{4}[-/.]\d{1,2}[-/.]\d{1,2}.*", text):
        return "date"
    if (text.startswith("{") and text.endswith("}")) or (text.startswith("[") and text.endswith("]")):
        try:
            json.loads(text)
            return "json"
        except json.JSONDecodeError:
            return "text"
    return "text"


def _infer_column_type(values: list[Any]) -> str:
    typed = [_infer_cell_type(value) for value in values if not _is_empty_cell(value)]
    if not typed:
        return "empty"
    counts = {kind: typed.count(kind) for kind in set(typed)}
    preferred = ["integer", "number", "boolean", "date", "json", "text"]
    dominant = max(preferred, key=lambda item: (counts.get(item, 0), -preferred.index(item)))
    if dominant == "integer" and counts.get("number"):
        return "number"
    return dominant if counts.get(dominant, 0) / max(len(typed), 1) >= 0.7 else "text"


def _suggest_semantic_role(column: str, inferred_type: str) -> tuple[str, float]:
    normalized = re.sub(r"[^0-9a-z가-힣]+", "_", column.strip().lower()).strip("_")
    role_patterns: list[tuple[str, set[str], float]] = [
        ("question", {"question", "query", "intent", "utterance", "질문", "문의", "의도"}, 0.92),
        ("answer", {"answer", "response", "reply", "body", "content", "답변", "응답", "내용"}, 0.92),
        ("title", {"title", "name", "subject", "제목", "이름", "명칭"}, 0.86),
        ("category", {"category", "type", "group", "분류", "유형", "카테고리"}, 0.84),
        ("status", {"status", "state", "상태"}, 0.82),
        ("valid_from", {"valid_from", "start", "from", "시작", "시작일"}, 0.8),
        ("valid_until", {"valid_until", "valid_to", "end", "until", "종료", "종료일", "만료"}, 0.8),
        ("id", {"id", "key", "code", "uid", "번호", "코드"}, 0.78),
    ]
    for role, names, confidence in role_patterns:
        if normalized in names or any(part in names for part in normalized.split("_")):
            return role, confidence
    if inferred_type == "date":
        return "metadata", 0.45
    return "metadata", 0.35


def _profile_structured_source(
    source_type: StructuredSourceType,
    raw_content: str,
    *,
    sample_limit: int = 20,
) -> StructuredProfileResponse:
    rows, kind, warnings = _parse_structured_rows(source_type, raw_content)
    columns: list[str] = []
    for row in rows:
        for column in row.keys():
            if column not in columns:
                columns.append(column)

    fields: list[StructuredFieldProfile] = []
    mapping_suggestions: dict[str, Optional[str]] = {
        "id": None,
        "title": None,
        "question": None,
        "answer": None,
        "category": None,
        "status": None,
        "valid_from": None,
        "valid_until": None,
    }
    for column in columns:
        values = [row.get(column) for row in rows]
        non_empty_values = [value for value in values if not _is_empty_cell(value)]
        inferred_type = _infer_column_type(values)
        role, confidence = _suggest_semantic_role(column, inferred_type)
        sample_values = []
        seen_samples: set[str] = set()
        for value in non_empty_values:
            text = str(value)
            if text in seen_samples:
                continue
            sample_values.append(text[:160])
            seen_samples.add(text)
            if len(sample_values) >= 5:
                break
        null_count = len(values) - len(non_empty_values)
        fields.append(
            StructuredFieldProfile(
                name=column,
                inferred_type=inferred_type,
                null_count=null_count,
                null_rate=round(null_count / max(len(values), 1), 4),
                distinct_count=len({str(value) for value in non_empty_values}),
                sample_values=sample_values,
                semantic_role=role,
                confidence=confidence,
            )
        )
        if role in mapping_suggestions and mapping_suggestions[role] is None and confidence >= 0.7:
            mapping_suggestions[role] = column

    return StructuredProfileResponse(
        source_type=source_type,
        kind=kind,
        row_count=len(rows),
        columns=columns,
        fields=fields,
        sample_rows=rows[:sample_limit],
        mapping_suggestions=mapping_suggestions,
        warnings=warnings,
    )


def _guidance_from_structured_profile(
    rows: list[dict[str, Any]],
    mapping: dict[str, str],
    guidance_columns: list[str],
    conversion_purpose: StructuredConversionPurpose = "faq",
) -> list[SourceGuidanceCandidate]:
    candidates: list[SourceGuidanceCandidate] = []
    seen: set[tuple[str, str]] = set()

    def add(guidance_type: GuidanceType, value: Any, weight: float, source: str) -> None:
        text = str(value or "").strip()
        if len(text) < 2:
            return
        key = (guidance_type, text.lower())
        if key in seen:
            return
        seen.add(key)
        candidates.append(
            SourceGuidanceCandidate(
                guidance_type=guidance_type,
                text=text[:500],
                weight=weight,
                source=source,
                metadata={"created_from": "structured_materialize"},
            )
        )

    question_column = mapping.get("question")
    title_column = mapping.get("title")
    category_column = mapping.get("category")
    answer_column = mapping.get("answer")
    lookup_columns = _id_lookup_searchable_columns(mapping, guidance_columns)
    for row in rows[:100]:
        if conversion_purpose == "id_lookup":
            combined_details = " ".join(
                dict.fromkeys(
                    str(row.get(column) or "").strip()
                    for column in lookup_columns
                    if str(row.get(column) or "").strip()
                )
            )
            add("question", combined_details, 1.35, "id_lookup_details")
        if question_column:
            add("question", row.get(question_column), 1.25, question_column)
        if title_column:
            add("keyword", row.get(title_column), 0.95, title_column)
        if category_column:
            add("keyword", row.get(category_column), 0.85, category_column)
        if answer_column:
            add("note", row.get(answer_column), 0.45, answer_column)
        for column in guidance_columns:
            column_name = _normalise_text(column)
            if any(
                marker in column_name
                for marker in ("synonym", "alias", "동의어", "유사어", "별칭")
            ):
                guidance_type: GuidanceType = "synonym"
                weight = 1.2
            elif any(
                marker in column_name
                for marker in (
                    "negative",
                    "exclude",
                    "block",
                    "제외어",
                    "금지어",
                    "차단어",
                )
            ):
                guidance_type = "negative_keyword"
                weight = 1.0
            else:
                guidance_type = "keyword"
                weight = 1.0

            raw_guidance = str(row.get(column) or "").strip()
            split_values = [
                value.strip()
                for value in re.split(r"[,;|\n]+", raw_guidance)
                if value.strip()
            ]
            for value in split_values or [raw_guidance]:
                add(guidance_type, value, weight, column)
        if len(candidates) >= 80:
            break
    return candidates[:80]


def _id_lookup_searchable_columns(
    mapping: dict[str, str],
    guidance_columns: list[str],
) -> list[str]:
    id_column = mapping.get("id")
    columns: list[str] = []
    for role in ("question", "title", "category", "answer"):
        column = mapping.get(role)
        if column and column != id_column and column not in columns:
            columns.append(column)
    for column in guidance_columns:
        if column and column != id_column and column not in columns:
            columns.append(column)
    return columns


def _validate_id_lookup_rows(
    rows: list[dict[str, Any]],
    mapping: dict[str, str],
    guidance_columns: list[str],
    conversion_purpose: StructuredConversionPurpose,
) -> StructuredIdLookupValidation:
    if conversion_purpose != "id_lookup":
        return StructuredIdLookupValidation()

    id_column = mapping.get("id")
    searchable_columns = _id_lookup_searchable_columns(mapping, guidance_columns)
    blank_id_rows: list[int] = []
    id_rows: dict[str, list[int]] = {}
    detail_groups: dict[tuple[str, ...], dict[str, Any]] = {}

    if id_column:
        for row_index, row in enumerate(rows, start=1):
            source_id = str(row.get(id_column) or "").strip()
            if not source_id:
                blank_id_rows.append(row_index)
                continue
            id_rows.setdefault(source_id, []).append(row_index)
            fingerprint = tuple(
                _normalise_text(str(row.get(column) or ""))
                for column in searchable_columns
            )
            if any(fingerprint):
                group = detail_groups.setdefault(
                    fingerprint,
                    {
                        "values": {
                            column: str(row.get(column) or "").strip()
                            for column in searchable_columns
                        },
                        "ids": set(),
                        "rows": [],
                    },
                )
                group["ids"].add(source_id)
                group["rows"].append(row_index)

    duplicate_ids = sorted(
        source_id
        for source_id, source_rows in id_rows.items()
        if len(source_rows) > 1
    )
    ambiguous_detail_groups = [
        {
            "values": group["values"],
            "ids": sorted(group["ids"]),
            "rows": group["rows"][:20],
        }
        for group in detail_groups.values()
        if len(group["ids"]) > 1
    ][:20]

    warnings: list[str] = []
    if not id_column:
        warnings.append("Select the business ID column to return.")
    if not searchable_columns:
        warnings.append("Select at least one detail column to search.")
    if blank_id_rows:
        warnings.append(f"{len(blank_id_rows)} rows have an empty business ID.")
    if duplicate_ids:
        warnings.append(
            f"{len(duplicate_ids)} business IDs appear in more than one row. "
            "They can be created, but duplicate FAQ candidates may be returned."
        )
    if ambiguous_detail_groups:
        warnings.append(
            f"{len(ambiguous_detail_groups)} detail combinations point to different business IDs."
        )

    return StructuredIdLookupValidation(
        enabled=True,
        ready=bool(id_column and searchable_columns)
        and not blank_id_rows
        and not ambiguous_detail_groups,
        id_column=id_column,
        searchable_columns=searchable_columns,
        row_count=len(rows),
        valid_id_count=sum(len(source_rows) for source_rows in id_rows.values()),
        blank_id_rows=blank_id_rows[:100],
        duplicate_ids=duplicate_ids[:100],
        ambiguous_detail_groups=ambiguous_detail_groups,
        warnings=warnings,
    )


def _structured_row_text(row: dict[str, Any], column: Optional[str]) -> str:
    if not column:
        return ""
    return str(row.get(column) or "").strip()


def _title_for_structured_row(
    base_title: str,
    row: dict[str, Any],
    mapping: dict[str, str],
    row_index: int,
) -> str:
    for role in ("title", "question", "id"):
        text = _structured_row_text(row, mapping.get(role))
        if text:
            return text[:500]
    return f"{base_title} #{row_index + 1}"


def _body_for_structured_row(
    row: dict[str, Any],
    mapping: dict[str, str],
    conversion_purpose: StructuredConversionPurpose = "faq",
) -> str:
    if conversion_purpose == "id_lookup":
        return _structured_row_text(row, mapping.get("id"))
    body = _structured_row_text(row, mapping.get("answer"))
    if body:
        return body
    return json.dumps(row, ensure_ascii=False, indent=2)


def _summary_for_structured_row(
    row: dict[str, Any],
    mapping: dict[str, str],
    fallback: Optional[str],
    conversion_purpose: StructuredConversionPurpose = "faq",
) -> Optional[str]:
    if conversion_purpose == "id_lookup":
        return None
    summary = _structured_row_text(row, mapping.get("answer"))
    return summary or fallback


def _tags_for_structured_row(
    base_tags: list[str],
    row: dict[str, Any],
    mapping: dict[str, str],
    source_type: str,
) -> list[str]:
    tags: list[str] = []
    for tag in [*base_tags, "structured", source_type]:
        text = str(tag or "").strip()
        if text and text not in tags:
            tags.append(text)
    category = _structured_row_text(row, mapping.get("category"))
    if category and category not in tags:
        tags.append(category)
    return tags


def _status_for_structured_row(
    row: dict[str, Any],
    mapping: dict[str, str],
    fallback: AnswerStatus,
) -> AnswerStatus:
    text = _structured_row_text(row, mapping.get("status")).strip().lower()
    if not text:
        return fallback
    status_aliases: dict[str, AnswerStatus] = {
        "draft": "draft",
        "초안": "draft",
        "검토": "draft",
        "검토중": "draft",
        "검토 중": "draft",
        "published": "published",
        "publish": "published",
        "active": "published",
        "게시": "published",
        "게시됨": "published",
        "사용": "published",
        "운영": "published",
        "archived": "archived",
        "archive": "archived",
        "보관": "archived",
        "보관됨": "archived",
        "숨김": "archived",
        "expired": "expired",
        "expire": "expired",
        "만료": "expired",
        "종료": "expired",
    }
    return status_aliases.get(text, fallback)


def _datetime_for_structured_row(
    row: dict[str, Any],
    mapping: dict[str, str],
    role: Literal["valid_from", "valid_until"],
) -> Optional[datetime]:
    text = _structured_row_text(row, mapping.get(role))
    if not text:
        return None
    normalized = text.strip().replace("Z", "+00:00")
    if re.fullmatch(r"\d{4}[./]\d{1,2}[./]\d{1,2}", normalized):
        normalized = normalized.replace(".", "-").replace("/", "-")
    try:
        parsed = datetime.fromisoformat(normalized)
    except ValueError:
        return None
    if parsed.tzinfo is None:
        return parsed.replace(tzinfo=timezone.utc)
    return parsed


def _rows_from_answer(answer: AnswerItem) -> tuple[list[dict[str, Any]], list[str], str]:
    body = answer.body.strip()
    rows: list[dict[str, Any]] = []
    kind = "text"

    if body:
        try:
            rows, kind, _ = _parse_structured_rows("json", body)
        except HTTPException:
            rows = []

    if not rows and body:
        try:
            rows, kind, _ = _parse_structured_rows("csv", body)
        except HTTPException:
            rows = []

    profile = _structured_profile(answer)
    if not rows and isinstance(answer.metadata, dict) and isinstance(answer.metadata.get("source_row"), dict):
        rows = [{str(key): _row_value(value) for key, value in answer.metadata["source_row"].items()}]
        kind = "structured_row"

    if not rows:
        rows = [
            {
                "answer_id": answer.answer_id,
                "title": answer.title,
                "approved_summary": answer.approved_summary or "",
                "body": answer.body,
                "status": answer.status,
                "priority": answer.priority,
                "tags": ", ".join(answer.tags),
                "source_type": answer.metadata.get("source_type") or answer.metadata.get("created_from") or "",
                "source_uri": answer.metadata.get("source_uri") or "",
            }
        ]
        kind = profile["kind"] or "answer"

    columns: list[str] = []
    for row in rows:
        for column in row.keys():
            if column not in columns:
                columns.append(column)
    for column in profile["columns"]:
        if column not in columns:
            columns.append(column)
            for row in rows:
                row.setdefault(column, "")

    return rows, columns, kind


def _dataset_from_answer(answer: AnswerItem) -> StructuredDataset:
    rows, columns, parsed_kind = _rows_from_answer(answer)
    profile = _structured_profile(answer)
    row_count = max(profile["row_count"], len(rows))
    return StructuredDataset(
        answer_id=answer.answer_id,
        title=answer.title,
        status=answer.status,
        source_type=str(answer.metadata.get("source_type") or answer.metadata.get("created_from") or "-"),
        source_uri=answer.metadata.get("source_uri"),
        kind=str(profile["kind"] or parsed_kind),
        columns=columns,
        row_count=row_count,
        sample_rows=rows[:5],
        metadata={
            "content_format": answer.content_format,
            "display_policy": answer.display_policy,
            "version": answer.version,
            "materialization_mode": answer.metadata.get("materialization_mode"),
        },
    )


def _matches_filter(row: dict[str, Any], item: StructuredFilter) -> bool:
    value = str(row.get(item.field, "") or "").lower()
    expected = str(item.value or "").lower()
    if item.operator == "equals":
        return value == expected
    if item.operator == "starts_with":
        return value.startswith(expected)
    if item.operator == "ends_with":
        return value.endswith(expected)
    return expected in value


def _pseudo_sql(dataset: StructuredDataset, filters: list[StructuredFilter], limit: int) -> str:
    columns = ", ".join(f'"{column}"' for column in dataset.columns[:20]) or "*"
    where_parts = []
    for item in filters:
        escaped = item.value.replace("'", "''")
        if item.operator == "equals":
            where_parts.append(f'"{item.field}" = \'{escaped}\'')
        elif item.operator == "starts_with":
            where_parts.append(f'"{item.field}" LIKE \'{escaped}%\'')
        elif item.operator == "ends_with":
            where_parts.append(f'"{item.field}" LIKE \'%{escaped}\'')
        else:
            where_parts.append(f'"{item.field}" ILIKE \'%{escaped}%\'')
    where_sql = f" WHERE {' AND '.join(where_parts)}" if where_parts else ""
    return f'SELECT {columns} FROM "{dataset.answer_id}"{where_sql} LIMIT {limit};'


def _normalise_answer_snapshot(snapshot: dict[str, Any]) -> dict[str, Any]:
    tags = _coerce_json(snapshot.get("tags"), [])
    metadata = _coerce_json(snapshot.get("metadata"), {})
    snapshot["tags"] = tags if isinstance(tags, list) else []
    snapshot["metadata"] = metadata if isinstance(metadata, dict) else {}
    if "assets" in snapshot:
        assets = _coerce_json(snapshot.get("assets"), [])
        snapshot["assets"] = assets if isinstance(assets, list) else []
    return snapshot


def _normalise_text(value: Any) -> str:
    normalized = unicodedata.normalize("NFKC", str(value or ""))
    return " ".join(normalized.lower().split())


def _tokens(value: Any) -> list[str]:
    return [token for token in TOKEN_RE.findall(_normalise_text(value)) if len(token) >= 2]


def _candidate_query_terms(value: Any) -> list[str]:
    terms: list[str] = []
    for token in _tokens(value):
        variants = [token]
        if len(token) >= 4 and re.fullmatch(r"[가-힣]+", token):
            variants.extend([token[:-1], token[:-2]])
        for variant in variants:
            if len(variant) >= 2 and variant not in terms:
                terms.append(variant)
    return terms[:64]


def _builtin_alias_groups() -> list[AnswerAliasGroup]:
    return [
        AnswerAliasGroup(
            alias_id=str(item["alias_id"]),
            workspace="*",
            canonical_term=str(item["canonical_term"]),
            aliases=[str(alias) for alias in item["aliases"]],
            enabled=True,
            source="builtin",
            metadata={"read_only": True},
        )
        for item in BUILTIN_ANSWER_ALIAS_GROUPS
    ]


def _alias_member_matches_query(normalized_query: str, member: str) -> bool:
    normalized_member = _normalise_text(member)
    if len(normalized_member) < 2:
        return False
    if re.search(r"[가-힣]", normalized_member):
        return normalized_member in normalized_query
    pattern = rf"(?<![0-9a-z_]){re.escape(normalized_member)}(?![0-9a-z_])"
    return re.search(pattern, normalized_query) is not None


async def _answer_alias_groups(db, workspace: str) -> list[AnswerAliasGroup]:
    rows = await db.query(
        """
        SELECT alias_id, workspace, canonical_term, aliases, enabled, source,
               metadata, create_time, update_time
        FROM LIGHTRAG_ANSWER_TERM_ALIASES
        WHERE workspace = $1 AND enabled = TRUE
        ORDER BY LOWER(canonical_term), alias_id
        """,
        [workspace],
        multirows=True,
    )
    return [*_builtin_alias_groups(), *[_alias_group_from_row(dict(row)) for row in rows or []]]


async def _expand_query_aliases(
    db,
    workspace: str,
    query: str,
) -> list[AnswerAliasExpansion]:
    normalized_query = _normalise_text(query)
    expansions: list[AnswerAliasExpansion] = []
    seen_canonical: set[str] = set()
    for group in await _answer_alias_groups(db, workspace):
        members = [group.canonical_term, *group.aliases]
        matched_term = next(
            (
                member
                for member in members
                if _alias_member_matches_query(normalized_query, member)
            ),
            None,
        )
        canonical_key = _normalise_text(group.canonical_term)
        if not matched_term or canonical_key in seen_canonical:
            continue
        seen_canonical.add(canonical_key)
        expanded_terms: list[str] = []
        for member in members:
            value = " ".join(str(member or "").split())
            if value and _normalise_text(value) not in {
                _normalise_text(item) for item in expanded_terms
            }:
                expanded_terms.append(value)
        expansions.append(
            AnswerAliasExpansion(
                canonical_term=group.canonical_term,
                matched_term=matched_term,
                expanded_terms=expanded_terms,
                source=group.source,
            )
        )
    return expansions


def _alias_query_variants(
    query: str,
    alias_expansions: Optional[list[AnswerAliasExpansion]] = None,
) -> list[str]:
    variants = [_normalise_text(query)]
    for expansion in alias_expansions or []:
        matched_term = _normalise_text(expansion.matched_term)
        if not matched_term:
            continue
        current_variants = list(variants)
        for current in current_variants:
            if not _alias_member_matches_query(current, matched_term):
                continue
            for expanded_term in expansion.expanded_terms:
                replacement = _normalise_text(expanded_term)
                if not replacement:
                    continue
                if re.search(r"[가-힣]", matched_term):
                    candidate = current.replace(matched_term, replacement)
                else:
                    pattern = rf"(?<![0-9a-z_]){re.escape(matched_term)}(?![0-9a-z_])"
                    candidate = re.sub(pattern, replacement, current)
                candidate = " ".join(candidate.split())
                if candidate and candidate not in variants:
                    variants.append(candidate)
                if len(variants) >= 32:
                    return variants
    return variants


def _partial_overlap(query_tokens: list[str], text: str) -> float:
    if not query_tokens:
        return 0.0
    normalized_text = _normalise_text(text)
    text_tokens = _tokens(text)
    matches = 0
    for query_token in query_tokens:
        if query_token in normalized_text:
            matches += 1
            continue
        if any(query_token in text_token or text_token in query_token for text_token in text_tokens):
            matches += 1
    return matches / max(len(query_tokens), 1)


def _add_score(details: dict[str, float], key: str, amount: float) -> None:
    if amount == 0:
        return
    details[key] = round(details.get(key, 0.0) + amount, 4)


def _safe_float(value: Any, default: float, minimum: Optional[float] = None, maximum: Optional[float] = None) -> float:
    try:
        number = float(value)
    except (TypeError, ValueError):
        number = default
    if minimum is not None:
        number = max(minimum, number)
    if maximum is not None:
        number = min(maximum, number)
    return number


def _compact_answer_metadata(metadata: Any) -> dict[str, Any]:
    if not isinstance(metadata, dict):
        return {}
    return {
        key: value
        for key, value in metadata.items()
        if key not in {"source_profile", "raw_content"}
    }


def _search_answer_view(answer: AnswerItem) -> AnswerItem:
    compact_metadata = _compact_answer_metadata(answer.metadata)
    if hasattr(answer, "model_copy"):
        return answer.model_copy(update={"metadata": compact_metadata})
    return answer.copy(update={"metadata": compact_metadata})


def _matched_business_id(answer: Optional[AnswerItem]) -> Optional[str]:
    if answer is None or not isinstance(answer.metadata, dict):
        return None
    if answer.metadata.get("conversion_purpose") != "id_lookup":
        return None
    matched_id = answer.metadata.get("matched_id") or answer.metadata.get("source_id")
    text = str(matched_id or "").strip()
    return text or None


def _embedding_func_from_rag(rag):
    if hasattr(rag, "embedding_func") and rag.embedding_func:
        return rag.embedding_func
    if hasattr(rag, "text_chunks") and hasattr(rag.text_chunks, "embedding_func"):
        return rag.text_chunks.embedding_func
    return None


def _embedding_to_list(value: Any) -> list[float]:
    if hasattr(value, "tolist"):
        value = value.tolist()
    if isinstance(value, tuple):
        value = list(value)
    if isinstance(value, list) and value and isinstance(value[0], (list, tuple)):
        value = value[0]
    if not isinstance(value, list):
        return []
    converted: list[float] = []
    for item in value:
        try:
            converted.append(float(item))
        except (TypeError, ValueError):
            return []
    return converted


def _embedding_to_pgvector(value: list[float]) -> str:
    return json.dumps(value, ensure_ascii=True, separators=(",", ":"))


def _graph_config_from_row(workspace: str, row: Optional[dict[str, Any]]) -> AnswerGraphConfig:
    if not row:
        return AnswerGraphConfig(workspace=workspace)
    entity_types = _coerce_json(row.get("entity_types"), [])
    relation_types = _coerce_json(row.get("relation_types"), [])
    return AnswerGraphConfig(
        workspace=workspace,
        enabled=bool(row.get("enabled")),
        auto_sync=bool(row.get("auto_sync")),
        graph_weight=_safe_float(row.get("graph_weight"), 0.35, 0.0, 1.0),
        min_similarity=_safe_float(row.get("min_similarity"), 0.60, 0.0, 1.0),
        max_hops=max(1, min(3, int(row.get("max_hops") or 2))),
        precision_mode=bool(row.get("precision_mode")),
        precision_min_score=_safe_float(
            row.get("precision_min_score"), 0.35, 0.0, 1.0
        ),
        min_score_margin=_safe_float(
            row.get("min_score_margin"), 0.08, 0.0, 1.0
        ),
        min_category_margin=_safe_float(
            row.get("min_category_margin"), 0.05, 0.0, 1.0
        ),
        min_evidence_sources=max(
            1, min(3, int(row.get("min_evidence_sources") or 2))
        ),
        llm_min_confidence=_safe_float(
            row.get("llm_min_confidence"), 0.80, 0.0, 1.0
        ),
        entity_types=(
            [str(item) for item in entity_types if str(item).strip()]
            if isinstance(entity_types, list) and entity_types
            else list(DEFAULT_ANSWER_GRAPH_ENTITY_TYPES)
        ),
        relation_types=(
            [str(item) for item in relation_types if str(item).strip()]
            if isinstance(relation_types, list) and relation_types
            else list(DEFAULT_ANSWER_GRAPH_RELATION_TYPES)
        ),
        extraction_prompt=str(row.get("extraction_prompt") or DEFAULT_ANSWER_GRAPH_PROMPT),
        ai_extraction_strategy=(
            str(row.get("ai_extraction_strategy"))
            if str(row.get("ai_extraction_strategy")) in {"fast", "adaptive", "deep"}
            else "fast"
        ),
        ai_retry_max_tokens=max(
            1024, min(32768, int(row.get("ai_retry_max_tokens") or 8192))
        ),
        ai_min_relations=max(
            1, min(12, int(row.get("ai_min_relations") or 4))
        ),
        ai_min_relation_types=max(
            1, min(8, int(row.get("ai_min_relation_types") or 3))
        ),
        schema_version=int(row.get("schema_version") or 1),
        create_time=_iso(row.get("create_time")),
        update_time=_iso(row.get("update_time")),
    )


async def _get_answer_graph_config(db, workspace: str) -> AnswerGraphConfig:
    row = await db.query(
        """
        SELECT workspace, enabled, auto_sync, graph_weight, min_similarity, max_hops,
               precision_mode, precision_min_score, min_score_margin,
               min_category_margin, min_evidence_sources, llm_min_confidence,
               entity_types, relation_types, extraction_prompt,
               ai_extraction_strategy, ai_retry_max_tokens, ai_min_relations,
               ai_min_relation_types, schema_version,
               create_time, update_time
        FROM LIGHTRAG_ANSWER_GRAPH_CONFIG
        WHERE workspace = $1
        """,
        [workspace],
    )
    return _graph_config_from_row(workspace, dict(row) if row else None)


def _graph_projection_from_row(row: dict[str, Any]) -> AnswerGraphProjection:
    raw_nodes = _coerce_json(row.get("nodes"), [])
    raw_relations = _coerce_json(row.get("relations"), [])
    return AnswerGraphProjection(
        workspace=str(row["workspace"]),
        answer_id=str(row["answer_id"]),
        answer_version=int(row.get("answer_version") or 1),
        schema_version=int(row.get("schema_version") or 1),
        status=str(row.get("status") or "pending"),
        content_hash=str(row.get("content_hash") or ""),
        nodes=[
            AnswerGraphNode.model_validate(item)
            for item in raw_nodes
            if isinstance(item, dict)
        ],
        relations=[
            AnswerGraphRelation.model_validate(item)
            for item in raw_relations
            if isinstance(item, dict)
        ],
        error=row.get("error"),
        built_at=_iso(row.get("built_at")),
        update_time=_iso(row.get("update_time")),
    )


def _graph_clean_label(value: Any, max_length: int = 160) -> str:
    label = " ".join(str(value or "").replace("::", " ").split()).strip()
    return label[:max_length]


def _graph_node_id(entity_type: str, label: str) -> str:
    safe_type = re.sub(r"[^0-9A-Za-z_]", "", entity_type) or "Term"
    safe_label = _graph_clean_label(label)
    return f"{ANSWER_GRAPH_NODE_PREFIX}::{safe_type.upper()}::{safe_label}"


def _answer_graph_node_id(answer_id: str) -> str:
    return _graph_node_id("FAQAnswer", answer_id)


def _answer_id_from_graph_node(node_id: str) -> Optional[str]:
    prefix = f"{ANSWER_GRAPH_NODE_PREFIX}::FAQANSWER::"
    return node_id[len(prefix) :] if node_id.startswith(prefix) else None


def _graph_label_from_node_id(node_id: str) -> str:
    return node_id.split("::", 2)[-1]


def _graph_values(value: Any) -> list[str]:
    if isinstance(value, list):
        values = value
    elif isinstance(value, (str, int, float)):
        values = re.split(r"[,|;/\n]+", str(value))
    else:
        values = []
    unique: list[str] = []
    for item in values:
        label = _graph_clean_label(item)
        if len(label) >= 2 and _normalise_text(label) not in {
            _normalise_text(existing) for existing in unique
        }:
            unique.append(label)
    return unique[:20]


def _graph_content_hash(
    answer: AnswerItem,
    guidance: list[AnswerGuidance],
    config: AnswerGraphConfig,
) -> str:
    source = {
        "answer_id": answer.answer_id,
        "version": answer.version,
        "title": answer.title,
        "body": answer.body,
        "tags": answer.tags,
        "metadata": answer.metadata,
        "assets": [
            {
                "asset_id": asset.asset_id,
                "type": asset.asset_type,
                "caption": asset.caption,
                "alt_text": asset.alt_text,
                "search_text": asset.search_text,
                "content_text": asset.content_text,
            }
            for asset in answer.assets
            if asset.is_active
        ],
        "guidance": [
            {
                "type": item.guidance_type,
                "text": item.text,
                "weight": item.weight,
            }
            for item in guidance
        ],
        "schema_version": config.schema_version,
        "entity_types": config.entity_types,
        "relation_types": config.relation_types,
    }
    return hashlib.sha256(
        json.dumps(source, ensure_ascii=False, sort_keys=True).encode("utf-8")
    ).hexdigest()


def _deterministic_answer_graph_projection(
    answer: AnswerItem,
    guidance: list[AnswerGuidance],
    alias_groups: list[AnswerAliasGroup],
    config: AnswerGraphConfig,
) -> tuple[list[AnswerGraphNode], list[AnswerGraphRelation]]:
    allowed_entities = set(config.entity_types)
    allowed_relations = set(config.relation_types)
    answer_node_id = _answer_graph_node_id(answer.answer_id)
    nodes: dict[str, AnswerGraphNode] = {
        answer_node_id: AnswerGraphNode(
            node_id=answer_node_id,
            entity_type="FAQAnswer",
            label=answer.title,
            description="\n".join(
                part
                for part in [
                    answer.title,
                    answer.body[:1800],
                    f"Tags: {', '.join(answer.tags)}" if answer.tags else "",
                    (
                        f"Attached media: {_answer_asset_search_content(answer.assets)[:1200]}"
                        if _answer_asset_search_content(answer.assets)
                        else ""
                    ),
                ]
                if part
            ),
        )
    }
    relations: dict[tuple[str, str], AnswerGraphRelation] = {}

    def add_concept(
        entity_type: str,
        label: str,
        relation_type: str,
        description: str,
        *,
        source: str = "deterministic",
        weight: float = 1.0,
    ) -> Optional[str]:
        clean_label = _graph_clean_label(label)
        if (
            entity_type not in allowed_entities
            or relation_type not in allowed_relations
            or len(clean_label) < 2
        ):
            return None
        node_id = _graph_node_id(entity_type, clean_label)
        nodes.setdefault(
            node_id,
            AnswerGraphNode(
                node_id=node_id,
                entity_type=entity_type,
                label=clean_label,
                description=description,
                source=source,
            ),
        )
        key = tuple(sorted((answer_node_id, node_id)))
        relations[key] = AnswerGraphRelation(
            source_id=answer_node_id,
            target_id=node_id,
            relation_type=relation_type,
            description=description,
            weight=weight,
            source=source,
        )
        return node_id

    guidance_relation_map = {
        "question": ("Intent", "REPRESENTS_QUESTION"),
        "keyword": ("Term", "HAS_TERM"),
        "synonym": ("Term", "HAS_TERM"),
        "note": ("Procedure", "RESOLVES"),
    }
    for item in guidance:
        if item.guidance_type == "negative_keyword":
            continue
        entity_type, relation_type = guidance_relation_map.get(
            item.guidance_type, ("Term", "HAS_TERM")
        )
        add_concept(
            entity_type,
            item.text,
            relation_type,
            f"{item.guidance_type} finding hint for {answer.title}",
            weight=max(0.1, min(float(item.weight), 3.0)),
        )

    for tag in answer.tags:
        add_concept(
            "Category",
            tag,
            "BELONGS_TO",
            f"Category tag for {answer.title}",
        )

    metadata_fields = {
        "category": ("Category", "BELONGS_TO"),
        "service_area": ("System", "APPLIES_TO"),
        "system": ("System", "APPLIES_TO"),
        "product": ("Product", "APPLIES_TO"),
        "application": ("Product", "APPLIES_TO"),
        "symptom": ("Symptom", "HAS_SYMPTOM"),
        "symptoms": ("Symptom", "HAS_SYMPTOM"),
        "error_code": ("ErrorCode", "HAS_TERM"),
        "code": ("ErrorCode", "HAS_TERM"),
        "procedure": ("Procedure", "RESOLVES"),
    }
    for key, (entity_type, relation_type) in metadata_fields.items():
        for value in _graph_values(answer.metadata.get(key)):
            add_concept(
                entity_type,
                value,
                relation_type,
                f"{key} metadata for {answer.title}",
            )

    searchable_text = _normalise_text(
        " ".join(
            [
                answer.title,
                answer.body,
                *answer.tags,
                *(item.text for item in guidance if item.guidance_type != "negative_keyword"),
            ]
        )
    )
    for group in alias_groups:
        members = [group.canonical_term, *group.aliases]
        if not any(
            _alias_member_matches_query(searchable_text, member) for member in members
        ):
            continue
        canonical_id = add_concept(
            "Term",
            group.canonical_term,
            "HAS_TERM",
            f"Standard term for {answer.title}",
            source=group.source,
        )
        if not canonical_id or "ALIAS_OF" not in allowed_relations:
            continue
        for alias in group.aliases:
            alias_label = _graph_clean_label(alias)
            if len(alias_label) < 2:
                continue
            alias_id = _graph_node_id("Term", alias_label)
            nodes.setdefault(
                alias_id,
                AnswerGraphNode(
                    node_id=alias_id,
                    entity_type="Term",
                    label=alias_label,
                    description=f"Alias of {group.canonical_term}",
                    source=group.source,
                ),
            )
            key = tuple(sorted((alias_id, canonical_id)))
            relations[key] = AnswerGraphRelation(
                source_id=alias_id,
                target_id=canonical_id,
                relation_type="ALIAS_OF",
                description=f"{alias_label} is an alias of {group.canonical_term}",
                weight=1.0,
                source=group.source,
            )

    return list(nodes.values()), list(relations.values())


def _parse_answer_graph_llm_result(
    raw: str,
    config: AnswerGraphConfig,
    *,
    source: str = "llm",
) -> tuple[list[AnswerGraphNode], list[AnswerGraphRelation]]:
    parsed = _coerce_json(raw.strip().removeprefix("```json").removesuffix("```").strip(), {})
    if not isinstance(parsed, dict):
        return [], []
    allowed_entities = set(config.entity_types)
    allowed_relations = set(config.relation_types)
    nodes: dict[tuple[str, str], AnswerGraphNode] = {}
    for item in parsed.get("entities", [])[:40]:
        if not isinstance(item, dict):
            continue
        entity_type = str(item.get("type") or "")
        label = _graph_clean_label(item.get("label"))
        if entity_type not in allowed_entities or entity_type == "FAQAnswer" or len(label) < 2:
            continue
        node = AnswerGraphNode(
            node_id=_graph_node_id(entity_type, label),
            entity_type=entity_type,
            label=label,
            description=_graph_clean_label(item.get("description"), 500),
            source=source,
        )
        nodes[(entity_type, _normalise_text(label))] = node

    relations: list[AnswerGraphRelation] = []
    for item in parsed.get("relations", [])[:60]:
        if not isinstance(item, dict):
            continue
        relation_type = str(item.get("type") or "")
        source_type = str(item.get("source_type") or "")
        target_type = str(item.get("target_type") or "")
        source_label = _graph_clean_label(item.get("source_label"))
        target_label = _graph_clean_label(item.get("target_label"))
        if (
            relation_type not in allowed_relations
            or source_type not in allowed_entities
            or target_type not in allowed_entities
            or len(source_label) < 2
            or len(target_label) < 2
        ):
            continue
        source_id = _graph_node_id(source_type, source_label)
        target_id = _graph_node_id(target_type, target_label)
        if source_type != "FAQAnswer":
            nodes.setdefault(
                (source_type, _normalise_text(source_label)),
                AnswerGraphNode(
                    node_id=source_id,
                    entity_type=source_type,
                    label=source_label,
                    source=source,
                ),
            )
        if target_type != "FAQAnswer":
            nodes.setdefault(
                (target_type, _normalise_text(target_label)),
                AnswerGraphNode(
                    node_id=target_id,
                    entity_type=target_type,
                    label=target_label,
                    source=source,
                ),
            )
        relations.append(
            AnswerGraphRelation(
                source_id=source_id,
                target_id=target_id,
                relation_type=relation_type,
                description=_graph_clean_label(item.get("description"), 500),
                source=source,
            )
        )
    return list(nodes.values()), relations


def _answer_graph_llm_quality(
    answer_node_id: str,
    nodes: list[AnswerGraphNode],
    relations: list[AnswerGraphRelation],
    config: AnswerGraphConfig,
    expected_relation_types: set[str] | None = None,
) -> tuple[bool, tuple[int, int, int, int, int], list[str]]:
    connected = [
        relation
        for relation in relations
        if answer_node_id in {relation.source_id, relation.target_id}
    ]
    relation_types = {relation.relation_type for relation in connected}
    expected = expected_relation_types or set()
    satisfied_expected = expected.intersection(relation_types)
    reasons: list[str] = []
    if len(connected) < config.ai_min_relations:
        reasons.append(
            f"answer_relations={len(connected)}<{config.ai_min_relations}"
        )
    if len(relation_types) < config.ai_min_relation_types:
        reasons.append(
            f"relation_types={len(relation_types)}<{config.ai_min_relation_types}"
        )
    missing_expected = sorted(expected - relation_types)
    if missing_expected:
        reasons.append(f"missing_expected={','.join(missing_expected)}")
    score = (
        1 if not reasons else 0,
        len(satisfied_expected),
        len(relation_types),
        len(connected),
        len(nodes) + len(relations),
    )
    return not reasons, score, reasons


def _expected_answer_graph_relation_types(
    answer: AnswerItem,
    config: AnswerGraphConfig,
) -> set[str]:
    allowed = set(config.relation_types)
    expected = {"REPRESENTS_QUESTION"}.intersection(allowed)
    title = _normalise_text(answer.title)
    if "HAS_SYMPTOM" in allowed and re.search(
        r"(안\s|않|오류|실패|불가|끊|느리|없|분실|잠김|인식|신호|소진|부족|못|문제)",
        title,
    ):
        expected.add("HAS_SYMPTOM")
    if "RESOLVES" in allowed and re.search(
        r"(방법|싶|신청|변경|교체|발급|해지|정지|차단|환불|등록|전환|확인|조회|납부|설정)",
        title,
    ):
        expected.add("RESOLVES")
    return expected


async def _llm_answer_graph_projection(
    rag,
    answer: AnswerItem,
    guidance: list[AnswerGuidance],
    config: AnswerGraphConfig,
) -> tuple[list[AnswerGraphNode], list[AnswerGraphRelation]]:
    llm_func = getattr(rag, "llm_model_func", None)
    if llm_func is None:
        return [], []
    answer_node_id = _answer_graph_node_id(answer.answer_id)
    expected_relation_types = _expected_answer_graph_relation_types(answer, config)
    payload = {
        "answer_node": {
            "type": "FAQAnswer",
            "label": answer.answer_id,
            "node_id": answer_node_id,
        },
        "title": answer.title,
        "body": answer.body[:4000],
        "tags": answer.tags,
        "metadata": answer.metadata,
        "assets": [
            {
                "type": asset.asset_type,
                "caption": asset.caption,
                "alt_text": asset.alt_text,
                "search_text": asset.search_text,
                "content_text": asset.content_text,
            }
            for asset in answer.assets
            if asset.is_active
        ],
        "finding_hints": [
            {"type": item.guidance_type, "text": item.text}
            for item in guidance
            if item.guidance_type != "negative_keyword"
        ],
        "allowed_entity_types": config.entity_types,
        "allowed_relation_types": config.relation_types,
    }
    base_prompt = (
        f"{config.extraction_prompt}\n\n"
        "Use the supplied FAQAnswer label exactly for every relation connected to the answer.\n"
        "Relation rules: REPRESENTS_QUESTION connects the FAQAnswer to its intent; "
        "HAS_TERM connects important domain terms; HAS_SYMPTOM is only for an explicit "
        "symptom; APPLIES_TO is only for a named product or system; BELONGS_TO is for "
        "categories; RESOLVES is for procedures or actions in the answer. Prefer three "
        "to eight supported relations and at least two relation types.\n"
        f"INPUT:\n{json.dumps(payload, ensure_ascii=False)}"
    )

    async def extract(*, deep: bool) -> tuple[list[AnswerGraphNode], list[AnswerGraphRelation]]:
        prompt = base_prompt
        source = "llm"
        options: dict[str, Any] = {"_llm_thinking_override": deep}
        if deep:
            source = "llm_deep"
            options["max_tokens"] = config.ai_retry_max_tokens
            prompt = (
                f"{base_prompt}\n\n"
                "Review the FAQ carefully before returning the final JSON. Check that "
                "the FAQAnswer is connected to its intent and, when supported, its "
                "terms, symptoms, product/system, category, and resolution procedure. "
                "The following relation types are expected for this FAQ when supported "
                f"by its title: {', '.join(sorted(expected_relation_types))}."
            )
        raw = await llm_func(
            prompt,
            system_prompt="You build bounded FAQ retrieval graphs and return valid JSON only.",
            _llm_purpose="knowledge_structure",
            **options,
        )
        return _parse_answer_graph_llm_result(str(raw), config, source=source)

    strategy = config.ai_extraction_strategy
    if strategy == "deep":
        try:
            return await extract(deep=True)
        except Exception as exc:
            logger.warning(
                "[Answers] Deep FAQ graph extraction failed for %s; retrying fast: %s",
                answer.answer_id,
                exc,
            )
            try:
                return await extract(deep=False)
            except Exception as fallback_exc:
                logger.warning(
                    "[Answers] FAQ graph fallback extraction failed for %s: %s",
                    answer.answer_id,
                    fallback_exc,
                )
                return [], []

    try:
        fast_nodes, fast_relations = await extract(deep=False)
    except Exception as exc:
        logger.warning(
            "[Answers] FAQ graph LLM extraction failed for %s: %s",
            answer.answer_id,
            exc,
        )
        fast_nodes, fast_relations = [], []
    if strategy == "fast":
        return fast_nodes, fast_relations

    fast_valid, fast_score, reasons = _answer_graph_llm_quality(
        answer_node_id,
        fast_nodes,
        fast_relations,
        config,
        expected_relation_types,
    )
    if fast_valid:
        return fast_nodes, fast_relations

    logger.info(
        "[Answers] Adaptive FAQ graph retry answer=%s reasons=%s",
        answer.answer_id,
        ",".join(reasons),
    )
    try:
        deep_nodes, deep_relations = await extract(deep=True)
    except Exception as exc:
        logger.warning(
            "[Answers] Adaptive FAQ graph retry failed for %s; keeping fast result: %s",
            answer.answer_id,
            exc,
        )
        return fast_nodes, fast_relations
    _, deep_score, _ = _answer_graph_llm_quality(
        answer_node_id,
        deep_nodes,
        deep_relations,
        config,
        expected_relation_types,
    )
    if deep_score > fast_score:
        return deep_nodes, deep_relations
    return fast_nodes, fast_relations


async def _build_answer_graph_preview(
    db,
    workspace: str,
    rag,
    answer: AnswerItem,
    guidance: list[AnswerGuidance],
    config: AnswerGraphConfig,
    *,
    use_llm: bool,
) -> AnswerGraphProjection:
    aliases = await _answer_alias_groups(db, workspace)
    nodes, relations = _deterministic_answer_graph_projection(
        answer, guidance, aliases, config
    )
    if use_llm:
        llm_nodes, llm_relations = await _llm_answer_graph_projection(
            rag, answer, guidance, config
        )
        node_map = {node.node_id: node for node in nodes}
        node_map.update({node.node_id: node for node in llm_nodes})
        relation_map = {
            (
                relation.source_id,
                relation.target_id,
                relation.relation_type,
            ): relation
            for relation in relations
        }
        relation_map.update(
            {
                (
                    relation.source_id,
                    relation.target_id,
                    relation.relation_type,
                ): relation
                for relation in llm_relations
            }
        )
        nodes = list(node_map.values())
        relations = list(relation_map.values())
    return AnswerGraphProjection(
        workspace=workspace,
        answer_id=answer.answer_id,
        answer_version=answer.version,
        schema_version=config.schema_version,
        status="preview",
        content_hash=_graph_content_hash(answer, guidance, config),
        nodes=nodes,
        relations=relations,
    )


async def _mark_answer_graph_stale(
    db,
    workspace: str,
    answer_id: str,
    answer_version: int,
) -> None:
    config = await _get_answer_graph_config(db, workspace)
    await db.query(
        """
        INSERT INTO LIGHTRAG_ANSWER_GRAPH_PROJECTIONS
            (
                workspace, answer_id, answer_version, schema_version,
                status, content_hash, nodes, relations
            )
        VALUES ($1, $2, $3, $4, 'stale', '', '[]'::jsonb, '[]'::jsonb)
        ON CONFLICT (workspace, answer_id)
        DO UPDATE SET
            answer_version = EXCLUDED.answer_version,
            schema_version = EXCLUDED.schema_version,
            status = 'stale',
            error = NULL,
            update_time = NOW()
        """,
        [workspace, answer_id, answer_version, config.schema_version],
    )


async def _mark_all_answer_graph_stale(db, workspace: str) -> None:
    await db.execute(
        """
        UPDATE LIGHTRAG_ANSWER_GRAPH_PROJECTIONS
        SET status = 'stale', error = NULL, update_time = NOW()
        WHERE workspace = $1
        """,
        {"workspace": workspace},
    )


async def _upsert_graph_entity(rag, node: AnswerGraphNode, source_id: str) -> None:
    storage = rag.chunk_entity_relation_graph
    if await storage.has_node(node.node_id):
        if node.entity_type == "FAQAnswer":
            await rag.aedit_entity(
                node.node_id,
                {
                    "entity_type": node.entity_type,
                    "description": node.description,
                    "source_id": source_id,
                    "file_path": "faq_graph",
                },
                allow_rename=False,
            )
        return
    await rag.acreate_entity(
        node.node_id,
        {
            "entity_type": node.entity_type,
            "description": node.description,
            "source_id": source_id,
            "file_path": "faq_graph",
        },
    )


async def _upsert_graph_relation(
    rag,
    relation: AnswerGraphRelation,
    source_id: str,
) -> None:
    storage = rag.chunk_entity_relation_graph
    data = {
        "description": relation.description,
        "keywords": relation.relation_type,
        "source_id": source_id,
        "weight": relation.weight,
        "file_path": "faq_graph",
    }
    if await storage.has_edge(relation.source_id, relation.target_id):
        await rag.aedit_relation(relation.source_id, relation.target_id, data)
    else:
        await rag.acreate_relation(relation.source_id, relation.target_id, data)


async def _persist_answer_graph_projection(
    *,
    db,
    workspace: str,
    rag,
    projection: AnswerGraphProjection,
) -> AnswerGraphProjection:
    answer_node_id = _answer_graph_node_id(projection.answer_id)
    source_id = (
        f"faq-graph:{projection.answer_id}:v{projection.answer_version}:"
        f"s{projection.schema_version}"
    )
    try:
        if await rag.chunk_entity_relation_graph.has_node(answer_node_id):
            await rag.adelete_by_entity(answer_node_id)
        for node in projection.nodes:
            await _upsert_graph_entity(rag, node, source_id)
        for relation in projection.relations:
            if not await rag.chunk_entity_relation_graph.has_node(relation.source_id):
                continue
            if not await rag.chunk_entity_relation_graph.has_node(relation.target_id):
                continue
            await _upsert_graph_relation(rag, relation, source_id)

        projection.status = "ready"
        projection.built_at = _iso(_now())
        projection.error = None
        await db.query(
            """
            INSERT INTO LIGHTRAG_ANSWER_GRAPH_PROJECTIONS
                (
                    workspace, answer_id, answer_version, schema_version,
                    status, content_hash, nodes, relations, error, built_at
                )
            VALUES ($1, $2, $3, $4, 'ready', $5, $6::jsonb, $7::jsonb, NULL, NOW())
            ON CONFLICT (workspace, answer_id)
            DO UPDATE SET
                answer_version = EXCLUDED.answer_version,
                schema_version = EXCLUDED.schema_version,
                status = 'ready',
                content_hash = EXCLUDED.content_hash,
                nodes = EXCLUDED.nodes,
                relations = EXCLUDED.relations,
                error = NULL,
                built_at = NOW(),
                update_time = NOW()
            """,
            [
                workspace,
                projection.answer_id,
                projection.answer_version,
                projection.schema_version,
                projection.content_hash,
                _json_list([item.model_dump() for item in projection.nodes]),
                _json_list([item.model_dump() for item in projection.relations]),
            ],
        )
        return projection
    except Exception as exc:
        projection.status = "failed"
        projection.error = str(exc)
        await db.query(
            """
            INSERT INTO LIGHTRAG_ANSWER_GRAPH_PROJECTIONS
                (
                    workspace, answer_id, answer_version, schema_version,
                    status, content_hash, nodes, relations, error
                )
            VALUES ($1, $2, $3, $4, 'failed', $5, $6::jsonb, $7::jsonb, $8)
            ON CONFLICT (workspace, answer_id)
            DO UPDATE SET
                answer_version = EXCLUDED.answer_version,
                schema_version = EXCLUDED.schema_version,
                status = 'failed',
                content_hash = EXCLUDED.content_hash,
                nodes = EXCLUDED.nodes,
                relations = EXCLUDED.relations,
                error = EXCLUDED.error,
                update_time = NOW()
            """,
            [
                workspace,
                projection.answer_id,
                projection.answer_version,
                projection.schema_version,
                projection.content_hash,
                _json_list([item.model_dump() for item in projection.nodes]),
                _json_list([item.model_dump() for item in projection.relations]),
                str(exc),
            ],
        )
        raise


async def _answer_rows_and_guidance(
    db,
    workspace: str,
    answer_ids: list[str],
) -> tuple[list[AnswerItem], dict[str, list[AnswerGuidance]]]:
    if not answer_ids:
        return [], {}
    rows = await db.query(
        """
        SELECT workspace, answer_id, title, body, approved_summary, content_format,
               display_policy, status, version, valid_from, valid_until, priority,
               tags, metadata, publish_time, create_time, update_time
        FROM LIGHTRAG_ANSWER_ITEMS
        WHERE workspace = $1 AND answer_id = ANY($2::text[])
        ORDER BY update_time DESC, answer_id
        """,
        [workspace, answer_ids],
        multirows=True,
    )
    guidance_rows = await db.query(
        """
        SELECT guidance_id, workspace, answer_id, guidance_type, text, weight,
               metadata, create_time
        FROM LIGHTRAG_ANSWER_GUIDANCE
        WHERE workspace = $1 AND answer_id = ANY($2::text[])
        """,
        [workspace, answer_ids],
        multirows=True,
    )
    guidance_by_answer: dict[str, list[AnswerGuidance]] = {}
    for row in guidance_rows or []:
        guidance = _guidance_from_row(dict(row))
        guidance_by_answer.setdefault(guidance.answer_id, []).append(guidance)
    answers = [_answer_from_row(dict(row)) for row in rows or []]
    await _hydrate_answer_assets(db, workspace, answers)
    return answers, guidance_by_answer


async def _rebuild_answer_graph_for_ids(
    *,
    db,
    workspace: str,
    rag,
    answer_ids: list[str],
    use_llm: bool,
    task_id: Optional[str] = None,
    trigger_source: str = "manual_api",
) -> dict[str, Any]:
    config = await _get_answer_graph_config(db, workspace)
    answers, guidance_by_answer = await _answer_rows_and_guidance(
        db, workspace, answer_ids
    )
    service = None
    if task_id:
        from lightrag.api.task_manager import TaskStatus, get_task_service

        service = get_task_service()
    rebuilt: list[str] = []
    failed: dict[str, str] = {}
    total = len(answers)
    processed = 0
    cancelled = False
    logger.info(
        "[Answers] FAQ graph rebuild started workspace=%s trigger=%s task=%s "
        "answers=%d schema_version=%d use_llm=%s",
        workspace,
        trigger_source,
        task_id or "-",
        total,
        config.schema_version,
        use_llm,
    )
    try:
        rebuild_concurrency = int(os.getenv("FAQ_GRAPH_REBUILD_CONCURRENCY", "4"))
    except ValueError:
        rebuild_concurrency = 4
    rebuild_concurrency = max(1, min(rebuild_concurrency, 16))
    semaphore = asyncio.Semaphore(rebuild_concurrency)
    logger.info(
        "[Answers] FAQ graph rebuild concurrency workspace=%s task=%s concurrency=%d",
        workspace,
        task_id or "-",
        rebuild_concurrency,
    )

    async def build_one(answer: AnswerItem) -> tuple[str, Optional[str], bool]:
        async with semaphore:
            if service and task_id:
                current_task = service.get_task(task_id)
                if current_task and current_task.status == TaskStatus.CANCELLED:
                    return answer.answer_id, None, True
            for attempt in range(2):
                try:
                    projection = await _build_answer_graph_preview(
                        db,
                        workspace,
                        rag,
                        answer,
                        guidance_by_answer.get(answer.answer_id, []),
                        config,
                        use_llm=use_llm,
                    )
                    await _persist_answer_graph_projection(
                        db=db,
                        workspace=workspace,
                        rag=rag,
                        projection=projection,
                    )
                    return answer.answer_id, None, False
                except asyncio.CancelledError:
                    raise
                except Exception as exc:
                    duplicate_race = "already exists" in str(exc).lower()
                    if attempt == 0 and duplicate_race:
                        logger.warning(
                            "[Answers] Retrying FAQ graph projection after a concurrent "
                            "entity write workspace=%s answer=%s: %s",
                            workspace,
                            answer.answer_id,
                            exc,
                        )
                        await asyncio.sleep(0.25)
                        continue
                    logger.error(
                        "[Answers] FAQ graph projection failed workspace=%s answer=%s: %s",
                        workspace,
                        answer.answer_id,
                        exc,
                    )
                    return answer.answer_id, str(exc), False
            return answer.answer_id, "FAQ graph projection retry exhausted", False

    build_tasks = [asyncio.create_task(build_one(answer)) for answer in answers]
    try:
        for completed in asyncio.as_completed(build_tasks):
            answer_id, error, task_cancelled = await completed
            if task_cancelled:
                cancelled = True
                break
            processed += 1
            if error:
                failed[answer_id] = error
            else:
                rebuilt.append(answer_id)
            if service and task_id:
                current_task = service.get_task(task_id)
                if current_task and current_task.status == TaskStatus.CANCELLED:
                    cancelled = True
                    break
                await service.update_progress(
                    task_id,
                    min(99.0, (processed / max(total, 1)) * 95.0),
                    f"FAQ graph {processed}/{total} prepared",
                    detail={
                        "graph_processed": processed,
                        "graph_rebuilt": len(rebuilt),
                        "graph_failed": len(failed),
                    },
                )
    finally:
        if cancelled:
            for build_task in build_tasks:
                if not build_task.done():
                    build_task.cancel()
        await asyncio.gather(*build_tasks, return_exceptions=True)
    result = {
        "processed": processed,
        "total": total,
        "rebuilt": rebuilt,
        "failed": failed,
        "cancelled": cancelled,
        "schema_version": config.schema_version,
        "trigger_source": trigger_source,
        "task_id": task_id,
    }
    logger.info(
        "[Answers] FAQ graph rebuild finished workspace=%s trigger=%s task=%s "
        "processed=%d rebuilt=%d failed=%d cancelled=%s schema_version=%d",
        workspace,
        trigger_source,
        task_id or "-",
        processed,
        len(rebuilt),
        len(failed),
        cancelled,
        config.schema_version,
    )
    return result


async def _rebuild_answer_graph_background(
    *,
    task_id: str,
    workspace: str,
    db,
    rag,
    answer_ids: list[str],
    use_llm: bool,
) -> None:
    from lightrag.api.task_manager import TaskStatus, get_task_service

    service = get_task_service()
    result = await _rebuild_answer_graph_for_ids(
        db=db,
        workspace=workspace,
        rag=rag,
        answer_ids=answer_ids,
        use_llm=use_llm,
        task_id=task_id,
        trigger_source="manual_api",
    )
    current_task = service.get_task(task_id)
    if result["cancelled"] or (
        current_task and current_task.status == TaskStatus.CANCELLED
    ):
        return
    await service.complete_task(task_id, result=result)


async def _refresh_answer_graph_safely(
    *,
    db,
    workspace: str,
    answer_id: str,
    trigger_source: str,
) -> None:
    try:
        config = await _get_answer_graph_config(db, workspace)
        if not (config.enabled and config.auto_sync):
            return
        rag = await get_workspace_rag(workspace)
        if rag is None:
            return
        await _rebuild_answer_graph_for_ids(
            db=db,
            workspace=workspace,
            rag=rag,
            answer_ids=[answer_id],
            use_llm=False,
            trigger_source=trigger_source,
        )
    except Exception as exc:
        logger.warning(
            "[Answers] Automatic graph refresh failed workspace=%s answer=%s "
            "trigger=%s: %s",
            workspace,
            answer_id,
            trigger_source,
            exc,
        )


def _graph_result_similarity(item: dict[str, Any]) -> float:
    return _safe_float(
        item.get("similarity", item.get("distance")),
        0.0,
        0.0,
        1.0,
    )


async def _answer_graph_scores(
    db,
    workspace: str,
    rag,
    query: str,
    top_k: int,
    allowed_answer_ids: Optional[list[str]],
    config: AnswerGraphConfig,
) -> tuple[dict[str, float], dict[str, list[AnswerGraphEvidence]], str]:
    if not config.enabled:
        return {}, {}, "graph_disabled"
    if rag is None or getattr(rag, "entities_vdb", None) is None:
        return {}, {}, "graph_unavailable"
    try:
        results = await rag.entities_vdb.query(query, top_k=max(12, top_k))
    except Exception as exc:
        logger.warning("[Answers] FAQ graph query failed: %s", exc)
        return {}, {}, "graph_query_failed"
    faq_results = [
        (
            str(item.get("entity_name") or ""),
            _graph_result_similarity(item),
        )
        for item in results or []
        if str(item.get("entity_name") or "").startswith(f"{ANSWER_GRAPH_NODE_PREFIX}::")
    ]
    if not faq_results:
        return {}, {}, "graph_no_nodes"

    answer_scores: dict[str, float] = {}
    evidence: dict[str, list[AnswerGraphEvidence]] = {}
    eligible_start_nodes = 0
    for start_node, start_score in faq_results:
        if start_score < config.min_similarity:
            continue
        eligible_start_nodes += 1
        queue: list[tuple[str, list[str], list[str], int]] = [
            (start_node, [start_node], [], 0)
        ]
        visited = {start_node}
        while queue:
            node_id, path, relation_types, depth = queue.pop(0)
            answer_id = _answer_id_from_graph_node(node_id)
            if answer_id:
                if allowed_answer_ids is None or answer_id in allowed_answer_ids:
                    score = round(start_score * (0.86**depth), 4)
                    if score > answer_scores.get(answer_id, 0.0):
                        answer_scores[answer_id] = score
                    evidence.setdefault(answer_id, []).append(
                        AnswerGraphEvidence(
                            answer_id=answer_id,
                            score=score,
                            start_node=_graph_label_from_node_id(start_node),
                            path=[_graph_label_from_node_id(item) for item in path],
                            relation_types=relation_types,
                        )
                    )
                continue
            if depth >= config.max_hops:
                continue
            edges = (
                await rag.chunk_entity_relation_graph.get_node_edges(node_id) or []
            )
            for source_id, target_id in edges:
                neighbour = target_id if source_id == node_id else source_id
                if neighbour in visited or not neighbour.startswith(
                    f"{ANSWER_GRAPH_NODE_PREFIX}::"
                ):
                    continue
                edge_data = await rag.chunk_entity_relation_graph.get_edge(
                    source_id, target_id
                )
                relation_type = str((edge_data or {}).get("keywords") or "RELATED_TO")
                visited.add(neighbour)
                queue.append(
                    (
                        neighbour,
                        [*path, neighbour],
                        [*relation_types, relation_type],
                        depth + 1,
                    )
                )

    if eligible_start_nodes == 0:
        return {}, {}, "graph_below_similarity"
    if not answer_scores:
        return {}, {}, "graph_no_answers"
    ready_rows = await db.query(
        """
        SELECT projections.answer_id
        FROM LIGHTRAG_ANSWER_GRAPH_PROJECTIONS AS projections
        JOIN LIGHTRAG_ANSWER_ITEMS AS answers
          ON answers.workspace = projections.workspace
         AND answers.answer_id = projections.answer_id
         AND answers.version = projections.answer_version
        WHERE projections.workspace = $1
          AND projections.status = 'ready'
          AND projections.schema_version = $2
          AND projections.answer_id = ANY($3::text[])
        """,
        [workspace, config.schema_version, list(answer_scores)],
        multirows=True,
    )
    ready_ids = {str(row["answer_id"]) for row in ready_rows or []}
    return (
        {answer_id: score for answer_id, score in answer_scores.items() if answer_id in ready_ids},
        {answer_id: paths[:3] for answer_id, paths in evidence.items() if answer_id in ready_ids},
        "graph_ready" if ready_ids else "graph_stale",
    )


def _effective_answer_retrieval(
    requested_mode: RetrievalMode,
    graph_status: str,
) -> tuple[RetrievalMode, Optional[str]]:
    if requested_mode != "graph_hybrid":
        return requested_mode, None
    if graph_status == "graph_ready":
        return requested_mode, None
    return "hybrid", graph_status


def _answer_vector_content(answer: AnswerItem, guidance: list[AnswerGuidance]) -> str:
    guidance_text = "\n".join(
        f"{item.guidance_type}: {item.text}"
        for item in guidance
        if item.guidance_type != "negative_keyword"
    )
    tags = ", ".join(answer.tags)
    body_preview = answer.body[:1600]
    asset_text = _answer_asset_search_content(answer.assets)[:1600]
    return "\n".join(
        part
        for part in [
            f"Title: {answer.title}",
            f"Tags: {tags}",
            f"Matching hints:\n{guidance_text}",
            f"Answer:\n{body_preview}",
            f"Attached media descriptions:\n{asset_text}",
        ]
        if part.strip()
    )


def _answer_vector_hash(answer: AnswerItem, content: str) -> str:
    raw = f"{answer.answer_id}:{answer.version}:{content}"
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


async def _delete_answer_vectors(db, workspace: str, answer_id: str) -> None:
    await db.execute(
        """
        DELETE FROM LIGHTRAG_ANSWER_VECTORS
        WHERE workspace = $1 AND answer_id = $2
        """,
        {"workspace": workspace, "answer_id": answer_id},
    )


async def _ensure_answer_vector(
    db,
    workspace: str,
    rag,
    answer: AnswerItem,
    guidance: list[AnswerGuidance],
) -> Optional[list[float]]:
    content = _answer_vector_content(answer, guidance)
    content_hash = _answer_vector_hash(answer, content)
    existing = await db.query(
        """
        SELECT embedding, content_hash, metadata
        FROM LIGHTRAG_ANSWER_VECTORS
        WHERE workspace = $1 AND answer_id = $2 AND vector_kind = 'combined'
        """,
        [workspace, answer.answer_id],
    )
    existing_metadata = _coerce_json(existing.get("metadata"), {}) if existing else {}
    if (
        existing
        and existing.get("content_hash") == content_hash
        and existing_metadata.get("projection_version") == ANSWER_VECTOR_PROJECTION_VERSION
    ):
        embedding = _coerce_json(existing.get("embedding"), [])
        return _embedding_to_list(embedding)

    embedding_func = _embedding_func_from_rag(rag)
    if embedding_func is None:
        return None

    try:
        embedding_result = await embedding_func([content])
        embedding = _embedding_to_list(embedding_result)
    except Exception as exc:
        logger.warning("[Answers] Failed to build answer vector for %s: %s", answer.answer_id, exc)
        return None

    if not embedding:
        return None

    await db.query(
        """
        INSERT INTO LIGHTRAG_ANSWER_VECTORS
            (
                vector_id, workspace, answer_id, answer_version, vector_kind,
                content_hash, content, embedding, embedding_vector, metadata
            )
        VALUES ($1, $2, $3, $4, 'combined', $5, $6, $7::jsonb, $8::vector, $9::jsonb)
        ON CONFLICT (workspace, answer_id, vector_kind)
        DO UPDATE SET
            answer_version = EXCLUDED.answer_version,
            content_hash = EXCLUDED.content_hash,
            content = EXCLUDED.content,
            embedding = EXCLUDED.embedding,
            embedding_vector = EXCLUDED.embedding_vector,
            metadata = EXCLUDED.metadata,
            update_time = NOW()
        """,
        [
            f"avec-{uuid.uuid4().hex}",
            workspace,
            answer.answer_id,
            answer.version,
            content_hash,
            content,
            _json_list(embedding),
            _embedding_to_pgvector(embedding),
            _json(
                {
                    "source": "answer_catalog_hybrid",
                    "projection_version": ANSWER_VECTOR_PROJECTION_VERSION,
                }
            ),
        ],
    )
    return embedding


async def _rebuild_answer_vectors_for_ids(
    *,
    db,
    workspace: str,
    rag,
    answer_ids: list[str],
    task_id: Optional[str] = None,
    progress_start: float = 0.0,
    progress_span: float = 95.0,
) -> dict[str, Any]:
    if not answer_ids:
        return {"processed": 0, "rebuilt": 0, "failed": []}
    if _embedding_func_from_rag(rag) is None:
        return {
            "processed": 0,
            "rebuilt": 0,
            "failed": list(answer_ids),
            "status": "embedding_unavailable",
        }
    rows = await db.query(
        """
        SELECT workspace, answer_id, title, body, approved_summary, content_format,
               display_policy, status, version, valid_from, valid_until, priority,
               tags, metadata, publish_time, create_time, update_time
        FROM LIGHTRAG_ANSWER_ITEMS
        WHERE workspace = $1 AND answer_id = ANY($2::text[])
        ORDER BY update_time DESC, answer_id
        """,
        [workspace, answer_ids],
        multirows=True,
    )
    existing_ids = [str(row["answer_id"]) for row in rows or []]
    answers = [_answer_from_row(dict(row)) for row in rows or []]
    await _hydrate_answer_assets(db, workspace, answers)
    guidance_rows = await db.query(
        """
        SELECT guidance_id, workspace, answer_id, guidance_type, text, weight,
               metadata, create_time
        FROM LIGHTRAG_ANSWER_GUIDANCE
        WHERE workspace = $1 AND answer_id = ANY($2::text[])
        """,
        [workspace, existing_ids],
        multirows=True,
    )
    guidance_by_answer: dict[str, list[AnswerGuidance]] = {}
    for row in guidance_rows or []:
        guidance = _guidance_from_row(dict(row))
        guidance_by_answer.setdefault(guidance.answer_id, []).append(guidance)

    service = None
    if task_id:
        from lightrag.api.task_manager import get_task_service

        service = get_task_service()
    rebuilt = 0
    failed: list[str] = []
    total = len(answers)
    for index, answer in enumerate(answers):
        embedding = await _ensure_answer_vector(
            db,
            workspace,
            rag,
            answer,
            guidance_by_answer.get(answer.answer_id, []),
        )
        if embedding:
            rebuilt += 1
        else:
            failed.append(answer.answer_id)
        if service and task_id:
            await service.update_progress(
                task_id,
                min(
                    99.0,
                    progress_start + ((index + 1) / max(total, 1)) * progress_span,
                ),
                f"FAQ vector {index + 1}/{total} prepared",
                detail={
                    "vector_processed": index + 1,
                    "vector_rebuilt": rebuilt,
                    "vector_failed": len(failed),
                },
            )
    return {
        "processed": total,
        "rebuilt": rebuilt,
        "failed": failed,
    }


async def _rebuild_answer_vectors_background(
    *,
    task_id: str,
    workspace: str,
    db,
    rag,
    answer_ids: list[str],
) -> None:
    from lightrag.api.task_manager import get_task_service

    service = get_task_service()
    result = await _rebuild_answer_vectors_for_ids(
        db=db,
        workspace=workspace,
        rag=rag,
        answer_ids=answer_ids,
        task_id=task_id,
    )
    await service.complete_task(task_id, result=result)


async def _refresh_answer_vector_safely(
    *,
    db,
    workspace: str,
    answer_id: str,
) -> None:
    try:
        workspace_rag = await get_workspace_rag(workspace)
        if workspace_rag is None:
            return
        await _rebuild_answer_vectors_for_ids(
            db=db,
            workspace=workspace,
            rag=workspace_rag,
            answer_ids=[answer_id],
        )
    except Exception as exc:
        logger.warning(
            "[Answers] Automatic vector refresh failed workspace=%s answer=%s: %s",
            workspace,
            answer_id,
            exc,
        )


async def _answer_vector_scores(
    db,
    workspace: str,
    rag,
    query: str,
    top_k: int,
    status_filter: list[str],
    allowed_answer_ids: Optional[list[str]],
) -> tuple[dict[str, float], str]:
    embedding_func = _embedding_func_from_rag(rag)
    if embedding_func is None:
        return {}, "vector_unavailable"

    try:
        query_embedding = _embedding_to_list(await embedding_func([query]))
    except Exception as exc:
        logger.warning("[Answers] Failed to embed answer query: %s", exc)
        return {}, "vector_query_failed"

    if not query_embedding:
        return {}, "vector_query_empty"

    vector_rows = await db.query(
        """
        SELECT
            vectors.answer_id,
            GREATEST(
                0.0,
                1.0 - (vectors.embedding_vector <=> $4::vector)
            ) AS score
        FROM LIGHTRAG_ANSWER_VECTORS AS vectors
        JOIN LIGHTRAG_ANSWER_ITEMS AS answers
          ON answers.workspace = vectors.workspace
         AND answers.answer_id = vectors.answer_id
         AND answers.version = vectors.answer_version
        WHERE vectors.workspace = $1
          AND vectors.vector_kind = 'combined'
          AND answers.status = ANY($2::text[])
          AND (answers.valid_from IS NULL OR answers.valid_from <= NOW())
          AND (answers.valid_until IS NULL OR answers.valid_until >= NOW())
          AND ($3::text[] IS NULL OR answers.answer_id = ANY($3::text[]))
          AND vectors.embedding_vector IS NOT NULL
          AND COALESCE((vectors.metadata->>'projection_version')::integer, 0) = $7
          AND vector_dims(vectors.embedding_vector) = $5
        ORDER BY vectors.embedding_vector <=> $4::vector
        LIMIT $6
        """,
        [
            workspace,
            status_filter,
            allowed_answer_ids,
            _embedding_to_pgvector(query_embedding),
            len(query_embedding),
            top_k,
            ANSWER_VECTOR_PROJECTION_VERSION,
        ],
        multirows=True,
    )
    if not vector_rows:
        return {}, "vector_not_built"

    scores = {
        str(row["answer_id"]): round(float(row.get("score") or 0.0), 4)
        for row in vector_rows
        if float(row.get("score") or 0.0) > 0
    }
    return scores, "vector_ready"


def _extract_json_object(text: str) -> Optional[dict[str, Any]]:
    if not text:
        return None
    fenced = re.search(r"```(?:json)?\s*(\{.*?\})\s*```", text, flags=re.DOTALL)
    candidate = fenced.group(1) if fenced else text
    if "{" in candidate and "}" in candidate:
        candidate = candidate[candidate.find("{") : candidate.rfind("}") + 1]
    try:
        parsed = json.loads(candidate)
    except json.JSONDecodeError:
        return None
    return parsed if isinstance(parsed, dict) else None


def _extract_json_array(text: str) -> list[Any]:
    if not text:
        return []
    fenced = re.search(r"```(?:json)?\s*(\[.*?\])\s*```", text, flags=re.DOTALL)
    candidate = fenced.group(1) if fenced else text
    if "[" in candidate and "]" in candidate:
        candidate = candidate[candidate.find("[") : candidate.rfind("]") + 1]
    try:
        parsed = json.loads(candidate)
    except json.JSONDecodeError:
        return []
    return parsed if isinstance(parsed, list) else []


def _parse_term_discovery_candidates(
    raw: str,
    known_groups: list[AnswerAliasGroup],
    evidence_lookup: dict[str, dict[str, Any]],
) -> list[dict[str, Any]]:
    member_to_canonical: dict[str, str] = {}
    canonical_members: dict[str, set[str]] = {}
    for group in known_groups:
        canonical_key = _normalise_text(group.canonical_term)
        members = {canonical_key}
        members.update(_normalise_text(alias) for alias in group.aliases)
        canonical_members.setdefault(canonical_key, set()).update(members)
        for member in members:
            member_to_canonical.setdefault(member, group.canonical_term)

    candidates_by_key: dict[str, dict[str, Any]] = {}
    allowed_types = {"synonym", "abbreviation", "neologism"}
    for item in _extract_json_array(raw):
        if not isinstance(item, dict):
            continue
        canonical_term = " ".join(str(item.get("canonical_term") or "").split())
        canonical_key = _normalise_text(canonical_term)
        if len(canonical_term) < 2 or len(canonical_term) > 200:
            continue

        known_canonical = member_to_canonical.get(canonical_key)
        if known_canonical:
            canonical_term = known_canonical
            canonical_key = _normalise_text(known_canonical)

        aliases: list[str] = []
        seen = {canonical_key}
        for raw_alias in item.get("aliases") or []:
            alias = " ".join(str(raw_alias or "").split())
            alias_key = _normalise_text(alias)
            if len(alias) < 2 or len(alias) > 200 or alias_key in seen:
                continue
            existing_canonical = member_to_canonical.get(alias_key)
            if (
                existing_canonical
                and _normalise_text(existing_canonical) != canonical_key
            ):
                continue
            if alias_key in canonical_members.get(canonical_key, set()):
                continue
            seen.add(alias_key)
            aliases.append(alias)
        if not aliases:
            continue

        evidence_refs = item.get("evidence_refs") or []
        evidence = [
            evidence_lookup[str(ref)]
            for ref in evidence_refs
            if str(ref) in evidence_lookup
        ][:20]
        term_type = str(item.get("term_type") or "synonym")
        if term_type not in allowed_types:
            term_type = "synonym"
        confidence = _safe_float(item.get("confidence"), 0.5, 0.0, 1.0)
        rationale = " ".join(str(item.get("rationale") or "").split())[:1000]

        existing = candidates_by_key.get(canonical_key)
        if existing is None:
            candidates_by_key[canonical_key] = {
                "candidate_key": canonical_key,
                "canonical_term": canonical_term,
                "aliases": aliases,
                "term_type": term_type,
                "confidence": confidence,
                "rationale": rationale,
                "evidence": evidence,
            }
            continue

        existing_alias_keys = {
            _normalise_text(alias) for alias in existing["aliases"]
        }
        existing["aliases"].extend(
            alias
            for alias in aliases
            if _normalise_text(alias) not in existing_alias_keys
        )
        existing["confidence"] = max(existing["confidence"], confidence)
        if len(rationale) > len(existing["rationale"]):
            existing["rationale"] = rationale
        existing_evidence_refs = {
            str(entry.get("ref")) for entry in existing["evidence"]
        }
        existing["evidence"].extend(
            entry
            for entry in evidence
            if str(entry.get("ref")) not in existing_evidence_refs
        )

    return list(candidates_by_key.values())


async def _upsert_term_discovery_candidate(
    *,
    db,
    workspace: str,
    candidate: dict[str, Any],
    task_id: str,
) -> bool:
    existing = await db.query(
        """
        SELECT candidate_id, status, aliases, evidence, confidence, rationale
        FROM LIGHTRAG_ANSWER_TERM_CANDIDATES
        WHERE workspace = $1 AND candidate_key = $2
        """,
        [workspace, candidate["candidate_key"]],
    )
    if existing and str(existing.get("status")) in {"approved", "rejected"}:
        return False

    aliases = list(candidate["aliases"])
    evidence = list(candidate["evidence"])
    confidence = float(candidate["confidence"])
    rationale = str(candidate["rationale"])
    if existing:
        existing_aliases = _coerce_json(existing.get("aliases"), [])
        alias_keys = {_normalise_text(alias) for alias in aliases}
        for alias in existing_aliases if isinstance(existing_aliases, list) else []:
            if _normalise_text(alias) not in alias_keys:
                aliases.append(str(alias))
                alias_keys.add(_normalise_text(alias))
        existing_evidence = _coerce_json(existing.get("evidence"), [])
        evidence_refs = {str(item.get("ref")) for item in evidence}
        for item in existing_evidence if isinstance(existing_evidence, list) else []:
            if isinstance(item, dict) and str(item.get("ref")) not in evidence_refs:
                evidence.append(item)
        confidence = max(
            confidence,
            _safe_float(existing.get("confidence"), 0.0, 0.0, 1.0),
        )
        if not rationale:
            rationale = str(existing.get("rationale") or "")

    await db.query(
        """
        INSERT INTO LIGHTRAG_ANSWER_TERM_CANDIDATES
            (
                candidate_id, workspace, candidate_key, canonical_term, aliases,
                term_type, status, confidence, rationale, evidence, source, metadata
            )
        VALUES ($1, $2, $3, $4, $5::jsonb, $6, 'suggested', $7, $8,
                $9::jsonb, 'llm', $10::jsonb)
        ON CONFLICT (workspace, candidate_key)
        DO UPDATE SET
            canonical_term = EXCLUDED.canonical_term,
            aliases = EXCLUDED.aliases,
            term_type = EXCLUDED.term_type,
            confidence = EXCLUDED.confidence,
            rationale = EXCLUDED.rationale,
            evidence = EXCLUDED.evidence,
            metadata = LIGHTRAG_ANSWER_TERM_CANDIDATES.metadata || EXCLUDED.metadata,
            update_time = NOW()
        """,
        [
            str(existing.get("candidate_id"))
            if existing
            else f"aterm-{uuid.uuid4().hex[:16]}",
            workspace,
            candidate["candidate_key"],
            candidate["canonical_term"],
            _json_list(aliases[:30]),
            candidate["term_type"],
            confidence,
            rationale,
            _json(evidence[:30]),
            _json({"last_task_id": task_id}),
        ],
    )
    return True


async def _discover_answer_terms_background(
    *,
    task_id: str,
    workspace: str,
    db,
    rag,
    include_drafts: bool,
    include_no_match_queries: bool,
    answer_limit: int,
    event_limit: int,
    batch_size: int,
) -> None:
    from lightrag.api.task_manager import TaskStatus, get_task_service

    service = get_task_service()
    llm_func = getattr(rag, "llm_model_func", None)
    if llm_func is None:
        raise RuntimeError("LLM is not configured for this workspace")

    statuses = ["published", "draft"] if include_drafts else ["published"]
    answer_rows = await db.query(
        """
        SELECT answer_id, title, LEFT(body, 1200) AS body, tags
        FROM LIGHTRAG_ANSWER_ITEMS
        WHERE workspace = $1 AND status = ANY($2::text[])
        ORDER BY update_time DESC, answer_id
        LIMIT $3
        """,
        [workspace, statuses, answer_limit],
        multirows=True,
    )
    answer_ids = [str(row["answer_id"]) for row in answer_rows or []]
    guidance_rows = []
    if answer_ids:
        guidance_rows = await db.query(
            """
            SELECT answer_id, guidance_type, text
            FROM LIGHTRAG_ANSWER_GUIDANCE
            WHERE workspace = $1
              AND answer_id = ANY($2::text[])
              AND guidance_type IN ('question', 'keyword', 'synonym')
            ORDER BY create_time DESC
            """,
            [workspace, answer_ids],
            multirows=True,
        )
    guidance_by_answer: dict[str, list[str]] = {}
    for row in guidance_rows or []:
        guidance_by_answer.setdefault(str(row["answer_id"]), []).append(
            str(row["text"])
        )

    discovery_items: list[dict[str, Any]] = []
    evidence_lookup: dict[str, dict[str, Any]] = {}
    for row in answer_rows or []:
        ref = f"faq:{row['answer_id']}"
        item = {
            "ref": ref,
            "kind": "faq",
            "answer_id": str(row["answer_id"]),
            "title": str(row.get("title") or ""),
            "body": str(row.get("body") or ""),
            "tags": _coerce_json(row.get("tags"), []),
            "search_hints": guidance_by_answer.get(str(row["answer_id"]), [])[
                :12
            ],
        }
        discovery_items.append(item)
        evidence_lookup[ref] = {
            "ref": ref,
            "kind": "faq",
            "answer_id": str(row["answer_id"]),
            "label": str(row.get("title") or ""),
        }

    if include_no_match_queries and event_limit > 0:
        event_rows = await db.query(
            """
            SELECT event_id, query, create_time
            FROM LIGHTRAG_ANSWER_EVENTS
            WHERE workspace = $1
              AND selected_answer_id IS NULL
              AND query IS NOT NULL
              AND LENGTH(TRIM(query)) >= 2
            ORDER BY create_time DESC
            LIMIT $2
            """,
            [workspace, event_limit],
            multirows=True,
        )
        for row in event_rows or []:
            ref = f"query:{row['event_id']}"
            item = {
                "ref": ref,
                "kind": "unmatched_query",
                "query": str(row.get("query") or ""),
            }
            discovery_items.append(item)
            evidence_lookup[ref] = {
                "ref": ref,
                "kind": "unmatched_query",
                "label": str(row.get("query") or ""),
                "create_time": _iso(row.get("create_time")),
            }

    if not discovery_items:
        await service.complete_task(
            task_id,
            result={
                "analyzed_count": 0,
                "suggested_count": 0,
                "message": "No FAQ or unmatched queries were available.",
            },
        )
        return

    known_groups = await _answer_alias_groups(db, workspace)
    known_payload = [
        {
            "canonical_term": group.canonical_term,
            "aliases": group.aliases,
        }
        for group in known_groups
    ]
    batch_count = (len(discovery_items) + batch_size - 1) // batch_size
    suggested_count = 0
    failed_batches: list[dict[str, Any]] = []
    for batch_index in range(batch_count):
        task = service.get_task(task_id)
        if task and task.status == TaskStatus.CANCELLED:
            return
        batch = discovery_items[
            batch_index * batch_size : (batch_index + 1) * batch_size
        ]
        prompt = (
            "Discover shared Korean FAQ search terms from the evidence below. "
            "Return only a JSON array. Each item must be "
            "{\"canonical_term\": string, \"aliases\": [string], "
            "\"term_type\": \"synonym|abbreviation|neologism\", "
            "\"confidence\": number, \"rationale\": string, "
            "\"evidence_refs\": [string]}. "
            "Propose only expressions that refer to the same product, service, feature, "
            "or concept. Include useful Korean-English names, abbreviations, common spoken "
            "forms, spacing/spelling variants, and new workplace expressions. "
            "Every alias must be safely interchangeable with the canonical term without "
            "changing the user's intent. Never include related actions, symptoms, causes, "
            "resolutions, commands, components, or broad category words as aliases. "
            "Do not repeat a known group unless the evidence contains a genuinely new alias. "
            "Use only provided evidence_refs, do not invent product names, and write the "
            "rationale in Korean.\n\n"
            f"Known term groups:\n{json.dumps(known_payload, ensure_ascii=False)}\n\n"
            f"Evidence:\n{json.dumps(batch, ensure_ascii=False)}"
        )
        try:
            raw = await llm_func(
                prompt,
                system_prompt=(
                    "You are a Korean enterprise terminology curator. Extract conservative "
                    "synonym, abbreviation, and neologism candidates for FAQ retrieval. "
                    "Return strict JSON only."
                ),
                _llm_purpose="knowledge_structure",
            )
            parsed = _parse_term_discovery_candidates(
                str(raw),
                known_groups,
                evidence_lookup,
            )
            for candidate in parsed:
                if await _upsert_term_discovery_candidate(
                    db=db,
                    workspace=workspace,
                    candidate=candidate,
                    task_id=task_id,
                ):
                    suggested_count += 1
        except Exception as exc:
            logger.warning(
                "[Answers] Term discovery failed task=%s batch=%s: %s",
                task_id,
                batch_index + 1,
                exc,
            )
            failed_batches.append(
                {"batch": batch_index + 1, "error": str(exc)[:500]}
            )

        await service.update_progress(
            task_id,
            ((batch_index + 1) / batch_count) * 100.0,
            f"AI term discovery batch {batch_index + 1}/{batch_count} completed",
            detail={
                "analyzed_count": min(
                    (batch_index + 1) * batch_size, len(discovery_items)
                ),
                "suggested_count": suggested_count,
                "failed_batch_count": len(failed_batches),
            },
        )

    if failed_batches and len(failed_batches) == batch_count:
        raise RuntimeError("All AI term discovery batches failed")
    await service.complete_task(
        task_id,
        result={
            "analyzed_count": len(discovery_items),
            "answer_count": len(answer_rows or []),
            "suggested_count": suggested_count,
            "failed_batches": failed_batches,
        },
    )


async def _llm_select_candidate(rag, query: str, candidates: list[ResolveCandidate]) -> tuple[Optional[str], dict[str, Any]]:
    llm_func = getattr(rag, "llm_model_func", None)
    if llm_func is None or not candidates:
        return None, {"status": "llm_unavailable"}

    candidate_payload = [
        {
            "answer_id": item.answer.answer_id,
            "title": item.answer.title,
            "score": item.score,
            "matched_guidance": item.matched_guidance,
            "reason": item.reason,
            "category": _candidate_category(item),
            "graph_evidence": [
                {
                    "score": evidence.score,
                    "path": evidence.path,
                    "relation_types": evidence.relation_types,
                }
                for evidence in item.graph_evidence[:3]
            ],
        }
        for item in candidates
    ]
    system_prompt = (
        "You are a conservative FAQ ID selector. Choose an approved FAQ answer ID "
        "only when the user intent is directly supported by the candidate title, "
        "finding hints, category, and graph evidence. Similar words alone are not "
        "enough. Do not reject a query merely because it is short: a query such as "
        "'유심 안돼' can be valid when its explicit service object maps clearly to one "
        "candidate. However, never fill in an omitted product, service, or task from "
        "the candidate text. Return null when an ordinary customer-support agent "
        "would need to ask what the user means. This applies even when multiple "
        "candidates belong to the same category: '유심 안돼' can mean recognition, "
        "compatibility, activation, or a locked SIM, and '로밍 안돼' can mean "
        "activation, a missing foreign signal, or a post-return network issue. "
        "Do not pick the most general-looking candidate by default. For example, "
        "'카드가 안 돼요' can mean "
        "a SIM card or a payment card, and '신청이 실패했어요' can mean activation, "
        "number portability, roaming, or eSIM registration, so both require null "
        "unless the query itself supplies clarifying context. Return null for "
        "out-of-scope or semantically ambiguous queries. "
        "When answer_id is null because the query is ambiguous, set decision_reason "
        "to 'ambiguous' and provide one concise clarification_question that helps "
        "distinguish the candidate intents. For out-of-scope queries, set "
        "decision_reason to 'out_of_scope' and clarification_question to null. "
        "Write rationale and clarification_question in the same language as the "
        "user query. "
        "Never create answer content. Return only strict JSON."
    )
    prompt = (
        "User query:\n"
        f"{query}\n\n"
        "Candidates:\n"
        f"{json.dumps(candidate_payload, ensure_ascii=False)}\n\n"
        'Return {"answer_id": string|null, "confidence": number, "rationale": string, '
        '"decision_reason": "match"|"ambiguous"|"out_of_scope", '
        '"clarification_question": string|null}. '
        "answer_id must be one of the candidate answer_id values, or null if none match."
    )
    try:
        raw = await llm_func(
            prompt,
            system_prompt=system_prompt,
            _llm_purpose="faq_selection",
        )
    except Exception as exc:
        logger.warning("[Answers] LLM answer selector failed: %s", exc)
        return None, {"status": "llm_failed", "error": str(exc)}

    parsed = _extract_json_object(str(raw))
    if not parsed:
        return None, {"status": "llm_invalid_json", "raw": str(raw)[:500]}

    allowed = {item.answer.answer_id for item in candidates}
    answer_id = parsed.get("answer_id")
    if answer_id is not None and answer_id not in allowed:
        return None, {"status": "llm_id_rejected", "answer_id": answer_id}

    return answer_id, {
        "status": "llm_ready",
        "answer_id": answer_id,
        "confidence": parsed.get("confidence"),
        "rationale": parsed.get("rationale"),
        "decision_reason": parsed.get("decision_reason"),
        "clarification_question": parsed.get("clarification_question"),
    }


def _sentence_candidates(text: str, limit: int = 4) -> list[str]:
    pieces = re.split(r"[\n\r.!?？。]+", text or "")
    return [piece.strip() for piece in pieces if 4 <= len(piece.strip()) <= 80][:limit]


def _heuristic_guidance_suggestions(
    answer: AnswerItem,
    existing_guidance: list[AnswerGuidance],
    query_examples: list[str],
    max_suggestions: int,
) -> list[SourceGuidanceCandidate]:
    seen = {_normalise_text(item.text) for item in existing_guidance}
    suggestions: list[SourceGuidanceCandidate] = []

    def add(guidance_type: GuidanceType, text: Any, weight: float, source: str) -> None:
        value = " ".join(str(text or "").split())
        if not value or len(value) < 2:
            return
        key = _normalise_text(value)
        if key in seen:
            return
        seen.add(key)
        suggestions.append(
            SourceGuidanceCandidate(
                guidance_type=guidance_type,
                text=value[:200],
                weight=weight,
                source=source,
                metadata={"suggested_by": "heuristic"},
            )
        )

    add("question" if "?" in answer.title or "？" in answer.title else "keyword", answer.title, 1.0, "title")
    add("question", f"{answer.title} 문의", 0.9, "title")
    add("question", f"{answer.title} 안내", 0.9, "title")
    for tag in answer.tags:
        add("keyword", tag, 0.85, "tag")
    for example in query_examples:
        add("question", example, 1.0, "query_example")
    for sentence in _sentence_candidates(answer.body, 4):
        add("synonym", sentence, 0.65, "body")

    metadata = answer.metadata if isinstance(answer.metadata, dict) else {}
    for key in ["category", "type", "detail_type", "question", "keywords", "source_title"]:
        value = metadata.get(key)
        if isinstance(value, list):
            for item in value[:6]:
                add("keyword", item, 0.8, f"metadata.{key}")
        elif value:
            add("keyword", value, 0.8, f"metadata.{key}")

    return suggestions[:max_suggestions]


async def _get_answer_row(db, workspace: str, answer_id: str) -> dict[str, Any]:
    row = await db.query(
        """
        SELECT workspace, answer_id, title, body, approved_summary, content_format,
               display_policy, status, version, valid_from, valid_until, priority,
               tags, metadata, publish_time, create_time, update_time
        FROM LIGHTRAG_ANSWER_ITEMS
        WHERE workspace = $1 AND answer_id = $2
        """,
        [workspace, answer_id],
    )
    if not row:
        raise HTTPException(status_code=404, detail=f"Answer '{answer_id}' not found")
    return row


async def _record_revision(db, row: dict[str, Any]) -> None:
    snapshot = dict(row)
    for key, value in list(snapshot.items()):
        if isinstance(value, datetime):
            snapshot[key] = value.isoformat()
    assets_by_answer = await _answer_assets_by_answer(
        db,
        str(row["workspace"]),
        [str(row["answer_id"])],
    )
    snapshot["assets"] = [
        asset.model_dump(mode="json")
        for asset in assets_by_answer.get(str(row["answer_id"]), [])
    ]
    snapshot = _normalise_answer_snapshot(snapshot)
    await db.execute(
        """
        INSERT INTO LIGHTRAG_ANSWER_REVISIONS
            (revision_id, workspace, answer_id, version, snapshot_json)
        VALUES ($1, $2, $3, $4, $5::jsonb)
        """,
        {
            "revision_id": f"rev-{uuid.uuid4().hex}",
            "workspace": row["workspace"],
            "answer_id": row["answer_id"],
            "version": row["version"],
            "snapshot_json": _json(snapshot),
        },
    )


async def _get_answer_asset_row(
    db,
    workspace: str,
    answer_id: str,
    asset_id: str,
    *,
    include_inactive: bool = False,
) -> dict[str, Any]:
    active_filter = "" if include_inactive else "AND is_active = TRUE"
    row = await db.query(
        f"""
        SELECT asset_id, workspace, answer_id, answer_version, asset_type,
               storage_type, storage_uri, file_name, mime_type, file_size,
               caption, alt_text, search_text, content_text, display_order,
               is_active, metadata, create_time, update_time
        FROM LIGHTRAG_ANSWER_ASSETS
        WHERE workspace = $1 AND answer_id = $2 AND asset_id = $3
          {active_filter}
        """,
        [workspace, answer_id, asset_id],
    )
    if not row:
        raise HTTPException(status_code=404, detail=f"Asset '{asset_id}' not found")
    return dict(row)


async def _touch_answer_after_asset_change(
    db,
    workspace: str,
    answer_id: str,
) -> AnswerItem:
    row = await db.query(
        """
        UPDATE LIGHTRAG_ANSWER_ITEMS
        SET version = version + 1,
            update_time = NOW()
        WHERE workspace = $1 AND answer_id = $2
        RETURNING workspace, answer_id, title, body, approved_summary, content_format,
                  display_policy, status, version, valid_from, valid_until, priority,
                  tags, metadata, publish_time, create_time, update_time
        """,
        [workspace, answer_id],
    )
    if not row:
        raise HTTPException(status_code=404, detail=f"Answer '{answer_id}' not found")
    await db.execute(
        """
        UPDATE LIGHTRAG_ANSWER_ASSETS
        SET answer_version = $3,
            update_time = NOW()
        WHERE workspace = $1 AND answer_id = $2 AND is_active = TRUE
        """,
        {
            "workspace": workspace,
            "answer_id": answer_id,
            "answer_version": int(row.get("version") or 1),
        },
    )
    await _record_revision(db, dict(row))
    await _delete_answer_vectors(db, workspace, answer_id)
    asyncio.create_task(
        _refresh_answer_vector_safely(
            db=db,
            workspace=workspace,
            answer_id=answer_id,
        )
    )
    await _mark_answer_graph_stale(
        db,
        workspace,
        answer_id,
        int(row.get("version") or 1),
    )
    asyncio.create_task(
        _refresh_answer_graph_safely(
            db=db,
            workspace=workspace,
            answer_id=answer_id,
            trigger_source="answer_asset_change",
        )
    )
    answer = _answer_from_row(dict(row))
    await _hydrate_answer_assets(db, workspace, [answer])
    return answer


async def _restore_answer_assets(
    db,
    workspace: str,
    answer_id: str,
    answer_version: int,
    snapshot_assets: list[Any],
) -> None:
    await db.execute(
        """
        UPDATE LIGHTRAG_ANSWER_ASSETS
        SET is_active = FALSE,
            update_time = NOW()
        WHERE workspace = $1 AND answer_id = $2
        """,
        {
            "workspace": workspace,
            "answer_id": answer_id,
        },
    )
    for item in snapshot_assets:
        if not isinstance(item, dict) or not str(item.get("asset_id") or "").strip():
            continue
        create_time = item.get("create_time")
        if isinstance(create_time, str):
            try:
                create_time = datetime.fromisoformat(
                    create_time.replace("Z", "+00:00")
                )
            except ValueError:
                create_time = None
        await db.execute(
            """
            INSERT INTO LIGHTRAG_ANSWER_ASSETS
                (asset_id, workspace, answer_id, answer_version, asset_type,
                 storage_type, storage_uri, file_name, mime_type, file_size,
                 caption, alt_text, search_text, content_text, display_order,
                 is_active, metadata, create_time, update_time)
            VALUES
                ($1, $2, $3, $4, $5, $6, $7, $8, $9, $10,
                 $11, $12, $13, $14, $15, TRUE, $16::jsonb,
                 COALESCE($17::timestamptz, NOW()), NOW())
            ON CONFLICT (asset_id) DO UPDATE
            SET workspace = EXCLUDED.workspace,
                answer_id = EXCLUDED.answer_id,
                answer_version = EXCLUDED.answer_version,
                asset_type = EXCLUDED.asset_type,
                storage_type = EXCLUDED.storage_type,
                storage_uri = EXCLUDED.storage_uri,
                file_name = EXCLUDED.file_name,
                mime_type = EXCLUDED.mime_type,
                file_size = EXCLUDED.file_size,
                caption = EXCLUDED.caption,
                alt_text = EXCLUDED.alt_text,
                search_text = EXCLUDED.search_text,
                content_text = EXCLUDED.content_text,
                display_order = EXCLUDED.display_order,
                is_active = TRUE,
                metadata = EXCLUDED.metadata,
                update_time = NOW()
            """,
            {
                "asset_id": str(item["asset_id"]),
                "workspace": workspace,
                "answer_id": answer_id,
                "answer_version": answer_version,
                "asset_type": str(item.get("asset_type") or "file"),
                "storage_type": str(item.get("storage_type") or "external"),
                "storage_uri": item.get("storage_uri"),
                "file_name": item.get("file_name"),
                "mime_type": item.get("mime_type"),
                "file_size": int(item.get("file_size") or 0),
                "caption": item.get("caption"),
                "alt_text": item.get("alt_text"),
                "search_text": item.get("search_text"),
                "content_text": item.get("content_text"),
                "display_order": int(item.get("display_order") or 0),
                "metadata": _json(item.get("metadata")),
                "create_time": create_time,
            },
        )


async def _log_event(
    db,
    workspace: str,
    event_type: str,
    query: Optional[str] = None,
    selected_answer_id: Optional[str] = None,
    candidate_ids: Optional[list[str]] = None,
    scores: Optional[dict[str, float]] = None,
    metadata: Optional[dict[str, Any]] = None,
) -> str:
    event_id = f"evt-{uuid.uuid4().hex}"
    await db.execute(
        """
        INSERT INTO LIGHTRAG_ANSWER_EVENTS
            (event_id, workspace, event_type, query, selected_answer_id, candidate_ids, scores, metadata)
        VALUES ($1, $2, $3, $4, $5, $6::jsonb, $7::jsonb, $8::jsonb)
        """,
        {
            "event_id": event_id,
            "workspace": workspace,
            "event_type": event_type,
            "query": query,
            "selected_answer_id": selected_answer_id,
            "candidate_ids": _json_list(candidate_ids or []),
            "scores": _json(scores or {}),
            "metadata": _json(metadata or {}),
        },
    )
    return event_id


async def _record_source_snapshot(
    db,
    workspace: str,
    *,
    source_type: str,
    source_uri: Optional[str],
    file_name: Optional[str],
    title: Optional[str],
    raw_content: str,
    profile: dict[str, Any],
    metadata: dict[str, Any],
    created_answer_ids: list[str],
    status: str = "captured",
    task_id: Optional[str] = None,
) -> AnswerSourceSnapshot:
    snapshot_id = f"src-{uuid.uuid4().hex}"
    content = raw_content or ""
    row = await db.query(
        """
        INSERT INTO LIGHTRAG_ANSWER_SOURCE_SNAPSHOTS
            (snapshot_id, workspace, source_type, source_uri, file_name, title,
             raw_content, content_hash, content_length, profile, metadata,
             created_answer_ids, status, task_id)
        VALUES
            ($1, $2, $3, $4, $5, $6, $7, $8, $9, $10::jsonb, $11::jsonb,
             $12::jsonb, $13, $14)
        RETURNING snapshot_id, workspace, source_type, source_uri, file_name, title,
                  LEFT(raw_content, 600) AS content_preview, content_hash, content_length,
                  profile, metadata, created_answer_ids, status, task_id, create_time
        """,
        [
            snapshot_id,
            workspace,
            source_type,
            source_uri,
            file_name,
            title,
            content,
            _content_hash(content),
            len(content),
            _json(profile),
            _json(metadata),
            _json_list(created_answer_ids),
            status,
            task_id,
        ],
    )
    if not row:
        raise HTTPException(status_code=500, detail="Failed to record source snapshot")
    return _source_snapshot_from_row(dict(row))


async def _record_source_link(
    db,
    workspace: str,
    *,
    answer_id: str,
    answer_version: int,
    snapshot_id: str,
    link_type: str = "created_from",
    metadata: Optional[dict[str, Any]] = None,
) -> AnswerSourceLink:
    row = await db.query(
        """
        INSERT INTO LIGHTRAG_ANSWER_SOURCE_LINKS
            (link_id, workspace, answer_id, answer_version, snapshot_id, link_type, metadata)
        VALUES ($1, $2, $3, $4, $5, $6, $7::jsonb)
        RETURNING link_id, workspace, answer_id, answer_version, snapshot_id,
                  link_type, metadata, create_time
        """,
        [
            f"asl-{uuid.uuid4().hex}",
            workspace,
            answer_id,
            answer_version,
            snapshot_id,
            link_type,
            _json(metadata or {}),
        ],
    )
    if not row:
        raise HTTPException(status_code=500, detail="Failed to record source link")
    return _source_link_from_row(dict(row))


async def _record_structured_lookup_log(
    db,
    workspace: str,
    *,
    dataset_id: str,
    dataset_title: str,
    pseudo_sql: str,
    filters: list[StructuredFilter],
    result_count: int,
    preview_only: bool,
    latency_ms: float,
    metadata: Optional[dict[str, Any]] = None,
) -> str:
    log_id = f"sll-{uuid.uuid4().hex}"
    filter_payload = [
        {
            "field": item.field,
            "operator": item.operator,
            "value": item.value,
        }
        for item in filters
    ]
    await db.execute(
        """
        INSERT INTO LIGHTRAG_STRUCTURED_LOOKUP_LOGS
            (log_id, workspace, dataset_id, dataset_title, pseudo_sql, filters,
             result_count, preview_only, latency_ms, metadata)
        VALUES ($1, $2, $3, $4, $5, $6::jsonb, $7, $8, $9, $10::jsonb)
        """,
        {
            "log_id": log_id,
            "workspace": workspace,
            "dataset_id": dataset_id,
            "dataset_title": dataset_title,
            "pseudo_sql": pseudo_sql,
            "filters": _json_list(filter_payload),
            "result_count": result_count,
            "preview_only": preview_only,
            "latency_ms": latency_ms,
            "metadata": _json(metadata or {}),
        },
    )
    return log_id


async def _insert_answer_item(
    db,
    workspace: str,
    *,
    answer_id: str,
    title: str,
    body: str,
    approved_summary: Optional[str],
    content_format: ContentFormat,
    display_policy: DisplayPolicy,
    status: AnswerStatus,
    priority: int,
    tags: list[str],
    metadata: dict[str, Any],
    valid_from: Optional[datetime] = None,
    valid_until: Optional[datetime] = None,
) -> AnswerItem:
    row = await db.query(
        """
        INSERT INTO LIGHTRAG_ANSWER_ITEMS
            (workspace, answer_id, title, body, approved_summary, content_format,
             display_policy, status, version, valid_from, valid_until, priority,
             tags, metadata, publish_time)
        VALUES
            ($1, $2, $3, $4, $5, $6, $7, $8, 1, $9, $10, $11,
             $12::jsonb, $13::jsonb, CASE WHEN $8 = 'published' THEN NOW() ELSE NULL END)
        RETURNING workspace, answer_id, title, body, approved_summary, content_format,
                  display_policy, status, version, valid_from, valid_until, priority,
                  tags, metadata, publish_time, create_time, update_time
        """,
        [
            workspace,
            answer_id,
            title,
            body,
            approved_summary,
            content_format,
            display_policy,
            status,
            valid_from,
            valid_until,
            priority,
            _json_list(tags),
            _json(metadata),
        ],
    )
    if not row:
        raise HTTPException(status_code=500, detail="Failed to create answer")
    await _record_revision(db, dict(row))
    return _answer_from_row(dict(row))


async def _insert_answer_guidance(
    db,
    workspace: str,
    answer_id: str,
    candidates: list[SourceGuidanceCandidate],
) -> list[AnswerGuidance]:
    created_guidance: list[AnswerGuidance] = []
    for item in candidates:
        if not item.text.strip():
            continue
        guidance_row = await db.query(
            """
            INSERT INTO LIGHTRAG_ANSWER_GUIDANCE
                (guidance_id, workspace, answer_id, guidance_type, text, weight, metadata)
            VALUES ($1, $2, $3, $4, $5, $6, $7::jsonb)
            RETURNING guidance_id, workspace, answer_id, guidance_type, text, weight, metadata, create_time
            """,
            [
                f"agd-{uuid.uuid4().hex}",
                workspace,
                answer_id,
                item.guidance_type,
                item.text.strip(),
                item.weight,
                _json({**item.metadata, "source": item.source}),
            ],
        )
        if guidance_row:
            created_guidance.append(_guidance_from_row(dict(guidance_row)))
    return created_guidance


def _guidance_needs_llm_enrichment(guidance: list[AnswerGuidance]) -> bool:
    positive_types = {
        item.guidance_type
        for item in guidance
        if item.guidance_type not in {"note", "negative_keyword"}
    }
    return "question" not in positive_types or not (
        positive_types & {"keyword", "synonym"}
    )


def _guidance_needs_coverage_enrichment(
    answer: AnswerItem,
    guidance: list[AnswerGuidance],
) -> bool:
    positive = [
        item
        for item in guidance
        if item.guidance_type not in {"note", "negative_keyword"}
    ]
    question_count = sum(item.guidance_type == "question" for item in positive)
    synonym_count = sum(item.guidance_type == "synonym" for item in positive)
    source_text = " ".join(
        [answer.title, *answer.tags, *(item.text for item in positive)]
    )
    has_latin_term = re.search(r"\b[A-Za-z][A-Za-z0-9.+-]{2,}\b", source_text) is not None
    has_korean_synonym = any(
        item.guidance_type == "synonym" and re.search(r"[가-힣]", item.text)
        for item in positive
    )
    return (
        len(positive) < 6
        or question_count < 2
        or synonym_count < 2
        or (has_latin_term and not has_korean_synonym)
    )


def _parse_batch_guidance_suggestions(
    raw: str,
    allowed_answer_ids: set[str],
    existing_texts: dict[str, set[str]],
    max_suggestions: int,
    task_id: str,
) -> dict[str, list[SourceGuidanceCandidate]]:
    parsed: dict[str, list[SourceGuidanceCandidate]] = {
        answer_id: [] for answer_id in allowed_answer_ids
    }
    allowed_types = {"question", "keyword", "synonym"}
    for item in _extract_json_array(raw):
        if not isinstance(item, dict):
            continue
        answer_id = str(item.get("answer_id") or "").strip()
        if answer_id not in allowed_answer_ids:
            continue
        suggestions = item.get("suggestions")
        if not isinstance(suggestions, list):
            continue
        seen = existing_texts.setdefault(answer_id, set())
        for suggestion in suggestions:
            if not isinstance(suggestion, dict):
                continue
            guidance_type = str(
                suggestion.get("guidance_type") or "keyword"
            ).strip()
            text = " ".join(str(suggestion.get("text") or "").split())
            if guidance_type not in allowed_types or len(text) < 2:
                continue
            normalized = _normalise_text(text)
            if not normalized or normalized in seen:
                continue
            seen.add(normalized)
            parsed[answer_id].append(
                SourceGuidanceCandidate(
                    guidance_type=guidance_type,
                    text=text[:200],
                    weight=_safe_float(
                        suggestion.get("weight"),
                        1.0,
                        0.1,
                        2.0,
                    ),
                    source="llm_batch",
                    metadata={
                        "suggested_by": "llm_batch",
                        "task_id": task_id,
                    },
                )
            )
            if len(parsed[answer_id]) >= max_suggestions:
                break
    return parsed


async def _enrich_answer_guidance_background(
    *,
    task_id: str,
    workspace: str,
    db,
    rag,
    answer_ids: list[str],
    scope: GuidanceEnrichmentScope,
    batch_size: int,
    max_suggestions: int,
) -> None:
    from lightrag.api.task_manager import TaskStatus, get_task_service

    service = get_task_service()
    llm_func = getattr(rag, "llm_model_func", None)
    if llm_func is None:
        raise RuntimeError("LLM is not configured for this workspace")

    answer_rows = await db.query(
        """
        SELECT workspace, answer_id, title, body, approved_summary, content_format,
               display_policy, status, version, valid_from, valid_until, priority,
               tags, metadata, publish_time, create_time, update_time
        FROM LIGHTRAG_ANSWER_ITEMS
        WHERE workspace = $1 AND answer_id = ANY($2::text[])
        ORDER BY create_time ASC, answer_id ASC
        """,
        [workspace, answer_ids],
        multirows=True,
    )
    guidance_rows = await db.query(
        """
        SELECT guidance_id, workspace, answer_id, guidance_type, text, weight,
               metadata, create_time
        FROM LIGHTRAG_ANSWER_GUIDANCE
        WHERE workspace = $1 AND answer_id = ANY($2::text[])
        ORDER BY create_time ASC
        """,
        [workspace, answer_ids],
        multirows=True,
    )
    guidance_by_answer: dict[str, list[AnswerGuidance]] = {}
    for row in guidance_rows or []:
        guidance = _guidance_from_row(dict(row))
        guidance_by_answer.setdefault(guidance.answer_id, []).append(guidance)

    answers = [_answer_from_row(dict(row)) for row in answer_rows or []]
    targets = [
        answer
        for answer in answers
        if scope == "all"
        or (
            scope == "coverage"
            and _guidance_needs_coverage_enrichment(
                answer,
                guidance_by_answer.get(answer.answer_id, []),
            )
        )
        or _guidance_needs_llm_enrichment(
            guidance_by_answer.get(answer.answer_id, [])
        )
    ]
    skipped = len(answers) - len(targets)
    if not targets:
        vector_result = await _rebuild_answer_vectors_for_ids(
            db=db,
            workspace=workspace,
            rag=rag,
            answer_ids=answer_ids,
            task_id=task_id,
        )
        await service.complete_task(
            task_id,
            result={
                "answer_count": len(answers),
                "target_count": 0,
                "skipped_count": skipped,
                "processed_count": 0,
                "guidance_created": 0,
                "failed_batches": [],
                "vector_result": vector_result,
            },
        )
        return

    guidance_created = 0
    processed_count = 0
    failed_batches: list[dict[str, Any]] = []
    batch_count = (len(targets) + batch_size - 1) // batch_size
    for batch_index in range(batch_count):
        task = service.get_task(task_id)
        if task and task.status == TaskStatus.CANCELLED:
            return
        batch = targets[
            batch_index * batch_size : (batch_index + 1) * batch_size
        ]
        existing_texts = {
            answer.answer_id: {
                _normalise_text(item.text)
                for item in guidance_by_answer.get(answer.answer_id, [])
            }
            for answer in batch
        }
        prompt_items = [
            {
                "answer_id": answer.answer_id,
                "title": answer.title,
                "body": answer.body[:1200],
                "tags": answer.tags[:10],
                "existing_hints": [
                    {
                        "guidance_type": item.guidance_type,
                        "text": item.text,
                    }
                    for item in guidance_by_answer.get(answer.answer_id, [])[:12]
                ],
            }
            for answer in batch
        ]
        prompt = (
            "Create short search hints for Korean users of each fixed FAQ answer below. "
            "Return only a JSON array. Preserve each provided answer_id exactly. "
            "For each answer return {\"answer_id\": string, \"suggestions\": "
            "[{\"guidance_type\": \"question|keyword|synonym\", "
            "\"text\": string, \"weight\": number}]}. "
            f"Return at most {max_suggestions} useful, non-duplicate suggestions "
            "per answer. Cover natural spoken questions, abbreviations, spacing or spelling "
            "variants, and established Korean-English product aliases. When an English product "
            "name appears, include its common Korean pronunciation or name when useful "
            "(for example Teams and 팀즈). Do not generate answer content, negative terms, "
            "or internal notes.\n\n"
            f"FAQ answers:\n{json.dumps(prompt_items, ensure_ascii=False)}"
        )
        try:
            raw = await llm_func(
                prompt,
                system_prompt=(
                    "You improve FAQ retrieval quality by adding concise representative "
                    "questions, keywords, colloquial expressions, and bilingual aliases for "
                    "Korean search. Return strict JSON only."
                ),
                _llm_purpose="knowledge_structure",
            )
            parsed = _parse_batch_guidance_suggestions(
                str(raw),
                {answer.answer_id for answer in batch},
                existing_texts,
                max_suggestions,
                task_id,
            )
            for answer in batch:
                created = await _insert_answer_guidance(
                    db,
                    workspace,
                    answer.answer_id,
                    parsed.get(answer.answer_id, []),
                )
                guidance_created += len(created)
                guidance_by_answer.setdefault(answer.answer_id, []).extend(created)
            processed_count += len(batch)
        except Exception as exc:
            logger.warning(
                "[Answers] Batch LLM guidance enrichment failed task=%s batch=%s: %s",
                task_id,
                batch_index + 1,
                exc,
            )
            failed_batches.append(
                {
                    "batch": batch_index + 1,
                    "answer_ids": [answer.answer_id for answer in batch],
                    "error": str(exc)[:500],
                }
            )

        progress = ((batch_index + 1) / batch_count) * 95.0
        await service.update_progress(
            task_id,
            progress,
            (
                f"FAQ hint enrichment batch {batch_index + 1}/{batch_count} "
                f"completed ({guidance_created} hints added)"
            ),
            detail={
                "processed_count": processed_count,
                "target_count": len(targets),
                "guidance_created": guidance_created,
                "failed_batch_count": len(failed_batches),
            },
        )

    if failed_batches and processed_count == 0:
        raise RuntimeError(
            f"All {len(failed_batches)} LLM guidance enrichment batches failed"
        )
    vector_result = await _rebuild_answer_vectors_for_ids(
        db=db,
        workspace=workspace,
        rag=rag,
        answer_ids=answer_ids,
        task_id=task_id,
        progress_start=95.0,
        progress_span=4.0,
    )
    await service.complete_task(
        task_id,
        result={
            "answer_count": len(answers),
            "target_count": len(targets),
            "skipped_count": skipped,
            "processed_count": processed_count,
            "guidance_created": guidance_created,
            "failed_batches": failed_batches,
            "vector_result": vector_result,
        },
    )


def _score_candidate(
    answer: AnswerItem,
    guidance: list[AnswerGuidance],
    query: str,
    strategy: ResolveStrategy = "balanced",
    alias_expansions: Optional[list[AnswerAliasExpansion]] = None,
) -> tuple[float, list[str], str, dict[str, float]]:
    query_variants = _alias_query_variants(query, alias_expansions)
    query_token_variants = [_tokens(variant) for variant in query_variants]
    details: dict[str, float] = {}
    matched_guidance: list[str] = []

    title = _normalise_text(answer.title)
    body = _normalise_text(answer.body)
    tags = _normalise_text(" ".join(answer.tags))
    asset_text = _normalise_text(_answer_asset_search_content(answer.assets))

    if any(variant and variant in title for variant in query_variants):
        _add_score(details, "title_exact", 0.42)
    if any(variant and variant in body for variant in query_variants):
        _add_score(details, "body_exact", 0.16)
    if any(variant and variant in tags for variant in query_variants):
        _add_score(details, "tag_exact", 0.2)
    if any(variant and variant in asset_text for variant in query_variants):
        _add_score(details, "asset_exact", 0.18)

    _add_score(
        details,
        "title_overlap",
        max((_partial_overlap(tokens, title) for tokens in query_token_variants), default=0.0)
        * 0.26,
    )
    _add_score(
        details,
        "body_overlap",
        max((_partial_overlap(tokens, body) for tokens in query_token_variants), default=0.0)
        * 0.1,
    )
    _add_score(
        details,
        "tag_overlap",
        max((_partial_overlap(tokens, tags) for tokens in query_token_variants), default=0.0)
        * 0.18,
    )
    _add_score(
        details,
        "asset_overlap",
        max(
            (_partial_overlap(tokens, asset_text) for tokens in query_token_variants),
            default=0.0,
        )
        * 0.12,
    )

    for item in guidance:
        guidance_text = _normalise_text(item.text)
        multiplier = GUIDANCE_TYPE_MULTIPLIER.get(item.guidance_type, 1.0)
        exact_match = bool(
            guidance_text
            and any(
                variant
                and (variant in guidance_text or guidance_text in variant)
                for variant in query_variants
            )
        )
        overlap = max(
            (
                _partial_overlap(tokens, guidance_text)
                for tokens in query_token_variants
            ),
            default=0.0,
        )
        if not exact_match and overlap <= 0:
            continue

        if item.guidance_type == "negative_keyword" and exact_match:
            _add_score(details, "negative_block", -1.0)
            matched_guidance.append(f"!{item.text}")
            return 0.0, matched_guidance[:8], "blocked_by_negative_guidance", details

        if exact_match:
            amount = min(0.42, 0.18 * item.weight) * multiplier
        else:
            amount = min(0.28, overlap * 0.24 * item.weight) * multiplier

        key = f"guidance_{item.guidance_type}"
        _add_score(details, key, amount)
        prefix = "!" if item.guidance_type == "negative_keyword" else ""
        matched_guidance.append(f"{prefix}{item.text}")

    positive_guidance_text = " ".join(
        item.text
        for item in guidance
        if item.guidance_type not in {"negative_keyword", "note"}
    )
    for expansion in alias_expansions or []:
        matched_variant: Optional[str] = None
        for variant in expansion.expanded_terms:
            variant_text = _normalise_text(variant)
            if not variant_text:
                continue
            if (
                variant_text in title
                or variant_text in _normalise_text(positive_guidance_text)
                or variant_text in tags
                or variant_text in body
            ):
                matched_variant = variant
                break
        if matched_variant:
            matched_guidance.append(
                f"{expansion.matched_term} ↔ {matched_variant}"
                if _normalise_text(expansion.matched_term)
                != _normalise_text(matched_variant)
                else matched_variant
            )

    if strategy == "balanced":
        metadata_text = _normalise_text(
            json.dumps(_compact_answer_metadata(answer.metadata), ensure_ascii=False)
        )
        _add_score(
            details,
            "metadata_overlap",
            max(
                (
                    _partial_overlap(tokens, metadata_text)
                    for tokens in query_token_variants
                ),
                default=0.0,
            )
            * 0.08,
        )

    _add_score(details, "priority", min(0.07, max(answer.priority, 0) * 0.01))

    raw_score = sum(details.values())
    score = min(max(raw_score, 0.0), 1.0)
    positive_details = {key: value for key, value in details.items() if value > 0}
    reason = ", ".join(f"{key}:{value:.2f}" for key, value in sorted(positive_details.items()))
    if not reason and details:
        reason = ", ".join(f"{key}:{value:.2f}" for key, value in sorted(details.items()))
    return score, matched_guidance[:8], reason or "weak_match", details


def _candidate_category(candidate: ResolveCandidate) -> Optional[str]:
    metadata = candidate.answer.metadata
    if not isinstance(metadata, dict):
        return None
    for key in ("category", "service_area", "product_family", "domain"):
        value = metadata.get(key)
        if isinstance(value, (str, int, float)) and str(value).strip():
            return f"{key}:{_normalise_text(str(value))}"
    return None


def _candidate_evidence_sources(candidate: ResolveCandidate) -> list[str]:
    details = candidate.score_details
    sources: list[str] = []
    if _safe_float(details.get("keyword_score"), 0.0) > 0:
        sources.append("keyword")
    if _safe_float(details.get("vector_score"), 0.0) > 0:
        sources.append("vector")
    if _safe_float(details.get("graph_score"), 0.0) > 0:
        sources.append("graph")
    return sources


def _precision_abstention_reason(
    selected: Optional[ResolveCandidate],
    candidates: list[ResolveCandidate],
    payload: ResolveRequest,
    config: AnswerGraphConfig,
    llm_selection: dict[str, Any],
) -> tuple[Optional[str], dict[str, Any]]:
    diagnostics: dict[str, Any] = {}
    if selected is None:
        return "no_candidates", diagnostics

    llm_status = str(llm_selection.get("status") or "")
    llm_confidence = _safe_float(llm_selection.get("confidence"), 0.0, 0.0, 1.0)
    confident_llm_selection = (
        payload.retrieval_mode == "llm_rerank"
        and selected.selected_by == "llm_id_selector"
        and llm_status == "llm_ready"
        and llm_selection.get("answer_id") == selected.answer.answer_id
        and llm_confidence >= config.llm_min_confidence
    )
    diagnostics["llm_status"] = llm_status
    diagnostics["llm_confidence"] = llm_confidence
    diagnostics["confident_llm_selection"] = confident_llm_selection

    effective_min_score = max(payload.min_score, config.precision_min_score)
    if confident_llm_selection:
        effective_min_score = payload.min_score
    diagnostics["effective_min_score"] = effective_min_score
    if selected.score < effective_min_score:
        return "low_score", diagnostics

    evidence_sources = _candidate_evidence_sources(selected)
    available_sources = (
        1
        if payload.retrieval_mode == "keyword"
        else (
            3
            if payload.retrieval_mode == "graph_hybrid"
            or (
                payload.retrieval_mode == "llm_rerank"
                and config.enabled
            )
            else 2
        )
    )
    required_sources = min(
        config.min_evidence_sources,
        available_sources,
    )
    diagnostics["evidence_sources"] = evidence_sources
    diagnostics["required_evidence_sources"] = required_sources
    if len(evidence_sources) < required_sources:
        return "insufficient_evidence", diagnostics

    selected_category = _candidate_category(selected)
    category_scores: dict[str, float] = {}
    for candidate in candidates:
        category = _candidate_category(candidate)
        if category:
            category_scores[category] = max(
                category_scores.get(category, 0.0),
                candidate.score,
            )
    diagnostics["selected_category"] = selected_category
    diagnostics["category_scores"] = category_scores
    if selected_category and len(category_scores) > 1:
        ranked_categories = sorted(
            category_scores.items(),
            key=lambda item: item[1],
            reverse=True,
        )
        category_margin = ranked_categories[0][1] - ranked_categories[1][1]
        diagnostics["category_margin"] = round(category_margin, 4)
        diagnostics["category_competition"] = (
            category_margin < config.min_category_margin
        )
        if selected_category != ranked_categories[0][0]:
            return "category_conflict", diagnostics

    same_category_competitors = [
        candidate
        for candidate in candidates
        if candidate.answer.answer_id != selected.answer.answer_id
        and (
            not selected_category
            or _candidate_category(candidate) == selected_category
        )
    ]
    nearest_score = max(
        (candidate.score for candidate in same_category_competitors),
        default=0.0,
    )
    score_margin = selected.score - nearest_score
    diagnostics["score_margin"] = round(score_margin, 4)
    diagnostics["candidate_competition"] = (
        score_margin < config.min_score_margin
    )
    exceptionally_strong_candidate = (
        selected.score >= 0.75
        and score_margin >= max(config.min_score_margin, 0.20)
        and len(evidence_sources) >= required_sources
    )
    diagnostics["exceptionally_strong_candidate"] = exceptionally_strong_candidate

    if payload.retrieval_mode == "llm_rerank":
        if llm_status != "llm_ready":
            return "llm_unavailable", diagnostics
        if llm_selection.get("answer_id") is None:
            if (
                llm_selection.get("decision_reason") == "ambiguous"
                and exceptionally_strong_candidate
            ):
                return None, diagnostics
            if llm_selection.get("decision_reason") == "ambiguous":
                return "ambiguous_intent", diagnostics
            return "llm_no_match", diagnostics
        if (
            selected.selected_by == "llm_id_selector"
            and llm_confidence < config.llm_min_confidence
        ):
            return "low_llm_confidence", diagnostics

    return None, diagnostics


def _clarification_question(
    candidates: list[ResolveCandidate],
    abstention_reason: Optional[str],
    llm_selection: dict[str, Any],
) -> Optional[str]:
    if abstention_reason not in {
        "no_candidates",
        "low_score",
        "ambiguous_intent",
        "ambiguous_category",
        "ambiguous_candidates",
        "category_conflict",
    }:
        return None

    if abstention_reason in {"no_candidates", "low_score"}:
        return "어떤 서비스에서 어떤 문제가 발생했는지 조금 더 알려주세요."

    llm_question = str(llm_selection.get("clarification_question") or "").strip()
    candidate_text = " ".join(candidate.answer.title for candidate in candidates)
    question_matches_candidate_language = (
        not re.search(r"[가-힣]", candidate_text)
        or bool(re.search(r"[가-힣]", llm_question))
    )
    if llm_question and question_matches_candidate_language:
        return llm_question[:300]

    titles: list[str] = []
    for candidate in candidates:
        title = candidate.answer.title.strip()
        if title and title not in titles:
            titles.append(title)
        if len(titles) >= 3:
            break
    if not titles:
        return "어떤 서비스에서 어떤 문제가 발생했는지 조금 더 알려주세요."
    return (
        "문의하신 내용이 여러 경우에 해당할 수 있습니다. "
        f"{', '.join(titles)} 중 어느 내용인지 알려주세요."
    )


async def _keyword_candidate_ids(
    db,
    workspace: str,
    query: str,
    status_filter: list[str],
    allowed_answer_ids: Optional[list[str]],
    alias_expansions: Optional[list[AnswerAliasExpansion]] = None,
) -> tuple[list[str], bool]:
    normalized_query = _normalise_text(query)
    query_terms = _candidate_query_terms(query)
    for expansion in alias_expansions or []:
        for term in expansion.expanded_terms:
            for candidate in _candidate_query_terms(term):
                if candidate not in query_terms:
                    query_terms.append(candidate)
    query_terms = query_terms[:64]
    rows = await db.query(
        """
        WITH guidance_text AS (
            SELECT answer_id, LOWER(STRING_AGG(text, ' ' ORDER BY guidance_id)) AS search_text
            FROM LIGHTRAG_ANSWER_GUIDANCE
            WHERE workspace = $1
            GROUP BY answer_id
        ),
        asset_text AS (
            SELECT
                answer_id,
                LOWER(
                    STRING_AGG(
                        CONCAT_WS(
                            ' ',
                            caption,
                            alt_text,
                            search_text,
                            content_text
                        ),
                        ' ' ORDER BY display_order, asset_id
                    )
                ) AS search_text
            FROM LIGHTRAG_ANSWER_ASSETS
            WHERE workspace = $1 AND is_active = TRUE
            GROUP BY answer_id
        ),
        eligible AS (
            SELECT
                answers.answer_id,
                answers.priority,
                answers.update_time,
                LOWER(
                    CONCAT_WS(
                        ' ',
                        answers.answer_id,
                        answers.title,
                        answers.body,
                        answers.tags::text,
                        (
                            answers.metadata
                            - 'source_profile'
                            - 'raw_content'
                        )::text
                    )
                ) AS answer_text,
                COALESCE(guidance_text.search_text, '') AS guidance_text,
                COALESCE(asset_text.search_text, '') AS asset_text
            FROM LIGHTRAG_ANSWER_ITEMS AS answers
            LEFT JOIN guidance_text ON guidance_text.answer_id = answers.answer_id
            LEFT JOIN asset_text ON asset_text.answer_id = answers.answer_id
            WHERE answers.workspace = $1
              AND answers.status = ANY($2::text[])
              AND (answers.valid_from IS NULL OR answers.valid_from <= NOW())
              AND (answers.valid_until IS NULL OR answers.valid_until >= NOW())
              AND ($3::text[] IS NULL OR answers.answer_id = ANY($3::text[]))
        ),
        ranked AS (
            SELECT
                answer_id,
                priority,
                update_time,
                (
                    CASE
                        WHEN $4 <> '' AND STRPOS(answer_text, $4) > 0 THEN 30
                        ELSE 0
                    END
                    + CASE
                        WHEN $4 <> '' AND STRPOS(guidance_text, $4) > 0 THEN 40
                        ELSE 0
                    END
                    + CASE
                        WHEN $4 <> '' AND STRPOS(asset_text, $4) > 0 THEN 24
                        ELSE 0
                    END
                    + (
                        SELECT COUNT(*) * 3
                        FROM UNNEST($5::text[]) AS term
                        WHERE STRPOS(answer_text, term) > 0
                    )
                    + (
                        SELECT COUNT(*) * 4
                        FROM UNNEST($5::text[]) AS term
                        WHERE STRPOS(guidance_text, term) > 0
                    )
                    + (
                        SELECT COUNT(*) * 3
                        FROM UNNEST($5::text[]) AS term
                        WHERE STRPOS(asset_text, term) > 0
                    )
                ) AS prefilter_score
            FROM eligible
        )
        SELECT answer_id
        FROM ranked
        WHERE prefilter_score > 0
        ORDER BY prefilter_score DESC, priority DESC, update_time DESC, answer_id ASC
        LIMIT $6
        """,
        [
            workspace,
            status_filter,
            allowed_answer_ids,
            normalized_query,
            query_terms,
            MAX_KEYWORD_SEARCH_CANDIDATES + 1,
        ],
        multirows=True,
    )
    candidate_ids = [str(row["answer_id"]) for row in rows or []]
    truncated = len(candidate_ids) > MAX_KEYWORD_SEARCH_CANDIDATES
    return candidate_ids[:MAX_KEYWORD_SEARCH_CANDIDATES], truncated


def _render_answer_response(answer: AnswerItem, display_policy: Optional[DisplayPolicy] = None) -> tuple[str, DisplayPolicy]:
    policy = display_policy or answer.display_policy
    summary = (answer.approved_summary or "").strip()
    body = answer.body.strip()
    if policy == "summary":
        return summary or body, policy
    if policy == "full":
        return body, policy
    if summary and body and summary != body:
        return f"{summary}\n\n{body}", policy
    return body or summary, policy


def create_answer_routes(rag, api_key: Optional[str] = None):
    combined_auth = get_combined_auth_dependency(api_key)

    async def db_for_request(request: Request):
        workspace = _get_workspace_from_request(request)
        base_db = _db_from_rag(rag)
        if base_db is None:
            raise HTTPException(status_code=500, detail="Database connection not available")
        await _ensure_tables(base_db)

        workspace_mode = await _workspace_mode_from_db(base_db, workspace)
        if workspace_mode not in ANSWER_WORKSPACE_MODES:
            raise HTTPException(
                status_code=409,
                detail=(
                    "Answer catalog APIs require an answer_catalog or hybrid workspace. "
                    f"Current workspace mode is '{workspace_mode}'."
                ),
            )

        workspace_rag = await get_workspace_rag(workspace)
        if workspace_rag is None:
            workspace_rag = rag
            logger.warning("[Answers] Using default RAG instance for workspace: %s", workspace)
        db = _db_from_rag(workspace_rag)
        if db is None:
            raise HTTPException(status_code=500, detail="Database connection not available")
        await _ensure_tables(db)
        return workspace, db

    async def rag_for_workspace(workspace: str):
        workspace_rag = await get_workspace_rag(workspace)
        if workspace_rag is None:
            workspace_rag = rag
        return workspace_rag

    @router.get(
        "/graph/config",
        response_model=AnswerGraphConfig,
        dependencies=[Depends(combined_auth)],
        summary="Get FAQ graph configuration",
        description=(
            "Returns workspace-scoped graph settings. Reading this endpoint does not "
            "build or modify graph data."
        ),
    )
    async def get_answer_graph_config(request: Request):
        workspace, db = await db_for_request(request)
        return await _get_answer_graph_config(db, workspace)

    @router.put(
        "/graph/config",
        response_model=AnswerGraphConfig,
        dependencies=[Depends(combined_auth)],
        summary="Update FAQ graph configuration",
        description=(
            "Enables optional graph_hybrid retrieval and updates its bounded schema. "
            "Changing entity types, relation types, or the extraction prompt increments "
            "the schema version and marks existing projections stale; call "
            "POST /api/answers/graph/rebuild afterward."
        ),
    )
    async def update_answer_graph_config(
        request: Request,
        payload: AnswerGraphConfigUpdate,
    ):
        workspace, db = await db_for_request(request)
        current = await _get_answer_graph_config(db, workspace)
        updates = payload.model_dump(exclude_unset=True)
        entity_types = updates.get("entity_types", current.entity_types)
        relation_types = updates.get("relation_types", current.relation_types)
        if "FAQAnswer" not in entity_types:
            entity_types = ["FAQAnswer", *entity_types]
        entity_types = list(dict.fromkeys(_graph_clean_label(item, 80) for item in entity_types if _graph_clean_label(item, 80)))
        relation_types = list(
            dict.fromkeys(
                _graph_clean_label(item, 80)
                for item in relation_types
                if _graph_clean_label(item, 80)
            )
        )
        schema_changed = (
            "entity_types" in updates and entity_types != current.entity_types
        ) or (
            "relation_types" in updates
            and relation_types != current.relation_types
        ) or (
            "extraction_prompt" in updates
            and updates["extraction_prompt"] != current.extraction_prompt
        ) or (
            "ai_extraction_strategy" in updates
            and updates["ai_extraction_strategy"] != current.ai_extraction_strategy
        ) or (
            "ai_retry_max_tokens" in updates
            and updates["ai_retry_max_tokens"] != current.ai_retry_max_tokens
        ) or (
            "ai_min_relations" in updates
            and updates["ai_min_relations"] != current.ai_min_relations
        ) or (
            "ai_min_relation_types" in updates
            and updates["ai_min_relation_types"] != current.ai_min_relation_types
        )
        next_schema_version = current.schema_version + (1 if schema_changed else 0)
        row = await db.query(
            """
            INSERT INTO LIGHTRAG_ANSWER_GRAPH_CONFIG
                (
                    workspace, enabled, auto_sync, graph_weight, min_similarity, max_hops,
                    precision_mode, precision_min_score, min_score_margin,
                    min_category_margin, min_evidence_sources, llm_min_confidence,
                    entity_types, relation_types, extraction_prompt,
                    ai_extraction_strategy, ai_retry_max_tokens, ai_min_relations,
                    ai_min_relation_types, schema_version
                )
            VALUES (
                $1, $2, $3, $4, $5, $6, $7, $8, $9, $10, $11, $12,
                $13::jsonb, $14::jsonb, $15, $16, $17, $18, $19, $20
            )
            ON CONFLICT (workspace)
            DO UPDATE SET
                enabled = EXCLUDED.enabled,
                auto_sync = EXCLUDED.auto_sync,
                graph_weight = EXCLUDED.graph_weight,
                min_similarity = EXCLUDED.min_similarity,
                max_hops = EXCLUDED.max_hops,
                precision_mode = EXCLUDED.precision_mode,
                precision_min_score = EXCLUDED.precision_min_score,
                min_score_margin = EXCLUDED.min_score_margin,
                min_category_margin = EXCLUDED.min_category_margin,
                min_evidence_sources = EXCLUDED.min_evidence_sources,
                llm_min_confidence = EXCLUDED.llm_min_confidence,
                entity_types = EXCLUDED.entity_types,
                relation_types = EXCLUDED.relation_types,
                extraction_prompt = EXCLUDED.extraction_prompt,
                ai_extraction_strategy = EXCLUDED.ai_extraction_strategy,
                ai_retry_max_tokens = EXCLUDED.ai_retry_max_tokens,
                ai_min_relations = EXCLUDED.ai_min_relations,
                ai_min_relation_types = EXCLUDED.ai_min_relation_types,
                schema_version = EXCLUDED.schema_version,
                update_time = NOW()
            RETURNING workspace, enabled, auto_sync, graph_weight, min_similarity, max_hops,
                      precision_mode, precision_min_score, min_score_margin,
                      min_category_margin, min_evidence_sources, llm_min_confidence,
                      entity_types, relation_types, extraction_prompt,
                      ai_extraction_strategy, ai_retry_max_tokens, ai_min_relations,
                      ai_min_relation_types, schema_version,
                      create_time, update_time
            """,
            [
                workspace,
                updates.get("enabled", current.enabled),
                updates.get("auto_sync", current.auto_sync),
                updates.get("graph_weight", current.graph_weight),
                updates.get("min_similarity", current.min_similarity),
                updates.get("max_hops", current.max_hops),
                updates.get("precision_mode", current.precision_mode),
                updates.get("precision_min_score", current.precision_min_score),
                updates.get("min_score_margin", current.min_score_margin),
                updates.get("min_category_margin", current.min_category_margin),
                updates.get("min_evidence_sources", current.min_evidence_sources),
                updates.get("llm_min_confidence", current.llm_min_confidence),
                _json_list(entity_types),
                _json_list(relation_types),
                updates.get("extraction_prompt", current.extraction_prompt),
                updates.get(
                    "ai_extraction_strategy", current.ai_extraction_strategy
                ),
                updates.get("ai_retry_max_tokens", current.ai_retry_max_tokens),
                updates.get("ai_min_relations", current.ai_min_relations),
                updates.get(
                    "ai_min_relation_types", current.ai_min_relation_types
                ),
                next_schema_version,
            ],
        )
        if schema_changed:
            await db.execute(
                """
                UPDATE LIGHTRAG_ANSWER_GRAPH_PROJECTIONS
                SET status = 'stale', schema_version = $2, update_time = NOW()
                WHERE workspace = $1
                """,
                {"workspace": workspace, "schema_version": next_schema_version},
            )
        return _graph_config_from_row(workspace, dict(row))

    @router.get(
        "/graph/status",
        response_model=AnswerGraphStatusResponse,
        dependencies=[Depends(combined_auth)],
        summary="Get FAQ graph build status",
        description=(
            "Counts ready, stale, pending, failed, and missing FAQ projections for the "
            "current workspace. It is safe to poll this endpoint."
        ),
    )
    async def get_answer_graph_status(request: Request):
        workspace, db = await db_for_request(request)
        config = await _get_answer_graph_config(db, workspace)
        row = await db.query(
            """
            SELECT
                COUNT(answers.answer_id)::INT AS total_answers,
                COUNT(*) FILTER (
                    WHERE projections.status = 'ready'
                      AND projections.answer_version = answers.version
                      AND projections.schema_version = $2
                )::INT AS ready,
                COUNT(*) FILTER (
                    WHERE projections.status = 'stale'
                       OR (
                            projections.status = 'ready'
                            AND (
                                projections.answer_version <> answers.version
                                OR projections.schema_version <> $2
                            )
                       )
                )::INT AS stale,
                COUNT(*) FILTER (WHERE projections.status = 'pending')::INT AS pending,
                COUNT(*) FILTER (WHERE projections.status = 'failed')::INT AS failed,
                COUNT(*) FILTER (WHERE projections.answer_id IS NULL)::INT AS missing,
                MAX(projections.built_at) AS last_built_at
            FROM LIGHTRAG_ANSWER_ITEMS AS answers
            LEFT JOIN LIGHTRAG_ANSWER_GRAPH_PROJECTIONS AS projections
              ON projections.workspace = answers.workspace
             AND projections.answer_id = answers.answer_id
            WHERE answers.workspace = $1
              AND answers.status IN ('draft', 'published')
            """,
            [workspace, config.schema_version],
        )
        return AnswerGraphStatusResponse(
            workspace=workspace,
            enabled=config.enabled,
            total_answers=int((row or {}).get("total_answers") or 0),
            ready=int((row or {}).get("ready") or 0),
            stale=int((row or {}).get("stale") or 0),
            pending=int((row or {}).get("pending") or 0),
            failed=int((row or {}).get("failed") or 0),
            missing=int((row or {}).get("missing") or 0),
            last_built_at=_iso((row or {}).get("last_built_at")),
        )

    @router.post(
        "/graph/preview",
        response_model=AnswerGraphProjection,
        dependencies=[Depends(combined_auth)],
        summary="Preview an FAQ graph projection",
        description=(
            "Builds a read-only projection for one FAQ and returns the nodes and "
            "relations that would be stored. The preview never writes graph or "
            "projection data. Set use_llm=true only when the workspace LLM is available."
        ),
    )
    async def preview_answer_graph(
        request: Request,
        payload: AnswerGraphPreviewRequest,
    ):
        workspace, db = await db_for_request(request)
        answers, guidance_by_answer = await _answer_rows_and_guidance(
            db, workspace, [payload.answer_id]
        )
        if not answers:
            raise HTTPException(
                status_code=404,
                detail=f"Answer '{payload.answer_id}' not found",
            )
        answer = answers[0]
        config = await _get_answer_graph_config(db, workspace)
        workspace_rag = await rag_for_workspace(workspace)
        if payload.use_llm and getattr(workspace_rag, "llm_model_func", None) is None:
            raise HTTPException(
                status_code=409,
                detail="LLM is not configured for this workspace.",
            )
        return await _build_answer_graph_preview(
            db,
            workspace,
            workspace_rag,
            answer,
            guidance_by_answer.get(payload.answer_id, []),
            config,
            use_llm=payload.use_llm,
        )

    @router.get(
        "/{answer_id}/graph",
        response_model=AnswerGraphProjection,
        dependencies=[Depends(combined_auth)],
        summary="Get a stored FAQ graph projection",
        description=(
            "Returns the last stored projection for one FAQ. A 404 means the FAQ has "
            "not been rebuilt yet; inspect GET /api/answers/graph/status or start a rebuild."
        ),
    )
    async def get_answer_graph_projection(request: Request, answer_id: str):
        workspace, db = await db_for_request(request)
        await _get_answer_row(db, workspace, answer_id)
        row = await db.query(
            """
            SELECT workspace, answer_id, answer_version, schema_version, status,
                   content_hash, nodes, relations, error, built_at, update_time
            FROM LIGHTRAG_ANSWER_GRAPH_PROJECTIONS
            WHERE workspace = $1 AND answer_id = $2
            """,
            [workspace, answer_id],
        )
        if not row:
            raise HTTPException(
                status_code=404,
                detail=f"FAQ graph projection for '{answer_id}' was not built.",
            )
        return _graph_projection_from_row(dict(row))

    @router.post(
        "/graph/rebuild",
        response_model=AnswerGraphRebuildResponse,
        dependencies=[Depends(combined_auth)],
        summary="Rebuild FAQ graph projections",
        description=(
            "Starts an asynchronous rebuild and returns a task ID immediately. By "
            "default, only published FAQs with missing or stale projections are selected. "
            "Track progress with GET /api/tasks/{task_id} or the returned stream_url. "
            "Existing keyword and vector retrieval remain available during the rebuild."
        ),
    )
    async def rebuild_answer_graph(
        request: Request,
        payload: AnswerGraphRebuildRequest,
    ):
        from lightrag.api.task_manager import TaskStatus, TaskType, get_task_service

        workspace, db = await db_for_request(request)
        config = await _get_answer_graph_config(db, workspace)
        if not config.enabled:
            raise HTTPException(
                status_code=409,
                detail="Enable the FAQ graph for this workspace before rebuilding it.",
            )
        service = get_task_service()
        active_tasks = [
            task
            for task in service.get_tasks_by_workspace(workspace)
            if task.task_type == TaskType.FAQ_GRAPH_REBUILD
            and task.status in (TaskStatus.PENDING, TaskStatus.RUNNING)
        ]
        if active_tasks:
            active_task = max(active_tasks, key=lambda task: task.updated_at)
            return AnswerGraphRebuildResponse(
                task_id=active_task.task_id,
                stream_url=f"/api/tasks/{active_task.task_id}/stream",
                answer_count=int(active_task.metadata.get("answer_count", 0)),
                message="An FAQ graph rebuild is already running.",
                reused=True,
            )
        workspace_rag = await rag_for_workspace(workspace)
        if payload.use_llm and getattr(workspace_rag, "llm_model_func", None) is None:
            raise HTTPException(
                status_code=409,
                detail="LLM is not configured for this workspace.",
            )
        status_filter = (
            ["draft", "published"] if payload.include_drafts else ["published"]
        )
        requested_ids = [
            str(answer_id)
            for answer_id in (payload.answer_ids or [])
            if str(answer_id).strip()
        ]
        rows = await db.query(
            """
            SELECT answers.answer_id
            FROM LIGHTRAG_ANSWER_ITEMS AS answers
            LEFT JOIN LIGHTRAG_ANSWER_GRAPH_PROJECTIONS AS projections
              ON projections.workspace = answers.workspace
             AND projections.answer_id = answers.answer_id
            WHERE answers.workspace = $1
              AND answers.status = ANY($2::text[])
              AND ($3::text[] IS NULL OR answers.answer_id = ANY($3::text[]))
              AND (
                    $4::boolean = FALSE
                    OR projections.answer_id IS NULL
                    OR projections.status <> 'ready'
                    OR projections.answer_version <> answers.version
                    OR projections.schema_version <> $5
              )
            ORDER BY answers.priority DESC, answers.update_time DESC
            LIMIT $6
            """,
            [
                workspace,
                status_filter,
                requested_ids or None,
                payload.only_stale,
                config.schema_version,
                payload.limit,
            ],
            multirows=True,
        )
        answer_ids = [str(row["answer_id"]) for row in rows or []]
        task = service.create_task(
            task_type=TaskType.FAQ_GRAPH_REBUILD,
            workspace=workspace,
            metadata={
                "answer_count": len(answer_ids),
                "use_llm": payload.use_llm,
                "schema_version": config.schema_version,
                "trigger_source": "manual_api",
            },
        )
        service.run_in_background(
            task.task_id,
            _rebuild_answer_graph_background,
            task_id=task.task_id,
            workspace=workspace,
            db=db,
            rag=workspace_rag,
            answer_ids=answer_ids,
            use_llm=payload.use_llm,
        )
        return AnswerGraphRebuildResponse(
            task_id=task.task_id,
            stream_url=f"/api/tasks/{task.task_id}/stream",
            answer_count=len(answer_ids),
            message="FAQ graph rebuild started.",
        )

    @router.get("", response_model=AnswerListResponse, dependencies=[Depends(combined_auth)])
    async def list_answers(
        request: Request,
        status: Optional[str] = Query(default=None),
        search: Optional[str] = Query(default=None),
        content_format: Optional[str] = Query(default=None),
        display_policy: Optional[str] = Query(default=None),
        validity: Optional[str] = Query(default=None),
        tag: Optional[str] = Query(default=None),
        source_type: Optional[str] = Query(default=None),
        has_guidance: Optional[bool] = Query(default=None),
        has_source: Optional[bool] = Query(default=None),
        min_priority: Optional[int] = Query(default=None),
        page: int = Query(default=1, ge=1),
        page_size: int = Query(default=20, ge=1, le=100),
    ):
        workspace, db = await db_for_request(request)
        params: list[Any] = [workspace]
        where = ["a.workspace = $1"]
        if status and status != "all":
            params.append(status)
            where.append(f"a.status = ${len(params)}")
        if search:
            params.append(f"%{search}%")
            p = f"${len(params)}"
            where.append(
                f"(a.title ILIKE {p} OR a.body ILIKE {p} OR a.approved_summary ILIKE {p} "
                f"OR a.answer_id ILIKE {p})"
            )
        if content_format and content_format != "all":
            params.append(content_format)
            where.append(f"a.content_format = ${len(params)}")
        if display_policy and display_policy != "all":
            params.append(display_policy)
            where.append(f"a.display_policy = ${len(params)}")
        if validity and validity != "all":
            if validity == "active":
                where.append(
                    "(a.status != 'expired' AND (a.valid_from IS NULL OR a.valid_from <= NOW()) "
                    "AND (a.valid_until IS NULL OR a.valid_until >= NOW()))"
                )
            elif validity == "scheduled":
                where.append("a.valid_from IS NOT NULL AND a.valid_from > NOW()")
            elif validity == "expired":
                where.append("(a.status = 'expired' OR (a.valid_until IS NOT NULL AND a.valid_until < NOW()))")
            elif validity == "no_period":
                where.append("a.valid_from IS NULL AND a.valid_until IS NULL")
        if tag:
            params.append(f"%{tag}%")
            where.append(
                "EXISTS ("
                "SELECT 1 FROM jsonb_array_elements_text(a.tags) AS tag_value "
                f"WHERE tag_value ILIKE ${len(params)}"
                ")"
            )
        if source_type and source_type != "all":
            params.append(source_type)
            where.append(
                "EXISTS ("
                "SELECT 1 FROM LIGHTRAG_ANSWER_SOURCE_LINKS l "
                "JOIN LIGHTRAG_ANSWER_SOURCE_SNAPSHOTS s "
                "ON s.workspace = l.workspace AND s.snapshot_id = l.snapshot_id "
                "WHERE l.workspace = a.workspace AND l.answer_id = a.answer_id "
                f"AND s.source_type = ${len(params)}"
                ")"
            )
        if has_guidance is not None:
            guidance_exists = (
                "EXISTS (SELECT 1 FROM LIGHTRAG_ANSWER_GUIDANCE g "
                "WHERE g.workspace = a.workspace AND g.answer_id = a.answer_id)"
            )
            where.append(guidance_exists if has_guidance else f"NOT {guidance_exists}")
        if has_source is not None:
            source_exists = (
                "EXISTS (SELECT 1 FROM LIGHTRAG_ANSWER_SOURCE_LINKS l "
                "WHERE l.workspace = a.workspace AND l.answer_id = a.answer_id)"
            )
            where.append(source_exists if has_source else f"NOT {source_exists}")
        if min_priority is not None:
            params.append(min_priority)
            where.append(f"a.priority >= ${len(params)}")
        where_sql = " AND ".join(where)
        count = await db.query(
            f"SELECT COUNT(*)::INT AS total FROM LIGHTRAG_ANSWER_ITEMS a WHERE {where_sql}",
            params,
        )
        total = int((count or {}).get("total") or 0)

        page_params = [*params, page_size, (page - 1) * page_size]
        rows = await db.query(
            f"""
            SELECT a.workspace, a.answer_id, a.title, a.body, a.approved_summary, a.content_format,
                   a.display_policy, a.status, a.version, a.valid_from, a.valid_until, a.priority,
                   a.tags, a.metadata, a.publish_time, a.create_time, a.update_time
            FROM LIGHTRAG_ANSWER_ITEMS a
            WHERE {where_sql}
            ORDER BY a.update_time DESC, a.answer_id ASC
            LIMIT ${len(page_params) - 1} OFFSET ${len(page_params)}
            """,
            page_params,
            multirows=True,
        )
        return AnswerListResponse(
            answers=[_answer_from_row(dict(row)) for row in rows or []],
            total=total,
            page=page,
            page_size=page_size,
        )

    @router.post("", response_model=AnswerItem, dependencies=[Depends(combined_auth)])
    async def create_answer(request: Request, payload: AnswerCreateRequest):
        workspace, db = await db_for_request(request)
        answer_id = payload.answer_id or f"ANS-{uuid.uuid4().hex[:12]}"
        try:
            row = await db.query(
                """
                INSERT INTO LIGHTRAG_ANSWER_ITEMS
                    (workspace, answer_id, title, body, approved_summary, content_format,
                     display_policy, status, version, valid_from, valid_until, priority,
                     tags, metadata, publish_time)
                VALUES
                    ($1, $2, $3, $4, $5, $6, $7, $8, 1, $9, $10, $11,
                     $12::jsonb, $13::jsonb, CASE WHEN $8 = 'published' THEN NOW() ELSE NULL END)
                RETURNING workspace, answer_id, title, body, approved_summary, content_format,
                          display_policy, status, version, valid_from, valid_until, priority,
                          tags, metadata, publish_time, create_time, update_time
                """,
                [
                    workspace,
                    answer_id,
                    payload.title,
                    payload.body,
                    payload.approved_summary,
                    payload.content_format,
                    payload.display_policy,
                    payload.status,
                    payload.valid_from,
                    payload.valid_until,
                    payload.priority,
                    _json_list(payload.tags),
                    _json(payload.metadata),
                ],
            )
            if not row:
                raise HTTPException(status_code=500, detail="Failed to create answer")
            await _record_revision(db, dict(row))
            for text in payload.guidance:
                if not text.strip():
                    continue
                await db.execute(
                    """
                    INSERT INTO LIGHTRAG_ANSWER_GUIDANCE
                        (guidance_id, workspace, answer_id, guidance_type, text, weight, metadata)
                    VALUES ($1, $2, $3, 'keyword', $4, 1.0, '{}'::jsonb)
                    """,
                    {
                        "guidance_id": f"agd-{uuid.uuid4().hex}",
                        "workspace": workspace,
                        "answer_id": answer_id,
                        "text": text.strip(),
                    },
                )
            asyncio.create_task(
                _refresh_answer_vector_safely(
                    db=db,
                    workspace=workspace,
                    answer_id=answer_id,
                )
            )
            await _mark_answer_graph_stale(db, workspace, answer_id, 1)
            asyncio.create_task(
                _refresh_answer_graph_safely(
                    db=db,
                    workspace=workspace,
                    answer_id=answer_id,
                    trigger_source="answer_create",
                )
            )
            return _answer_from_row(dict(row))
        except HTTPException:
            raise
        except Exception as e:
            message = str(e)
            if "duplicate" in message.lower() or "unique" in message.lower():
                raise HTTPException(status_code=409, detail=f"Answer '{answer_id}' already exists")
            logger.error("Failed to create answer: %s\n%s", e, traceback.format_exc())
            raise HTTPException(status_code=500, detail=message)

    @router.post(
        "/source-draft",
        response_model=AnswerSourceDraftResponse,
        dependencies=[Depends(combined_auth)],
    )
    async def create_answer_source_draft(request: Request, payload: AnswerSourceDraftRequest):
        workspace, db = await db_for_request(request)
        answer_id = payload.answer_id or f"ANS-{uuid.uuid4().hex[:12]}"
        snapshot_metadata = {
            **payload.metadata,
            "created_from": payload.metadata.get("created_from") or "answer_source_wizard",
            "content_format": payload.content_format,
            "tags": payload.tags,
        }
        snapshot = await _record_source_snapshot(
            db,
            workspace,
            source_type=payload.source_type,
            source_uri=payload.source_uri,
            file_name=payload.file_name,
            title=payload.title,
            raw_content=payload.body,
            profile=payload.source_profile,
            metadata=snapshot_metadata,
            created_answer_ids=[answer_id],
            status="materialized",
            task_id=payload.metadata.get("task_id"),
        )
        source_metadata = {
            **payload.metadata,
            "created_from": payload.metadata.get("created_from") or "answer_source_wizard",
            "source_type": payload.source_type,
            "source_uri": payload.source_uri,
            "file_name": payload.file_name,
            "source_profile": payload.source_profile,
            "source_snapshot_id": snapshot.snapshot_id,
            "source_content_hash": snapshot.content_hash,
        }
        try:
            row = await db.query(
                """
                INSERT INTO LIGHTRAG_ANSWER_ITEMS
                    (workspace, answer_id, title, body, approved_summary, content_format,
                     display_policy, status, version, valid_from, valid_until, priority,
                     tags, metadata, publish_time)
                VALUES
                    ($1, $2, $3, $4, $5, $6, $7, $8, 1, NULL, NULL, $9,
                     $10::jsonb, $11::jsonb, CASE WHEN $8 = 'published' THEN NOW() ELSE NULL END)
                RETURNING workspace, answer_id, title, body, approved_summary, content_format,
                          display_policy, status, version, valid_from, valid_until, priority,
                          tags, metadata, publish_time, create_time, update_time
                """,
                [
                    workspace,
                    answer_id,
                    payload.title,
                    payload.body,
                    payload.approved_summary,
                    payload.content_format,
                    payload.display_policy,
                    payload.status,
                    payload.priority,
                    _json_list(payload.tags),
                    _json(source_metadata),
                ],
            )
            if not row:
                raise HTTPException(status_code=500, detail="Failed to create answer draft")
            await _record_revision(db, dict(row))
            source_link = await _record_source_link(
                db,
                workspace,
                answer_id=answer_id,
                answer_version=int(row.get("version") or 1),
                snapshot_id=snapshot.snapshot_id,
                link_type="created_from",
                metadata={
                    "created_from": "answer_source_wizard",
                    "source_type": payload.source_type,
                    "source_uri": payload.source_uri,
                    "file_name": payload.file_name,
                },
            )

            created_guidance: list[AnswerGuidance] = []
            for item in payload.guidance:
                if not item.text.strip():
                    continue
                guidance_metadata = {
                    **item.metadata,
                    "created_from": "answer_source_wizard",
                    "source": item.source,
                }
                guidance_row = await db.query(
                    """
                    INSERT INTO LIGHTRAG_ANSWER_GUIDANCE
                        (guidance_id, workspace, answer_id, guidance_type, text, weight, metadata)
                    VALUES ($1, $2, $3, $4, $5, $6, $7::jsonb)
                    RETURNING guidance_id, workspace, answer_id, guidance_type, text, weight, metadata, create_time
                    """,
                    [
                        f"agd-{uuid.uuid4().hex}",
                        workspace,
                        answer_id,
                        item.guidance_type,
                        item.text.strip(),
                        item.weight,
                        _json(guidance_metadata),
                    ],
                )
                if guidance_row:
                    created_guidance.append(_guidance_from_row(dict(guidance_row)))

            asyncio.create_task(
                _refresh_answer_vector_safely(
                    db=db,
                    workspace=workspace,
                    answer_id=answer_id,
                )
            )
            return AnswerSourceDraftResponse(
                answer=_answer_from_row(dict(row)),
                guidance=created_guidance,
                snapshot=snapshot,
                source_link=source_link,
            )
        except HTTPException:
            raise
        except Exception as e:
            try:
                await db.execute(
                    "DELETE FROM LIGHTRAG_ANSWER_SOURCE_SNAPSHOTS WHERE workspace = $1 AND snapshot_id = $2",
                    {"workspace": workspace, "snapshot_id": snapshot.snapshot_id},
                )
            except Exception:
                pass
            message = str(e)
            if "duplicate" in message.lower() or "unique" in message.lower():
                raise HTTPException(status_code=409, detail=f"Answer '{answer_id}' already exists")
            logger.error("Failed to create answer source draft: %s\n%s", e, traceback.format_exc())
            raise HTTPException(status_code=500, detail=message)

    @router.get("/events", response_model=list[AnswerEvent], dependencies=[Depends(combined_auth)])
    async def list_events(
        request: Request,
        event_type: Optional[str] = Query(default=None),
        selected_answer_id: Optional[str] = Query(default=None),
        limit: int = Query(default=50, ge=1, le=1000),
    ):
        workspace, db = await db_for_request(request)
        params: list[Any] = [workspace]
        where = ["workspace = $1"]
        if event_type and event_type != "all":
            params.append(event_type)
            where.append(f"event_type = ${len(params)}")
        if selected_answer_id:
            params.append(selected_answer_id)
            where.append(f"selected_answer_id = ${len(params)}")
        params.append(limit)
        rows = await db.query(
            f"""
            SELECT event_id, workspace, event_type, query, selected_answer_id,
                   candidate_ids, scores, metadata, create_time
            FROM LIGHTRAG_ANSWER_EVENTS
            WHERE {' AND '.join(where)}
            ORDER BY create_time DESC
            LIMIT ${len(params)}
            """,
            params,
            multirows=True,
        )
        return [_event_from_row(dict(row)) for row in rows or []]

    @router.get("/events/page", response_model=AnswerEventListResponse, dependencies=[Depends(combined_auth)])
    async def list_events_page(
        request: Request,
        event_type: Optional[str] = Query(default=None),
        selected_answer_id: Optional[str] = Query(default=None),
        query: Optional[str] = Query(default=None),
        match_status: Optional[Literal["all", "matched", "no_match"]] = Query(default=None),
        source: Optional[str] = Query(default=None),
        mode: Optional[str] = Query(default=None),
        date: Optional[str] = Query(default=None),
        hour: Optional[str] = Query(default=None),
        date_from: Optional[datetime] = Query(default=None),
        date_to: Optional[datetime] = Query(default=None),
        min_confidence: Optional[float] = Query(default=None, ge=0.0, le=1.0),
        max_latency_ms: Optional[int] = Query(default=None, ge=0),
        search: Optional[str] = Query(default=None),
        timezone: str = Query(default="Asia/Seoul", max_length=64),
        page: int = Query(default=1, ge=1),
        page_size: int = Query(default=25, ge=5, le=100),
    ):
        workspace, db = await db_for_request(request)
        timezone_name = _safe_timezone(timezone)
        params: list[Any] = [workspace]
        where = ["e.workspace = $1"]
        _add_answer_event_filters(
            params,
            where,
            event_type=event_type,
            selected_answer_id=selected_answer_id,
            query_text=query,
            match_status=None if match_status == "all" else match_status,
            source=source,
            mode=mode,
            date_key=date,
            hour_key=hour,
            date_from=date_from,
            date_to=date_to,
            min_confidence=min_confidence,
            max_latency_ms=max_latency_ms,
            search=search.strip() if search else None,
            timezone_name=timezone_name,
        )
        where_sql = " AND ".join(where)
        from_sql = """
            FROM LIGHTRAG_ANSWER_EVENTS e
            LEFT JOIN LIGHTRAG_ANSWER_ITEMS a
              ON a.workspace = e.workspace AND a.answer_id = e.selected_answer_id
        """
        total_row = await db.query(
            f"SELECT COUNT(*)::INT AS total {from_sql} WHERE {where_sql}",
            params,
        )
        total = int((total_row or {}).get("total") or 0)
        offset = (page - 1) * page_size
        rows_params = [*params, page_size, offset]
        rows = await db.query(
            f"""
            SELECT e.event_id, e.workspace, e.event_type, e.query, e.selected_answer_id,
                   e.candidate_ids, e.scores, e.metadata, e.create_time
            {from_sql}
            WHERE {where_sql}
            ORDER BY e.create_time DESC
            LIMIT ${len(params) + 1}
            OFFSET ${len(params) + 2}
            """,
            rows_params,
            multirows=True,
        )
        events = [_event_from_row(dict(row)) for row in rows or []]
        answer_ids = {
            answer_id
            for event in events
            for answer_id in [event.selected_answer_id, *event.candidate_ids]
            if answer_id
        }
        answers: list[AnswerItem] = []
        if answer_ids:
            answer_rows = await db.query(
                """
                SELECT workspace, answer_id, title, body, approved_summary, content_format,
                       display_policy, status, version, valid_from, valid_until, priority,
                       tags, metadata, publish_time, create_time, update_time
                FROM LIGHTRAG_ANSWER_ITEMS
                WHERE workspace = $1 AND answer_id = ANY($2::text[])
                """,
                [workspace, list(answer_ids)],
                multirows=True,
            )
            answers = [_answer_from_row(dict(row)) for row in answer_rows or []]
        return AnswerEventListResponse(
            events=events,
            answers=answers,
            total=total,
            page=page,
            page_size=page_size,
        )

    @router.post(
        "/structured/profile",
        response_model=StructuredProfileResponse,
        dependencies=[Depends(combined_auth)],
    )
    async def profile_structured_source(request: Request, payload: StructuredProfileRequest):
        await db_for_request(request)
        return _profile_structured_source(
            payload.source_type,
            payload.raw_content,
            sample_limit=payload.sample_limit,
        )

    @router.post(
        "/structured/excel/preview",
        response_model=ExcelPreviewResponse,
        dependencies=[Depends(combined_auth)],
    )
    async def preview_excel_source(
        request: Request,
        file: UploadFile = File(...),
        sheet_name: Optional[str] = Form(default=None),
        header_row: int = Form(default=1, ge=1),
        data_start_row: Optional[int] = Form(default=None),
        sample_limit: int = Form(default=20, ge=1, le=100),
        max_rows: int = Form(default=MAX_EXCEL_ROWS_PER_PREVIEW, ge=1, le=MAX_EXCEL_ROWS_PER_PREVIEW),
    ):
        await db_for_request(request)
        filename = file.filename or "uploaded.xlsx"
        lower_name = filename.lower()
        if not lower_name.endswith((".xlsx", ".xlsm", ".xltx", ".xltm")):
            raise HTTPException(
                status_code=400,
                detail="Only .xlsx, .xlsm, .xltx, and .xltm Excel files are supported",
            )

        content, file_size = await _read_upload_limited(file, MAX_EXCEL_UPLOAD_BYTES)
        sheets, selected_sheet, header, start_row, rows, columns, warnings, truncated = _extract_excel_rows(
            content,
            sheet_name=sheet_name,
            header_row=header_row,
            data_start_row=data_start_row,
            max_rows=max_rows,
        )
        raw_content = json.dumps(rows, ensure_ascii=False)
        profile = _profile_structured_source("json", raw_content, sample_limit=sample_limit)
        profile.warnings.extend(warnings)
        source_uri = f"xlsx://{filename}/{selected_sheet}"
        return ExcelPreviewResponse(
            file_name=filename,
            file_size=file_size,
            sheets=sheets,
            selected_sheet=selected_sheet,
            header_row=header,
            data_start_row=start_row,
            row_count=len(rows),
            row_limit=max_rows,
            truncated=truncated,
            columns=columns,
            raw_content=raw_content,
            source_uri=source_uri,
            profile=profile,
            warnings=warnings,
        )

    @router.post(
        "/structured/materialize",
        response_model=StructuredMaterializeResponse,
        dependencies=[Depends(combined_auth)],
    )
    async def materialize_structured_source(request: Request, payload: StructuredMaterializeRequest):
        workspace, db = await db_for_request(request)
        enrichment_rag = None
        if payload.llm_guidance_enrichment.enabled:
            enrichment_rag = await rag_for_workspace(workspace)
            if getattr(enrichment_rag, "llm_model_func", None) is None:
                raise HTTPException(
                    status_code=503,
                    detail="LLM guidance enrichment is enabled, but no LLM is configured for this workspace.",
                )
        profile = _profile_structured_source(payload.source_type, payload.raw_content, sample_limit=20)
        rows, _, _ = _parse_structured_rows(payload.source_type, payload.raw_content)
        profile_dict = profile.model_dump() if hasattr(profile, "model_dump") else profile.dict()

        mapping: dict[str, str] = {}
        for role, column in profile.mapping_suggestions.items():
            if column and column in profile.columns:
                mapping[role] = column
        for role, column in payload.mapping.items():
            if column in profile.columns:
                mapping[str(role)] = column

        guidance_columns = [
            column for column in payload.guidance_columns
            if column in profile.columns and column not in mapping.values()
        ]
        if not guidance_columns:
            guidance_columns = [
                field.name
                for field in profile.fields
                if field.semantic_role in {"category", "title", "id"}
                and field.name != mapping.get("id")
            ][:5]

        if payload.conversion_purpose == "id_lookup" and payload.materialization_mode != "row_per_answer":
            raise HTTPException(
                status_code=400,
                detail="ID lookup conversion requires row_per_answer materialization.",
            )
        validation = _validate_id_lookup_rows(
            rows,
            mapping,
            guidance_columns,
            payload.conversion_purpose,
        )
        if validation.enabled and not validation.ready:
            raise HTTPException(
                status_code=400,
                detail={
                    "message": "ID lookup mapping validation failed.",
                    "validation": validation.model_dump(),
                },
            )

        tags: list[str] = []
        purpose_tag = "id-lookup" if payload.conversion_purpose == "id_lookup" else None
        for tag in [*payload.tags, "structured", payload.source_type, purpose_tag]:
            text = str(tag or "").strip()
            if text and text not in tags:
                tags.append(text)

        if payload.materialization_mode == "row_per_answer" and payload.source_truncated:
            raise HTTPException(
                status_code=400,
                detail=(
                    "The source preview was truncated. Split or filter the source before "
                    "creating one answer per row."
                ),
            )
        if payload.materialization_mode == "row_per_answer" and len(rows) > MAX_STRUCTURED_ROWS_PER_ANSWER_BATCH:
            raise HTTPException(
                status_code=400,
                detail=f"Row-per-answer materialization supports up to {MAX_STRUCTURED_ROWS_PER_ANSWER_BATCH} rows per request",
            )

        answer_ids = (
            [f"ANS-{uuid.uuid4().hex[:12]}" for _ in rows]
            if payload.materialization_mode == "row_per_answer"
            else [f"ANS-{uuid.uuid4().hex[:12]}"]
        )
        source_profile = {
            "structured": {
                **profile_dict,
                "mapping": mapping,
                "guidance_columns": guidance_columns,
                "materialization_mode": payload.materialization_mode,
                "conversion_purpose": payload.conversion_purpose,
            }
        }
        source_metadata = {
            **payload.metadata,
            "created_from": payload.metadata.get("created_from") or "structured_materialize",
            "source_type": "structured",
            "structured_source_type": payload.source_type,
            "materialization_mode": payload.materialization_mode,
            "conversion_purpose": payload.conversion_purpose,
            "source_uri": payload.source_uri,
            "file_name": payload.file_name,
            "source_profile": source_profile,
        }
        snapshot = await _record_source_snapshot(
            db,
            workspace,
            source_type="structured",
            source_uri=payload.source_uri,
            file_name=payload.file_name,
            title=payload.title,
            raw_content=payload.raw_content,
            profile=source_profile,
            metadata={
                **payload.metadata,
                "created_from": "structured_materialize",
                "structured_source_type": payload.source_type,
                "materialization_mode": payload.materialization_mode,
                "conversion_purpose": payload.conversion_purpose,
                "tags": tags,
            },
            created_answer_ids=answer_ids,
            status="materialized",
            task_id=payload.metadata.get("task_id"),
        )
        source_metadata["source_snapshot_id"] = snapshot.snapshot_id
        source_metadata["source_content_hash"] = snapshot.content_hash
        row_source_metadata = {
            key: value
            for key, value in source_metadata.items()
            if key != "source_profile"
        }
        row_source_metadata["source_profile_ref"] = snapshot.snapshot_id

        created_guidance: list[AnswerGuidance] = []
        created_answers: list[AnswerItem] = []
        created_datasets: list[StructuredDataset] = []
        first_source_link: Optional[AnswerSourceLink] = None
        try:
            if payload.materialization_mode == "table_as_dataset":
                answer = await _insert_answer_item(
                    db,
                    workspace,
                    answer_id=answer_ids[0],
                    title=payload.title,
                    body=payload.raw_content.strip(),
                    approved_summary=payload.approved_summary,
                    content_format="plain",
                    display_policy="both",
                    status=payload.status,
                    priority=payload.priority,
                    tags=tags,
                    metadata=source_metadata,
                )
                created_answers.append(answer)
                created_datasets.append(_dataset_from_answer(answer))
                first_source_link = await _record_source_link(
                    db,
                    workspace,
                    answer_id=answer.answer_id,
                    answer_version=answer.version,
                    snapshot_id=snapshot.snapshot_id,
                    link_type="created_from",
                    metadata={
                        "created_from": "structured_materialize",
                        "structured_source_type": payload.source_type,
                        "materialization_mode": payload.materialization_mode,
                        "conversion_purpose": payload.conversion_purpose,
                        "mapping": mapping,
                        "guidance_columns": guidance_columns,
                    },
                )
                created_guidance.extend(
                    await _insert_answer_guidance(
                        db,
                        workspace,
                        answer.answer_id,
                        _guidance_from_structured_profile(
                            rows,
                            mapping,
                            guidance_columns,
                            payload.conversion_purpose,
                        ),
                    )
                )
            else:
                for row_index, row_data in enumerate(rows):
                    row_answer_id = answer_ids[row_index]
                    row_status = _status_for_structured_row(row_data, mapping, payload.status)
                    row_valid_from = _datetime_for_structured_row(row_data, mapping, "valid_from")
                    row_valid_until = _datetime_for_structured_row(row_data, mapping, "valid_until")
                    source_id = _structured_row_text(row_data, mapping.get("id"))
                    row_metadata = {
                        **row_source_metadata,
                        "source_id": source_id or None,
                        "matched_id": source_id or None,
                        "conversion_purpose": payload.conversion_purpose,
                        "source_row_index": row_index,
                        "source_row_hash": _content_hash(json.dumps(row_data, ensure_ascii=False, sort_keys=True)),
                        "source_row": row_data,
                        "row_policy": {
                            "status": row_status,
                            "valid_from": row_valid_from.isoformat() if row_valid_from else None,
                            "valid_until": row_valid_until.isoformat() if row_valid_until else None,
                        },
                    }
                    answer = await _insert_answer_item(
                        db,
                        workspace,
                        answer_id=row_answer_id,
                        title=_title_for_structured_row(payload.title, row_data, mapping, row_index),
                        body=_body_for_structured_row(
                            row_data,
                            mapping,
                            payload.conversion_purpose,
                        ),
                        approved_summary=_summary_for_structured_row(
                            row_data,
                            mapping,
                            payload.approved_summary,
                            payload.conversion_purpose,
                        ),
                        content_format="plain",
                        display_policy="both",
                        status=row_status,
                        priority=payload.priority,
                        tags=_tags_for_structured_row(tags, row_data, mapping, payload.source_type),
                        metadata=row_metadata,
                        valid_from=row_valid_from,
                        valid_until=row_valid_until,
                    )
                    created_answers.append(answer)
                    created_datasets.append(_dataset_from_answer(answer))
                    source_link = await _record_source_link(
                        db,
                        workspace,
                        answer_id=answer.answer_id,
                        answer_version=answer.version,
                        snapshot_id=snapshot.snapshot_id,
                        link_type="created_from_row",
                        metadata={
                            "created_from": "structured_materialize",
                            "structured_source_type": payload.source_type,
                            "materialization_mode": payload.materialization_mode,
                            "conversion_purpose": payload.conversion_purpose,
                            "source_id": source_id or None,
                            "source_row_index": row_index,
                            "mapping": mapping,
                            "guidance_columns": guidance_columns,
                        },
                    )
                    first_source_link = first_source_link or source_link
                    created_guidance.extend(
                        await _insert_answer_guidance(
                            db,
                            workspace,
                            answer.answer_id,
                            _guidance_from_structured_profile(
                                [row_data],
                                mapping,
                                guidance_columns,
                                payload.conversion_purpose,
                            ),
                        )
                    )

            guidance_enrichment_task_id: Optional[str] = None
            guidance_enrichment_stream_url: Optional[str] = None
            vector_rebuild_task_id: Optional[str] = None
            vector_rebuild_stream_url: Optional[str] = None
            if payload.llm_guidance_enrichment.enabled and enrichment_rag is not None:
                from lightrag.api.task_manager import TaskType, get_task_service

                service = get_task_service()
                enrichment_task = service.create_task(
                    task_type=TaskType.FAQ_GUIDANCE_ENRICHMENT,
                    workspace=workspace,
                    metadata={
                        "answer_count": len(answer_ids),
                        "scope": payload.llm_guidance_enrichment.scope,
                        "batch_size": payload.llm_guidance_enrichment.batch_size,
                        "max_suggestions": payload.llm_guidance_enrichment.max_suggestions,
                        "source_snapshot_id": snapshot.snapshot_id,
                    },
                )
                guidance_enrichment_task_id = enrichment_task.task_id
                guidance_enrichment_stream_url = (
                    f"/api/tasks/{enrichment_task.task_id}/stream"
                )
                service.run_in_background(
                    enrichment_task.task_id,
                    _enrich_answer_guidance_background,
                    task_id=enrichment_task.task_id,
                    workspace=workspace,
                    db=db,
                    rag=enrichment_rag,
                    answer_ids=answer_ids,
                    scope=payload.llm_guidance_enrichment.scope,
                    batch_size=payload.llm_guidance_enrichment.batch_size,
                    max_suggestions=payload.llm_guidance_enrichment.max_suggestions,
                )
            else:
                try:
                    vector_rag = await rag_for_workspace(workspace)
                    if _embedding_func_from_rag(vector_rag) is not None:
                        from lightrag.api.task_manager import (
                            TaskType,
                            get_task_service,
                        )

                        service = get_task_service()
                        vector_task = service.create_task(
                            task_type=TaskType.FAQ_VECTOR_REBUILD,
                            workspace=workspace,
                            metadata={
                                "answer_count": len(answer_ids),
                                "source_snapshot_id": snapshot.snapshot_id,
                                "automatic": True,
                            },
                        )
                        vector_rebuild_task_id = vector_task.task_id
                        vector_rebuild_stream_url = (
                            f"/api/tasks/{vector_task.task_id}/stream"
                        )
                        service.run_in_background(
                            vector_task.task_id,
                            _rebuild_answer_vectors_background,
                            task_id=vector_task.task_id,
                            workspace=workspace,
                            db=db,
                            rag=vector_rag,
                            answer_ids=answer_ids,
                        )
                except Exception as exc:
                    logger.warning(
                        "[Answers] Automatic FAQ vector task could not start: %s",
                        exc,
                    )

            response_item_limit = 20
            return StructuredMaterializeResponse(
                answer=created_answers[0],
                dataset=created_datasets[0],
                answers=created_answers[:response_item_limit],
                datasets=created_datasets[:response_item_limit],
                answer_count=len(created_answers),
                answers_truncated=len(created_answers) > response_item_limit,
                profile=profile,
                guidance=created_guidance[:response_item_limit],
                guidance_count=len(created_guidance),
                guidance_truncated=len(created_guidance) > response_item_limit,
                snapshot=snapshot,
                source_link=first_source_link,
                validation=validation,
                guidance_enrichment_task_id=guidance_enrichment_task_id,
                guidance_enrichment_stream_url=guidance_enrichment_stream_url,
                vector_rebuild_task_id=vector_rebuild_task_id,
                vector_rebuild_stream_url=vector_rebuild_stream_url,
            )
        except HTTPException:
            raise
        except Exception as e:
            try:
                for cleanup_answer_id in answer_ids:
                    await db.execute(
                        "DELETE FROM LIGHTRAG_ANSWER_GUIDANCE WHERE workspace = $1 AND answer_id = $2",
                        {"workspace": workspace, "answer_id": cleanup_answer_id},
                    )
                    await db.execute(
                        "DELETE FROM LIGHTRAG_ANSWER_SOURCE_LINKS WHERE workspace = $1 AND answer_id = $2",
                        {"workspace": workspace, "answer_id": cleanup_answer_id},
                    )
                    await db.execute(
                        "DELETE FROM LIGHTRAG_ANSWER_REVISIONS WHERE workspace = $1 AND answer_id = $2",
                        {"workspace": workspace, "answer_id": cleanup_answer_id},
                    )
                    await db.execute(
                        "DELETE FROM LIGHTRAG_ANSWER_ITEMS WHERE workspace = $1 AND answer_id = $2",
                        {"workspace": workspace, "answer_id": cleanup_answer_id},
                    )
                await db.execute(
                    "DELETE FROM LIGHTRAG_ANSWER_SOURCE_SNAPSHOTS WHERE workspace = $1 AND snapshot_id = $2",
                    {"workspace": workspace, "snapshot_id": snapshot.snapshot_id},
                )
            except Exception:
                pass
            message = str(e)
            logger.error("Failed to materialize structured source: %s\n%s", e, traceback.format_exc())
            raise HTTPException(status_code=500, detail=message)

    @router.get("/structured/datasets", response_model=list[StructuredDataset], dependencies=[Depends(combined_auth)])
    async def list_structured_datasets(
        request: Request,
        structured_only: bool = Query(default=True),
        status: Optional[str] = Query(default="all"),
    ):
        workspace, db = await db_for_request(request)
        params: list[Any] = [workspace]
        where = ["workspace = $1"]
        if status and status != "all":
            params.append(status)
            where.append(f"status = ${len(params)}")
        rows = await db.query(
            f"""
            SELECT workspace, answer_id, title, body, approved_summary, content_format,
                   display_policy, status, version, valid_from, valid_until, priority,
                   tags, metadata, publish_time, create_time, update_time
            FROM LIGHTRAG_ANSWER_ITEMS
            WHERE {' AND '.join(where)}
            ORDER BY update_time DESC, answer_id ASC
            LIMIT 500
            """,
            params,
            multirows=True,
        )
        datasets = [_dataset_from_answer(_answer_from_row(dict(row))) for row in rows or []]
        if structured_only:
            datasets = [
                dataset
                for dataset in datasets
                if dataset.metadata.get("materialization_mode") != "row_per_answer"
                and (
                    dataset.row_count > 1
                    or len(dataset.columns) > 9
                    or dataset.kind in {"json", "table", "structured"}
                )
            ]
        return datasets

    @router.post("/structured/query", response_model=StructuredQueryResponse, dependencies=[Depends(combined_auth)])
    async def query_structured_dataset(request: Request, payload: StructuredQueryRequest):
        started_at = time.perf_counter()
        workspace, db = await db_for_request(request)
        row = await _get_answer_row(db, workspace, payload.answer_id)
        answer = _answer_from_row(dict(row))
        rows, columns, _ = _rows_from_answer(answer)
        dataset = _dataset_from_answer(answer)
        invalid_fields = [item.field for item in payload.filters if item.field not in columns]
        if invalid_fields:
            raise HTTPException(
                status_code=400,
                detail=f"Unknown structured field(s): {', '.join(invalid_fields)}",
            )

        filtered_rows = rows
        for item in payload.filters:
            filtered_rows = [row for row in filtered_rows if _matches_filter(row, item)]

        pseudo_sql = _pseudo_sql(dataset, payload.filters, payload.limit)
        latency_ms = round((time.perf_counter() - started_at) * 1000, 2)
        await _record_structured_lookup_log(
            db,
            workspace,
            dataset_id=answer.answer_id,
            dataset_title=answer.title,
            pseudo_sql=pseudo_sql,
            filters=payload.filters,
            result_count=len(filtered_rows),
            preview_only=payload.preview_only,
            latency_ms=latency_ms,
            metadata={
                "limit": payload.limit,
                "source_type": dataset.source_type,
                "kind": dataset.kind,
            },
        )
        return StructuredQueryResponse(
            answer_id=answer.answer_id,
            title=answer.title,
            pseudo_sql=pseudo_sql,
            columns=columns,
            rows=[] if payload.preview_only else filtered_rows[: payload.limit],
            row_count=len(filtered_rows),
            preview_only=payload.preview_only,
        )

    @router.get("/structured/query/logs", response_model=list[StructuredLookupLog], dependencies=[Depends(combined_auth)])
    async def list_structured_lookup_logs(
        request: Request,
        dataset_id: Optional[str] = Query(default=None),
        limit: int = Query(default=30, ge=1, le=200),
    ):
        workspace, db = await db_for_request(request)
        params: list[Any] = [workspace]
        where = ["workspace = $1"]
        if dataset_id:
            params.append(dataset_id)
            where.append(f"dataset_id = ${len(params)}")
        params.append(limit)
        rows = await db.query(
            f"""
            SELECT log_id, workspace, dataset_id, dataset_title, pseudo_sql, filters,
                   result_count, preview_only, latency_ms, metadata, create_time
            FROM LIGHTRAG_STRUCTURED_LOOKUP_LOGS
            WHERE {' AND '.join(where)}
            ORDER BY create_time DESC
            LIMIT ${len(params)}
            """,
            params,
            multirows=True,
        )
        return [_structured_lookup_log_from_row(dict(row)) for row in rows or []]

    async def _get_connector_row(db, workspace: str, connector_id: str) -> dict[str, Any]:
        row = await db.query(
            """
            SELECT connector_id, workspace, name, connector_type, status, config,
                   auth_ref, refresh_policy, enabled, metadata, create_time, update_time
            FROM LIGHTRAG_SOURCE_CONNECTORS
            WHERE workspace = $1 AND connector_id = $2
            """,
            [workspace, connector_id],
        )
        if not row:
            raise HTTPException(status_code=404, detail=f"Source connector '{connector_id}' not found")
        return dict(row)

    @router.get("/connectors", response_model=list[SourceConnector], dependencies=[Depends(combined_auth)])
    async def list_source_connectors(
        request: Request,
        connector_type: Optional[str] = Query(default=None),
        status: Optional[str] = Query(default=None),
        search: Optional[str] = Query(default=None),
        limit: int = Query(default=100, ge=1, le=500),
    ):
        workspace, db = await db_for_request(request)
        params: list[Any] = [workspace]
        where = ["workspace = $1"]
        if connector_type and connector_type != "all":
            params.append(connector_type)
            where.append(f"connector_type = ${len(params)}")
        if status and status != "all":
            params.append(status)
            where.append(f"status = ${len(params)}")
        if search:
            params.append(f"%{search}%")
            p = f"${len(params)}"
            where.append(f"(name ILIKE {p} OR connector_id ILIKE {p})")
        params.append(limit)
        rows = await db.query(
            f"""
            SELECT connector_id, workspace, name, connector_type, status, config,
                   auth_ref, refresh_policy, enabled, metadata, create_time, update_time
            FROM LIGHTRAG_SOURCE_CONNECTORS
            WHERE {' AND '.join(where)}
            ORDER BY update_time DESC, connector_id ASC
            LIMIT ${len(params)}
            """,
            params,
            multirows=True,
        )
        return [_source_connector_from_row(dict(row)) for row in rows or []]

    @router.post("/connectors", response_model=SourceConnector, dependencies=[Depends(combined_auth)])
    async def create_source_connector(request: Request, payload: SourceConnectorCreateRequest):
        workspace, db = await db_for_request(request)
        connector_id = payload.connector_id or f"conn-{uuid.uuid4().hex[:12]}"
        try:
            row = await db.query(
                """
                INSERT INTO LIGHTRAG_SOURCE_CONNECTORS
                    (connector_id, workspace, name, connector_type, status, config,
                     auth_ref, refresh_policy, enabled, metadata)
                VALUES ($1, $2, $3, $4, $5, $6::jsonb, $7, $8::jsonb, $9, $10::jsonb)
                RETURNING connector_id, workspace, name, connector_type, status, config,
                          auth_ref, refresh_policy, enabled, metadata, create_time, update_time
                """,
                [
                    connector_id,
                    workspace,
                    payload.name,
                    payload.connector_type,
                    payload.status,
                    _json(payload.config),
                    payload.auth_ref,
                    _json(payload.refresh_policy),
                    payload.enabled,
                    _json(payload.metadata),
                ],
            )
        except Exception as e:
            message = str(e)
            if "duplicate" in message.lower() or "unique" in message.lower():
                raise HTTPException(status_code=409, detail=f"Source connector '{connector_id}' already exists")
            logger.error("Failed to create source connector: %s\n%s", e, traceback.format_exc())
            raise HTTPException(status_code=500, detail=message)
        if not row:
            raise HTTPException(status_code=500, detail="Failed to create source connector")
        return _source_connector_from_row(dict(row))

    @router.get("/connectors/{connector_id}", response_model=SourceConnector, dependencies=[Depends(combined_auth)])
    async def get_source_connector(request: Request, connector_id: str):
        workspace, db = await db_for_request(request)
        return _source_connector_from_row(await _get_connector_row(db, workspace, connector_id))

    @router.patch("/connectors/{connector_id}", response_model=SourceConnector, dependencies=[Depends(combined_auth)])
    async def update_source_connector(request: Request, connector_id: str, payload: SourceConnectorUpdateRequest):
        workspace, db = await db_for_request(request)
        current = _source_connector_from_row(await _get_connector_row(db, workspace, connector_id))
        row = await db.query(
            """
            UPDATE LIGHTRAG_SOURCE_CONNECTORS
            SET name = $3,
                status = $4,
                config = $5::jsonb,
                auth_ref = $6,
                refresh_policy = $7::jsonb,
                enabled = $8,
                metadata = $9::jsonb,
                update_time = NOW()
            WHERE workspace = $1 AND connector_id = $2
            RETURNING connector_id, workspace, name, connector_type, status, config,
                      auth_ref, refresh_policy, enabled, metadata, create_time, update_time
            """,
            [
                workspace,
                connector_id,
                payload.name if payload.name is not None else current.name,
                payload.status if payload.status is not None else current.status,
                _json(payload.config if payload.config is not None else current.config),
                payload.auth_ref if payload.auth_ref is not None else current.auth_ref,
                _json(payload.refresh_policy if payload.refresh_policy is not None else current.refresh_policy),
                payload.enabled if payload.enabled is not None else current.enabled,
                _json(payload.metadata if payload.metadata is not None else current.metadata),
            ],
        )
        if not row:
            raise HTTPException(status_code=500, detail="Failed to update source connector")
        return _source_connector_from_row(dict(row))

    @router.post(
        "/connectors/{connector_id}/sample",
        response_model=SourceConnectorSampleResponse,
        dependencies=[Depends(combined_auth)],
    )
    async def sample_source_connector(
        request: Request,
        connector_id: str,
        payload: SourceConnectorSampleRequest = SourceConnectorSampleRequest(),
    ):
        workspace, db = await db_for_request(request)
        connector = _source_connector_from_row(await _get_connector_row(db, workspace, connector_id))
        if not connector.enabled:
            raise HTTPException(status_code=409, detail=f"Source connector '{connector_id}' is disabled")
        return _connector_sample_response_from_raw(
            connector,
            *await _connector_raw_content_from_db(db, connector, payload.limit),
            limit=payload.limit,
        )

    @router.post(
        "/connectors/{connector_id}/profile",
        response_model=StructuredProfileResponse,
        dependencies=[Depends(combined_auth)],
    )
    async def profile_source_connector(
        request: Request,
        connector_id: str,
        payload: SourceConnectorSampleRequest = SourceConnectorSampleRequest(),
    ):
        workspace, db = await db_for_request(request)
        connector = _source_connector_from_row(await _get_connector_row(db, workspace, connector_id))
        sample = _connector_sample_response_from_raw(
            connector,
            *await _connector_raw_content_from_db(db, connector, payload.limit),
            limit=payload.limit,
        )
        return _profile_structured_source(sample.source_type, sample.raw_content, sample_limit=min(payload.limit, 100))

    @router.post(
        "/connectors/{connector_id}/mapping/preview",
        response_model=SourceConnectorMappingPreviewResponse,
        dependencies=[Depends(combined_auth)],
    )
    async def preview_source_connector_mapping(
        request: Request,
        connector_id: str,
        payload: SourceConnectorMappingPreviewRequest = SourceConnectorMappingPreviewRequest(),
    ):
        workspace, db = await db_for_request(request)
        connector = _source_connector_from_row(await _get_connector_row(db, workspace, connector_id))
        sample = _connector_sample_response_from_raw(
            connector,
            *await _connector_raw_content_from_db(db, connector, 100),
            limit=100,
        )
        profile = _profile_structured_source(sample.source_type, sample.raw_content, sample_limit=20)
        mapping = _mapping_from_profile(profile, payload.mapping)
        guidance_columns = _guidance_columns_from_profile(
            profile,
            mapping,
            payload.guidance_columns,
        )
        validation = _validate_id_lookup_rows(
            sample.rows,
            mapping,
            guidance_columns,
            payload.conversion_purpose,
        )
        return SourceConnectorMappingPreviewResponse(
            connector=connector,
            sample=sample,
            profile=profile,
            mapping=mapping,
            guidance_columns=guidance_columns,
            materialization_modes=["table_as_dataset", "row_per_answer"],
            validation=validation,
        )

    @router.post(
        "/connectors/{connector_id}/materialize",
        response_model=SourceConnectorMaterializeResponse,
        dependencies=[Depends(combined_auth)],
    )
    async def materialize_source_connector(
        request: Request,
        connector_id: str,
        payload: SourceConnectorMaterializeRequest,
    ):
        workspace, db = await db_for_request(request)
        connector = _source_connector_from_row(await _get_connector_row(db, workspace, connector_id))
        if payload.conversion_purpose == "id_lookup" and payload.materialization_mode != "row_per_answer":
            raise HTTPException(
                status_code=400,
                detail="ID lookup conversion requires row_per_answer materialization.",
            )
        sample = _connector_sample_response_from_raw(
            connector,
            *await _connector_raw_content_from_db(db, connector, 1000),
            limit=1000,
        )
        if payload.materialization_mode == "row_per_answer" and sample.truncated:
            raise HTTPException(
                status_code=400,
                detail=(
                    f"DB row-per-answer materialization supports up to {sample.row_limit} rows per request. "
                    "Filter or split the source table before creating answer candidates."
                ),
            )
        profile = _profile_structured_source(sample.source_type, sample.raw_content, sample_limit=20)
        mapping = _mapping_from_profile(profile, payload.mapping)
        guidance_columns = _guidance_columns_from_profile(profile, mapping, payload.guidance_columns)
        tags = []
        for tag in [*payload.tags, "connector", connector.connector_type]:
            text = str(tag or "").strip()
            if text and text not in tags:
                tags.append(text)
        materialized = await materialize_structured_source(
            request,
            StructuredMaterializeRequest(
                source_type=sample.source_type,
                raw_content=sample.raw_content,
                title=payload.title or connector.name,
                approved_summary=None,
                source_uri=sample.source_uri or f"connector://{connector.connector_id}",
                file_name=None,
                status=payload.status,
                priority=0,
                tags=tags,
                mapping=mapping,
                guidance_columns=guidance_columns,
                materialization_mode=payload.materialization_mode,
                conversion_purpose=payload.conversion_purpose,
                source_truncated=sample.truncated,
                llm_guidance_enrichment=payload.llm_guidance_enrichment,
                metadata={
                    **payload.metadata,
                    "created_from": "source_connector",
                    "connector_id": connector.connector_id,
                    "connector_type": connector.connector_type,
                },
            ),
        )
        await db.query(
            """
            UPDATE LIGHTRAG_SOURCE_CONNECTORS
            SET status = CASE WHEN status = 'draft' THEN 'active' ELSE status END,
                metadata = metadata || $3::jsonb,
                update_time = NOW()
            WHERE workspace = $1 AND connector_id = $2
            RETURNING connector_id
            """,
            [
                workspace,
                connector.connector_id,
                _json({
                    "last_materialized_at": _now().isoformat(),
                    "last_materialized_answer_ids": [answer.answer_id for answer in materialized.answers],
                    "last_materialized_answer_count": materialized.answer_count,
                    "last_materialization_mode": payload.materialization_mode,
                }),
            ],
        )
        refreshed = _source_connector_from_row(await _get_connector_row(db, workspace, connector_id))
        return SourceConnectorMaterializeResponse(
            connector=refreshed,
            sample=_compact_connector_materialization_sample(sample, profile),
            materialized=materialized,
        )

    @router.get("/sources/snapshots", response_model=list[AnswerSourceSnapshot], dependencies=[Depends(combined_auth)])
    async def list_source_snapshots(
        request: Request,
        source_type: Optional[str] = Query(default=None),
        answer_id: Optional[str] = Query(default=None),
        search: Optional[str] = Query(default=None),
        limit: int = Query(default=50, ge=1, le=200),
    ):
        workspace, db = await db_for_request(request)
        params: list[Any] = [workspace]
        where = ["s.workspace = $1"]
        join_sql = ""
        if answer_id:
            join_sql = """
            JOIN LIGHTRAG_ANSWER_SOURCE_LINKS l
              ON l.workspace = s.workspace AND l.snapshot_id = s.snapshot_id
            """
            params.append(answer_id)
            where.append(f"l.answer_id = ${len(params)}")
        if source_type and source_type != "all":
            params.append(source_type)
            where.append(f"s.source_type = ${len(params)}")
        if search:
            params.append(f"%{search}%")
            p = f"${len(params)}"
            where.append(f"(s.title ILIKE {p} OR s.source_uri ILIKE {p} OR s.file_name ILIKE {p} OR s.raw_content ILIKE {p})")
        params.append(limit)
        rows = await db.query(
            f"""
            SELECT DISTINCT s.snapshot_id, s.workspace, s.source_type, s.source_uri,
                   s.file_name, s.title, LEFT(s.raw_content, 600) AS content_preview,
                   s.content_hash, s.content_length, s.profile, s.metadata,
                   s.created_answer_ids, s.status, s.task_id, s.create_time
            FROM LIGHTRAG_ANSWER_SOURCE_SNAPSHOTS s
            {join_sql}
            WHERE {' AND '.join(where)}
            ORDER BY s.create_time DESC
            LIMIT ${len(params)}
            """,
            params,
            multirows=True,
        )
        return [_source_snapshot_from_row(dict(row)) for row in rows or []]

    @router.get(
        "/aliases",
        response_model=list[AnswerAliasGroup],
        dependencies=[Depends(combined_auth)],
    )
    async def list_answer_aliases(
        request: Request,
        search: Optional[str] = Query(default=None),
    ):
        workspace, db = await db_for_request(request)
        groups = await _answer_alias_groups(db, workspace)
        if not search or not search.strip():
            return groups
        needle = _normalise_text(search)
        return [
            group
            for group in groups
            if needle in _normalise_text(group.canonical_term)
            or any(needle in _normalise_text(alias) for alias in group.aliases)
        ]

    @router.post(
        "/aliases",
        response_model=AnswerAliasGroup,
        dependencies=[Depends(combined_auth)],
    )
    async def create_answer_alias(
        request: Request,
        payload: AnswerAliasCreateRequest,
    ):
        workspace, db = await db_for_request(request)
        canonical_term = " ".join(payload.canonical_term.split())
        aliases: list[str] = []
        seen = {_normalise_text(canonical_term)}
        for alias in payload.aliases:
            value = " ".join(str(alias or "").split())
            key = _normalise_text(value)
            if len(value) < 2 or key in seen:
                continue
            seen.add(key)
            aliases.append(value[:200])
        if not aliases:
            raise HTTPException(
                status_code=400,
                detail="At least one distinct alias is required.",
            )
        try:
            row = await db.query(
                """
                INSERT INTO LIGHTRAG_ANSWER_TERM_ALIASES
                    (
                        alias_id, workspace, canonical_term, aliases, enabled,
                        source, metadata
                    )
                VALUES ($1, $2, $3, $4::jsonb, $5, 'workspace', $6::jsonb)
                RETURNING alias_id, workspace, canonical_term, aliases, enabled,
                          source, metadata, create_time, update_time
                """,
                [
                    f"aalias-{uuid.uuid4().hex[:16]}",
                    workspace,
                    canonical_term,
                    _json_list(aliases),
                    payload.enabled,
                    _json(payload.metadata),
                ],
            )
        except Exception as exc:
            message = str(exc)
            if "duplicate" in message.lower() or "unique" in message.lower():
                raise HTTPException(
                    status_code=409,
                    detail=f"Alias group '{canonical_term}' already exists.",
                )
            raise
        if not row:
            raise HTTPException(status_code=500, detail="Failed to create alias group.")
        await _mark_all_answer_graph_stale(db, workspace)
        return _alias_group_from_row(dict(row))

    @router.get(
        "/aliases/candidates",
        response_model=list[AnswerTermCandidate],
        dependencies=[Depends(combined_auth)],
    )
    async def list_answer_term_candidates(
        request: Request,
        status: Optional[TermCandidateStatus] = Query(default="suggested"),
        search: Optional[str] = Query(default=None),
        limit: int = Query(default=200, ge=1, le=1000),
    ):
        workspace, db = await db_for_request(request)
        params: list[Any] = [workspace]
        where = ["workspace = $1"]
        if status:
            params.append(status)
            where.append(f"status = ${len(params)}")
        if search and search.strip():
            params.append(f"%{search.strip()}%")
            where.append(
                f"(canonical_term ILIKE ${len(params)} "
                f"OR aliases::text ILIKE ${len(params)})"
            )
        params.append(limit)
        rows = await db.query(
            f"""
            SELECT candidate_id, workspace, canonical_term, aliases, term_type,
                   status, confidence, rationale, evidence, source, metadata,
                   create_time, update_time
            FROM LIGHTRAG_ANSWER_TERM_CANDIDATES
            WHERE {' AND '.join(where)}
            ORDER BY confidence DESC, update_time DESC
            LIMIT ${len(params)}
            """,
            params,
            multirows=True,
        )
        return [_term_candidate_from_row(dict(row)) for row in rows or []]

    @router.post(
        "/aliases/candidates/analyze",
        response_model=AnswerTermDiscoveryResponse,
        dependencies=[Depends(combined_auth)],
    )
    async def analyze_answer_terms(
        request: Request,
        payload: AnswerTermDiscoveryRequest,
    ):
        from lightrag.api.task_manager import TaskType, get_task_service

        workspace, db = await db_for_request(request)
        workspace_rag = await rag_for_workspace(workspace)
        if getattr(workspace_rag, "llm_model_func", None) is None:
            raise HTTPException(
                status_code=409,
                detail="LLM is not configured for this workspace.",
            )

        service = get_task_service()
        task = service.create_task(
            task_type=TaskType.FAQ_TERM_DISCOVERY,
            workspace=workspace,
            metadata={
                "include_drafts": payload.include_drafts,
                "include_no_match_queries": payload.include_no_match_queries,
                "answer_limit": payload.answer_limit,
                "event_limit": payload.event_limit,
                "batch_size": payload.batch_size,
            },
        )
        service.run_in_background(
            task.task_id,
            _discover_answer_terms_background,
            task_id=task.task_id,
            workspace=workspace,
            db=db,
            rag=workspace_rag,
            include_drafts=payload.include_drafts,
            include_no_match_queries=payload.include_no_match_queries,
            answer_limit=payload.answer_limit,
            event_limit=payload.event_limit,
            batch_size=payload.batch_size,
        )
        return AnswerTermDiscoveryResponse(
            task_id=task.task_id,
            stream_url=f"/api/tasks/{task.task_id}/stream",
            message="AI term discovery started.",
        )

    @router.post(
        "/aliases/candidates/{candidate_id}/approve",
        response_model=AnswerTermCandidateActionResponse,
        dependencies=[Depends(combined_auth)],
    )
    async def approve_answer_term_candidate(request: Request, candidate_id: str):
        workspace, db = await db_for_request(request)
        candidate_row = await db.query(
            """
            SELECT candidate_id, workspace, canonical_term, aliases, term_type,
                   status, confidence, rationale, evidence, source, metadata,
                   create_time, update_time
            FROM LIGHTRAG_ANSWER_TERM_CANDIDATES
            WHERE workspace = $1 AND candidate_id = $2
            """,
            [workspace, candidate_id],
        )
        if not candidate_row:
            raise HTTPException(
                status_code=404,
                detail=f"Term candidate '{candidate_id}' not found.",
            )
        candidate = _term_candidate_from_row(dict(candidate_row))

        groups = await _answer_alias_groups(db, workspace)
        member_to_group: dict[str, AnswerAliasGroup] = {}
        for group in groups:
            for member in [group.canonical_term, *group.aliases]:
                member_to_group.setdefault(_normalise_text(member), group)

        canonical_group = member_to_group.get(
            _normalise_text(candidate.canonical_term)
        )
        canonical_term = (
            canonical_group.canonical_term
            if canonical_group is not None
            else candidate.canonical_term
        )
        canonical_key = _normalise_text(canonical_term)
        aliases: list[str] = []
        seen = {canonical_key}
        for alias in candidate.aliases:
            alias_key = _normalise_text(alias)
            if alias_key in seen:
                continue
            conflicting_group = member_to_group.get(alias_key)
            if (
                conflicting_group is not None
                and _normalise_text(conflicting_group.canonical_term)
                != canonical_key
            ):
                continue
            seen.add(alias_key)
            aliases.append(alias)

        workspace_group_row = await db.query(
            """
            SELECT alias_id, workspace, canonical_term, aliases, enabled, source,
                   metadata, create_time, update_time
            FROM LIGHTRAG_ANSWER_TERM_ALIASES
            WHERE workspace = $1 AND LOWER(canonical_term) = LOWER($2)
            """,
            [workspace, canonical_term],
        )
        alias_group: Optional[AnswerAliasGroup] = None
        if workspace_group_row:
            workspace_group = _alias_group_from_row(dict(workspace_group_row))
            merged_aliases = list(workspace_group.aliases)
            merged_keys = {
                _normalise_text(value)
                for value in [workspace_group.canonical_term, *merged_aliases]
            }
            for alias in aliases:
                if _normalise_text(alias) not in merged_keys:
                    merged_aliases.append(alias)
                    merged_keys.add(_normalise_text(alias))
            updated_group = await db.query(
                """
                UPDATE LIGHTRAG_ANSWER_TERM_ALIASES
                SET aliases = $3::jsonb,
                    enabled = TRUE,
                    metadata = metadata || $4::jsonb,
                    update_time = NOW()
                WHERE workspace = $1 AND alias_id = $2
                RETURNING alias_id, workspace, canonical_term, aliases, enabled,
                          source, metadata, create_time, update_time
                """,
                [
                    workspace,
                    workspace_group.alias_id,
                    _json_list(merged_aliases[:30]),
                    _json(
                        {
                            "last_approved_candidate_id": candidate_id,
                            "updated_from": "ai_term_candidate",
                        }
                    ),
                ],
            )
            alias_group = _alias_group_from_row(dict(updated_group))
        elif aliases:
            created_group = await db.query(
                """
                INSERT INTO LIGHTRAG_ANSWER_TERM_ALIASES
                    (
                        alias_id, workspace, canonical_term, aliases, enabled,
                        source, metadata
                    )
                VALUES ($1, $2, $3, $4::jsonb, TRUE, 'ai_approved', $5::jsonb)
                RETURNING alias_id, workspace, canonical_term, aliases, enabled,
                          source, metadata, create_time, update_time
                """,
                [
                    f"aalias-{uuid.uuid4().hex[:16]}",
                    workspace,
                    canonical_term,
                    _json_list(aliases[:30]),
                    _json(
                        {
                            "approved_candidate_id": candidate_id,
                            "created_from": "ai_term_candidate",
                        }
                    ),
                ],
            )
            alias_group = _alias_group_from_row(dict(created_group))
        elif canonical_group is not None:
            alias_group = canonical_group
        else:
            raise HTTPException(
                status_code=409,
                detail="No new non-conflicting aliases remain to approve.",
            )

        approved_row = await db.query(
            """
            UPDATE LIGHTRAG_ANSWER_TERM_CANDIDATES
            SET status = 'approved',
                metadata = metadata || $3::jsonb,
                update_time = NOW()
            WHERE workspace = $1 AND candidate_id = $2
            RETURNING candidate_id, workspace, canonical_term, aliases, term_type,
                      status, confidence, rationale, evidence, source, metadata,
                      create_time, update_time
            """,
            [
                workspace,
                candidate_id,
                _json(
                    {
                        "approved_alias_id": alias_group.alias_id,
                        "approved_at": datetime.now(timezone.utc).isoformat(),
                    }
                ),
            ],
        )
        await _mark_all_answer_graph_stale(db, workspace)
        return AnswerTermCandidateActionResponse(
            candidate=_term_candidate_from_row(dict(approved_row)),
            alias_group=alias_group,
        )

    @router.post(
        "/aliases/candidates/{candidate_id}/reject",
        response_model=AnswerTermCandidateActionResponse,
        dependencies=[Depends(combined_auth)],
    )
    async def reject_answer_term_candidate(request: Request, candidate_id: str):
        workspace, db = await db_for_request(request)
        existing = await db.query(
            """
            SELECT status
            FROM LIGHTRAG_ANSWER_TERM_CANDIDATES
            WHERE workspace = $1 AND candidate_id = $2
            """,
            [workspace, candidate_id],
        )
        if not existing:
            raise HTTPException(
                status_code=404,
                detail=f"Term candidate '{candidate_id}' not found.",
            )
        if str(existing.get("status")) == "approved":
            raise HTTPException(
                status_code=409,
                detail="Approved terms must be removed from the active term groups.",
            )
        rejected_row = await db.query(
            """
            UPDATE LIGHTRAG_ANSWER_TERM_CANDIDATES
            SET status = 'rejected',
                metadata = metadata || $3::jsonb,
                update_time = NOW()
            WHERE workspace = $1 AND candidate_id = $2
            RETURNING candidate_id, workspace, canonical_term, aliases, term_type,
                      status, confidence, rationale, evidence, source, metadata,
                      create_time, update_time
            """,
            [
                workspace,
                candidate_id,
                _json({"rejected_at": datetime.now(timezone.utc).isoformat()}),
            ],
        )
        return AnswerTermCandidateActionResponse(
            candidate=_term_candidate_from_row(dict(rejected_row)),
        )

    @router.delete(
        "/aliases/{alias_id}",
        dependencies=[Depends(combined_auth)],
    )
    async def delete_answer_alias(request: Request, alias_id: str):
        workspace, db = await db_for_request(request)
        if alias_id.startswith("builtin-"):
            raise HTTPException(
                status_code=409,
                detail="Built-in alias groups cannot be deleted.",
            )
        row = await db.query(
            """
            DELETE FROM LIGHTRAG_ANSWER_TERM_ALIASES
            WHERE workspace = $1 AND alias_id = $2
            RETURNING alias_id
            """,
            [workspace, alias_id],
        )
        if not row:
            raise HTTPException(status_code=404, detail=f"Alias group '{alias_id}' not found.")
        await _mark_all_answer_graph_stale(db, workspace)
        return {"message": "Alias group deleted", "alias_id": alias_id}

    @router.get(
        "/{answer_id}/assets",
        response_model=list[AnswerAsset],
        dependencies=[Depends(combined_auth)],
    )
    async def list_answer_assets(
        request: Request,
        answer_id: str,
        include_inactive: bool = Query(False),
    ):
        workspace, db = await db_for_request(request)
        await _get_answer_row(db, workspace, answer_id)
        assets = await _answer_assets_by_answer(
            db,
            workspace,
            [answer_id],
            include_inactive=include_inactive,
        )
        return assets.get(answer_id, [])

    @router.post(
        "/{answer_id}/assets",
        response_model=AnswerAsset,
        dependencies=[Depends(combined_auth)],
    )
    async def create_answer_asset(
        request: Request,
        answer_id: str,
        payload: AnswerAssetCreateRequest,
    ):
        workspace, db = await db_for_request(request)
        current = await _get_answer_row(db, workspace, answer_id)
        external_url = str(payload.external_url or "").strip()
        content_text = str(payload.content_text or "").strip()
        if external_url:
            storage_type: AnswerAssetStorageType = "external"
            storage_uri = _validate_external_asset_url(external_url)
        elif payload.asset_type == "table" and content_text:
            storage_type = "inline"
            storage_uri = None
        else:
            raise HTTPException(
                status_code=422,
                detail=(
                    "Provide external_url, or provide content_text for an inline "
                    "table asset"
                ),
            )

        asset_id = f"asset-{uuid.uuid4().hex}"
        file_name = payload.file_name
        if not file_name and storage_uri:
            file_name = Path(urlparse(storage_uri).path).name or None
        mime_type = payload.mime_type
        if not mime_type and file_name:
            mime_type = mimetypes.guess_type(file_name)[0]
        await db.query(
            """
            INSERT INTO LIGHTRAG_ANSWER_ASSETS
                (asset_id, workspace, answer_id, answer_version, asset_type,
                 storage_type, storage_uri, file_name, mime_type, file_size,
                 caption, alt_text, search_text, content_text, display_order,
                 metadata)
            VALUES
                ($1, $2, $3, $4, $5, $6, $7, $8, $9, $10,
                 $11, $12, $13, $14, $15, $16::jsonb)
            RETURNING asset_id, workspace, answer_id, answer_version, asset_type,
                      storage_type, storage_uri, file_name, mime_type, file_size,
                      caption, alt_text, search_text, content_text, display_order,
                      is_active, metadata, create_time, update_time
            """,
            [
                asset_id,
                workspace,
                answer_id,
                int(current.get("version") or 1) + 1,
                payload.asset_type,
                storage_type,
                storage_uri,
                file_name,
                mime_type,
                len(content_text.encode("utf-8")) if content_text else 0,
                payload.caption,
                payload.alt_text,
                payload.search_text,
                payload.content_text,
                payload.display_order,
                _json(payload.metadata),
            ],
        )
        await _touch_answer_after_asset_change(db, workspace, answer_id)
        return _answer_asset_from_row(
            await _get_answer_asset_row(db, workspace, answer_id, asset_id)
        )

    @router.post(
        "/{answer_id}/assets/upload",
        response_model=AnswerAsset,
        dependencies=[Depends(combined_auth)],
    )
    async def upload_answer_asset(
        request: Request,
        answer_id: str,
        file: UploadFile = File(...),
        asset_type: Optional[AnswerAssetType] = Form(None),
        caption: Optional[str] = Form(None),
        alt_text: Optional[str] = Form(None),
        search_text: Optional[str] = Form(None),
        display_order: int = Form(0),
        metadata_json: str = Form("{}"),
    ):
        workspace, db = await db_for_request(request)
        current = await _get_answer_row(db, workspace, answer_id)
        safe_name = _validate_answer_asset_file(file.filename or "")
        content, file_size = await _read_answer_asset_upload(file)
        guessed_mime = mimetypes.guess_type(safe_name)[0]
        mime_type = str(file.content_type or guessed_mime or "application/octet-stream")
        resolved_asset_type = asset_type or _infer_answer_asset_type(
            safe_name,
            mime_type,
        )
        metadata = _coerce_json(metadata_json, {})
        if not isinstance(metadata, dict):
            raise HTTPException(status_code=422, detail="metadata_json must be a JSON object")

        asset_id = f"asset-{uuid.uuid4().hex}"
        storage_type: AnswerAssetStorageType = "local"
        storage_uri: Optional[str] = None
        s3_client = get_s3_client()
        if s3_client.is_enabled():
            temporary_path: Optional[Path] = None
            try:
                with tempfile.NamedTemporaryFile(
                    prefix="answer-asset-",
                    suffix=Path(safe_name).suffix,
                    delete=False,
                ) as temporary:
                    temporary.write(content)
                    temporary_path = Path(temporary.name)
                if resolved_asset_type == "image":
                    storage_uri = await s3_client.upload_image(
                        temporary_path,
                        safe_name,
                        workspace=workspace,
                        doc_id=answer_id,
                    )
                else:
                    storage_uri = await s3_client.upload_file(
                        temporary_path,
                        safe_name,
                        workspace=workspace,
                        doc_id=answer_id,
                    )
            finally:
                if temporary_path and temporary_path.exists():
                    temporary_path.unlink()
            if storage_uri:
                storage_type = "s3"

        if not storage_uri:
            relative_path = Path(
                _safe_asset_segment(workspace, "workspace"),
                _safe_asset_segment(answer_id, "answer"),
                asset_id,
                safe_name,
            )
            local_path = _local_answer_asset_path(str(relative_path))
            local_path.parent.mkdir(parents=True, exist_ok=True)
            local_path.write_bytes(content)
            storage_uri = str(relative_path)
            storage_type = "local"

        metadata = {
            **metadata,
            "original_file_name": file.filename or safe_name,
            "uploaded_at": datetime.now(timezone.utc).isoformat(),
        }
        await db.query(
            """
            INSERT INTO LIGHTRAG_ANSWER_ASSETS
                (asset_id, workspace, answer_id, answer_version, asset_type,
                 storage_type, storage_uri, file_name, mime_type, file_size,
                 caption, alt_text, search_text, display_order, metadata)
            VALUES
                ($1, $2, $3, $4, $5, $6, $7, $8, $9, $10,
                 $11, $12, $13, $14, $15::jsonb)
            RETURNING asset_id, workspace, answer_id, answer_version, asset_type,
                      storage_type, storage_uri, file_name, mime_type, file_size,
                      caption, alt_text, search_text, content_text, display_order,
                      is_active, metadata, create_time, update_time
            """,
            [
                asset_id,
                workspace,
                answer_id,
                int(current.get("version") or 1) + 1,
                resolved_asset_type,
                storage_type,
                storage_uri,
                safe_name,
                mime_type,
                file_size,
                caption,
                alt_text,
                search_text,
                display_order,
                _json(metadata),
            ],
        )
        await _touch_answer_after_asset_change(db, workspace, answer_id)
        return _answer_asset_from_row(
            await _get_answer_asset_row(db, workspace, answer_id, asset_id)
        )

    @router.patch(
        "/{answer_id}/assets/{asset_id}",
        response_model=AnswerAsset,
        dependencies=[Depends(combined_auth)],
    )
    async def update_answer_asset(
        request: Request,
        answer_id: str,
        asset_id: str,
        payload: AnswerAssetUpdateRequest,
    ):
        workspace, db = await db_for_request(request)
        current = await _get_answer_asset_row(db, workspace, answer_id, asset_id)
        values = dict(current)
        values.update(payload.model_dump(exclude_unset=True))
        await db.query(
            """
            UPDATE LIGHTRAG_ANSWER_ASSETS
            SET caption = $4,
                alt_text = $5,
                search_text = $6,
                content_text = $7,
                display_order = $8,
                metadata = $9::jsonb,
                update_time = NOW()
            WHERE workspace = $1 AND answer_id = $2 AND asset_id = $3
              AND is_active = TRUE
            RETURNING asset_id, workspace, answer_id, answer_version, asset_type,
                      storage_type, storage_uri, file_name, mime_type, file_size,
                      caption, alt_text, search_text, content_text, display_order,
                      is_active, metadata, create_time, update_time
            """,
            [
                workspace,
                answer_id,
                asset_id,
                values.get("caption"),
                values.get("alt_text"),
                values.get("search_text"),
                values.get("content_text"),
                int(values.get("display_order") or 0),
                _json(values.get("metadata")),
            ],
        )
        await _touch_answer_after_asset_change(db, workspace, answer_id)
        return _answer_asset_from_row(
            await _get_answer_asset_row(db, workspace, answer_id, asset_id)
        )

    @router.delete(
        "/{answer_id}/assets/{asset_id}",
        dependencies=[Depends(combined_auth)],
    )
    async def delete_answer_asset(
        request: Request,
        answer_id: str,
        asset_id: str,
    ):
        workspace, db = await db_for_request(request)
        await _get_answer_asset_row(db, workspace, answer_id, asset_id)
        row = await db.query(
            """
            UPDATE LIGHTRAG_ANSWER_ASSETS
            SET is_active = FALSE,
                update_time = NOW()
            WHERE workspace = $1 AND answer_id = $2 AND asset_id = $3
              AND is_active = TRUE
            RETURNING asset_id
            """,
            [workspace, answer_id, asset_id],
        )
        if not row:
            raise HTTPException(status_code=404, detail=f"Asset '{asset_id}' not found")
        await _touch_answer_after_asset_change(db, workspace, answer_id)
        return {
            "message": "Answer asset removed",
            "answer_id": answer_id,
            "asset_id": asset_id,
        }

    @router.get(
        "/{answer_id}/assets/{asset_id}/content",
        dependencies=[Depends(combined_auth)],
    )
    async def get_answer_asset_content(
        request: Request,
        answer_id: str,
        asset_id: str,
    ):
        workspace, db = await db_for_request(request)
        row = await _get_answer_asset_row(db, workspace, answer_id, asset_id)
        storage_type = str(row.get("storage_type") or "")
        storage_uri = str(row.get("storage_uri") or "").strip()
        if storage_type in {"external", "s3"}:
            return RedirectResponse(_validate_external_asset_url(storage_uri))
        if storage_type == "inline":
            return Response(
                content=str(row.get("content_text") or ""),
                media_type=str(row.get("mime_type") or "text/plain"),
            )
        if storage_type != "local" or not storage_uri:
            raise HTTPException(status_code=404, detail="Asset content is unavailable")
        local_path = _local_answer_asset_path(storage_uri)
        if not local_path.is_file():
            raise HTTPException(status_code=404, detail="Asset file is unavailable")
        return FileResponse(
            local_path,
            media_type=str(row.get("mime_type") or "application/octet-stream"),
            filename=str(row.get("file_name") or local_path.name),
            content_disposition_type="inline",
        )

    @router.get("/{answer_id}", response_model=AnswerItem, dependencies=[Depends(combined_auth)])
    async def get_answer(request: Request, answer_id: str):
        workspace, db = await db_for_request(request)
        row = await _get_answer_row(db, workspace, answer_id)
        answer = _answer_from_row(dict(row))
        await _hydrate_answer_assets(db, workspace, [answer])
        return answer

    @router.get("/{answer_id}/sources", response_model=list[AnswerSourceLink], dependencies=[Depends(combined_auth)])
    async def list_answer_sources(request: Request, answer_id: str):
        workspace, db = await db_for_request(request)
        await _get_answer_row(db, workspace, answer_id)
        rows = await db.query(
            """
            SELECT l.link_id, l.workspace, l.answer_id, l.answer_version, l.snapshot_id,
                   l.link_type, l.metadata, l.create_time,
                   s.source_type, s.source_uri, s.file_name, s.title,
                   LEFT(s.raw_content, 600) AS content_preview,
                   s.content_hash, s.content_length, s.profile AS snapshot_profile,
                   s.metadata AS snapshot_metadata, s.created_answer_ids, s.status,
                   s.task_id, s.create_time AS snapshot_create_time
            FROM LIGHTRAG_ANSWER_SOURCE_LINKS l
            LEFT JOIN LIGHTRAG_ANSWER_SOURCE_SNAPSHOTS s
              ON s.workspace = l.workspace AND s.snapshot_id = l.snapshot_id
            WHERE l.workspace = $1 AND l.answer_id = $2
            ORDER BY l.create_time DESC
            """,
            [workspace, answer_id],
            multirows=True,
        )
        links: list[AnswerSourceLink] = []
        for row in rows or []:
            row_dict = dict(row)
            snapshot = None
            if row_dict.get("source_type"):
                snapshot = _source_snapshot_from_row(
                    {
                        "snapshot_id": row_dict["snapshot_id"],
                        "workspace": row_dict["workspace"],
                        "source_type": row_dict.get("source_type"),
                        "source_uri": row_dict.get("source_uri"),
                        "file_name": row_dict.get("file_name"),
                        "title": row_dict.get("title"),
                        "content_preview": row_dict.get("content_preview"),
                        "content_hash": row_dict.get("content_hash"),
                        "content_length": row_dict.get("content_length"),
                        "profile": row_dict.get("snapshot_profile"),
                        "metadata": row_dict.get("snapshot_metadata"),
                        "created_answer_ids": row_dict.get("created_answer_ids"),
                        "status": row_dict.get("status"),
                        "task_id": row_dict.get("task_id"),
                        "create_time": row_dict.get("snapshot_create_time"),
                    }
                )
            links.append(_source_link_from_row(row_dict, snapshot=snapshot))
        return links

    @router.post("/{answer_id}/view", dependencies=[Depends(combined_auth)])
    async def record_answer_view(
        request: Request,
        answer_id: str,
        payload: Optional[AnswerViewRequest] = None,
    ):
        workspace, db = await db_for_request(request)
        await _get_answer_row(db, workspace, answer_id)
        event_payload = payload or AnswerViewRequest()
        event_id = await _log_event(
            db,
            workspace,
            "view",
            query=event_payload.query,
            selected_answer_id=answer_id,
            candidate_ids=[answer_id],
            scores={answer_id: 1.0},
            metadata=event_payload.metadata,
        )
        return {"message": "Answer view recorded", "event_id": event_id, "answer_id": answer_id}

    @router.post("/{answer_id}/feedback", dependencies=[Depends(combined_auth)])
    async def record_answer_feedback(
        request: Request,
        answer_id: str,
        payload: AnswerFeedbackRequest,
    ):
        workspace, db = await db_for_request(request)
        await _get_answer_row(db, workspace, answer_id)
        candidate_ids = [answer_id]
        if payload.selected_alternative_id:
            await _get_answer_row(db, workspace, payload.selected_alternative_id)
            candidate_ids.append(payload.selected_alternative_id)
        metadata = {
            **payload.metadata,
            "helpful": payload.helpful,
            "note": payload.note,
            "selected_alternative_id": payload.selected_alternative_id,
        }
        event_id = await _log_event(
            db,
            workspace,
            "feedback",
            query=payload.query,
            selected_answer_id=answer_id,
            candidate_ids=candidate_ids,
            scores={answer_id: 1.0},
            metadata=metadata,
        )
        return {"message": "Answer feedback recorded", "event_id": event_id, "answer_id": answer_id}

    @router.patch("/{answer_id}", response_model=AnswerItem, dependencies=[Depends(combined_auth)])
    async def update_answer(request: Request, answer_id: str, payload: AnswerUpdateRequest):
        workspace, db = await db_for_request(request)
        current = await _get_answer_row(db, workspace, answer_id)
        next_values = dict(current)
        for key, value in payload.model_dump(exclude_unset=True).items():
            next_values[key] = value
        next_version = int(current.get("version") or 1) + 1
        row = await db.query(
            """
            UPDATE LIGHTRAG_ANSWER_ITEMS
            SET title = $3,
                body = $4,
                approved_summary = $5,
                content_format = $6,
                display_policy = $7,
                status = $8,
                version = $9,
                valid_from = $10,
                valid_until = $11,
                priority = $12,
                tags = $13::jsonb,
                metadata = $14::jsonb,
                publish_time = CASE
                    WHEN $8 = 'published' AND publish_time IS NULL THEN NOW()
                    WHEN $8 <> 'published' THEN NULL
                    ELSE publish_time
                END,
                update_time = NOW()
            WHERE workspace = $1 AND answer_id = $2
            RETURNING workspace, answer_id, title, body, approved_summary, content_format,
                      display_policy, status, version, valid_from, valid_until, priority,
                      tags, metadata, publish_time, create_time, update_time
            """,
            [
                workspace,
                answer_id,
                next_values.get("title"),
                next_values.get("body"),
                next_values.get("approved_summary"),
                next_values.get("content_format") or "markdown",
                next_values.get("display_policy") or "both",
                next_values.get("status") or "draft",
                next_version,
                next_values.get("valid_from"),
                next_values.get("valid_until"),
                int(next_values.get("priority") or 0),
                _json_list(next_values.get("tags")),
                _json(next_values.get("metadata")),
            ],
        )
        if not row:
            raise HTTPException(status_code=404, detail=f"Answer '{answer_id}' not found")
        await _record_revision(db, dict(row))
        await _delete_answer_vectors(db, workspace, answer_id)
        asyncio.create_task(
            _refresh_answer_vector_safely(
                db=db,
                workspace=workspace,
                answer_id=answer_id,
            )
        )
        await _mark_answer_graph_stale(db, workspace, answer_id, next_version)
        asyncio.create_task(
            _refresh_answer_graph_safely(
                db=db,
                workspace=workspace,
                answer_id=answer_id,
                trigger_source="answer_update",
            )
        )
        answer = _answer_from_row(dict(row))
        await _hydrate_answer_assets(db, workspace, [answer])
        return answer

    @router.post("/{answer_id}/publish", response_model=AnswerItem, dependencies=[Depends(combined_auth)])
    async def publish_answer(request: Request, answer_id: str):
        return await update_answer(request, answer_id, AnswerUpdateRequest(status="published"))

    @router.post("/{answer_id}/archive", response_model=AnswerItem, dependencies=[Depends(combined_auth)])
    async def archive_answer(request: Request, answer_id: str):
        return await update_answer(request, answer_id, AnswerUpdateRequest(status="archived"))

    @router.get("/{answer_id}/revisions", response_model=list[AnswerRevision], dependencies=[Depends(combined_auth)])
    async def list_revisions(request: Request, answer_id: str):
        workspace, db = await db_for_request(request)
        rows = await db.query(
            """
            SELECT revision_id, workspace, answer_id, version, snapshot_json, created_at
            FROM LIGHTRAG_ANSWER_REVISIONS
            WHERE workspace = $1 AND answer_id = $2
            ORDER BY version DESC, created_at DESC
            """,
            [workspace, answer_id],
            multirows=True,
        )
        revisions = []
        for row in rows or []:
            snapshot = _coerce_json(row.get("snapshot_json"), {})
            if not isinstance(snapshot, dict):
                snapshot = {}
            revisions.append(
                AnswerRevision(
                    revision_id=str(row["revision_id"]),
                    workspace=str(row["workspace"]),
                    answer_id=str(row["answer_id"]),
                    version=int(row["version"]),
                    snapshot=_normalise_answer_snapshot(snapshot),
                    created_at=_iso(row.get("created_at")),
                )
            )
        return revisions

    @router.post(
        "/{answer_id}/revisions/{revision_id}/restore",
        response_model=AnswerItem,
        dependencies=[Depends(combined_auth)],
    )
    async def restore_revision(request: Request, answer_id: str, revision_id: str):
        workspace, db = await db_for_request(request)
        current = await _get_answer_row(db, workspace, answer_id)
        revision = await db.query(
            """
            SELECT snapshot_json
            FROM LIGHTRAG_ANSWER_REVISIONS
            WHERE workspace = $1 AND answer_id = $2 AND revision_id = $3
            """,
            [workspace, answer_id, revision_id],
        )
        if not revision:
            raise HTTPException(status_code=404, detail=f"Revision '{revision_id}' not found")

        snapshot = _coerce_json(revision.get("snapshot_json"), {})
        if not isinstance(snapshot, dict):
            raise HTTPException(status_code=409, detail="Revision snapshot is invalid")
        snapshot = _normalise_answer_snapshot(snapshot)
        next_version = int(current.get("version") or 1) + 1
        restored_status = snapshot.get("status") or current.get("status") or "draft"
        row = await db.query(
            """
            UPDATE LIGHTRAG_ANSWER_ITEMS
            SET title = $3,
                body = $4,
                approved_summary = $5,
                content_format = $6,
                display_policy = $7,
                status = $8,
                version = $9,
                valid_from = $10,
                valid_until = $11,
                priority = $12,
                tags = $13::jsonb,
                metadata = $14::jsonb,
                publish_time = CASE
                    WHEN $8 = 'published' THEN COALESCE(publish_time, NOW())
                    ELSE NULL
                END,
                update_time = NOW()
            WHERE workspace = $1 AND answer_id = $2
            RETURNING workspace, answer_id, title, body, approved_summary, content_format,
                      display_policy, status, version, valid_from, valid_until, priority,
                      tags, metadata, publish_time, create_time, update_time
            """,
            [
                workspace,
                answer_id,
                snapshot.get("title") or current.get("title"),
                snapshot.get("body") or current.get("body"),
                snapshot.get("approved_summary"),
                snapshot.get("content_format") or current.get("content_format") or "markdown",
                snapshot.get("display_policy") or current.get("display_policy") or "both",
                restored_status,
                next_version,
                snapshot.get("valid_from"),
                snapshot.get("valid_until"),
                int(snapshot.get("priority") or 0),
                _json_list(snapshot.get("tags")),
                _json(snapshot.get("metadata")),
            ],
        )
        if not row:
            raise HTTPException(status_code=404, detail=f"Answer '{answer_id}' not found")
        if "assets" in snapshot:
            await _restore_answer_assets(
                db,
                workspace,
                answer_id,
                next_version,
                snapshot.get("assets") or [],
            )
        await _record_revision(db, dict(row))
        await _delete_answer_vectors(db, workspace, answer_id)
        asyncio.create_task(
            _refresh_answer_vector_safely(
                db=db,
                workspace=workspace,
                answer_id=answer_id,
            )
        )
        await _mark_answer_graph_stale(db, workspace, answer_id, next_version)
        asyncio.create_task(
            _refresh_answer_graph_safely(
                db=db,
                workspace=workspace,
                answer_id=answer_id,
                trigger_source="answer_restore",
            )
        )
        answer = _answer_from_row(dict(row))
        await _hydrate_answer_assets(db, workspace, [answer])
        return answer

    @router.get("/{answer_id}/guidance", response_model=list[AnswerGuidance], dependencies=[Depends(combined_auth)])
    async def list_guidance(request: Request, answer_id: str):
        workspace, db = await db_for_request(request)
        await _get_answer_row(db, workspace, answer_id)
        rows = await db.query(
            """
            SELECT guidance_id, workspace, answer_id, guidance_type, text, weight, metadata, create_time
            FROM LIGHTRAG_ANSWER_GUIDANCE
            WHERE workspace = $1 AND answer_id = $2
            ORDER BY create_time DESC
            """,
            [workspace, answer_id],
            multirows=True,
        )
        return [_guidance_from_row(dict(row)) for row in rows or []]

    @router.post(
        "/{answer_id}/guidance/suggest",
        response_model=GuidanceSuggestResponse,
        dependencies=[Depends(combined_auth)],
    )
    async def suggest_guidance(request: Request, answer_id: str, payload: GuidanceSuggestRequest):
        workspace, db = await db_for_request(request)
        answer = _answer_from_row(await _get_answer_row(db, workspace, answer_id))
        guidance_rows = await db.query(
            """
            SELECT guidance_id, workspace, answer_id, guidance_type, text, weight, metadata, create_time
            FROM LIGHTRAG_ANSWER_GUIDANCE
            WHERE workspace = $1 AND answer_id = $2
            ORDER BY create_time DESC
            """,
            [workspace, answer_id],
            multirows=True,
        )
        existing_guidance = [_guidance_from_row(dict(row)) for row in guidance_rows or []]
        suggestions = _heuristic_guidance_suggestions(
            answer,
            existing_guidance,
            payload.query_examples,
            payload.max_suggestions,
        )
        mode = "heuristic"

        if payload.use_llm:
            workspace_rag = await rag_for_workspace(workspace)
            llm_func = getattr(workspace_rag, "llm_model_func", None)
            if llm_func is not None:
                prompt = (
                    "Create matching hints for a fixed FAQ answer. "
                    "Return only a JSON array. Each item must contain guidance_type, text, weight. "
                    "Allowed guidance_type values: question, keyword, synonym, negative_keyword, note.\n\n"
                    f"Answer title: {answer.title}\n"
                    f"Summary: {answer.approved_summary or ''}\n"
                    f"Body:\n{answer.body[:1800]}\n"
                    f"Existing hints: {json.dumps([item.text for item in existing_guidance], ensure_ascii=False)}\n"
                    f"Query examples: {json.dumps(payload.query_examples, ensure_ascii=False)}"
                )
                try:
                    raw = await llm_func(
                        prompt,
                        system_prompt=(
                            "You help FAQ operators add short representative questions, keywords, and synonyms. "
                            "Do not create final answer content."
                        ),
                        _llm_purpose="knowledge_structure",
                    )
                    seen = {_normalise_text(item.text) for item in existing_guidance + suggestions}
                    allowed_types = {"question", "keyword", "synonym", "negative_keyword", "note"}
                    for item in _extract_json_array(str(raw)):
                        if not isinstance(item, dict):
                            continue
                        guidance_type = item.get("guidance_type") or "keyword"
                        text = item.get("text")
                        if guidance_type not in allowed_types or not text:
                            continue
                        key = _normalise_text(text)
                        if key in seen:
                            continue
                        seen.add(key)
                        suggestions.append(
                            SourceGuidanceCandidate(
                                guidance_type=guidance_type,
                                text=str(text).strip()[:200],
                                weight=_safe_float(item.get("weight"), 1.0, 0.0, 10.0),
                                source="llm",
                                metadata={"suggested_by": "llm"},
                            )
                        )
                    if suggestions:
                        mode = "llm"
                except Exception as exc:
                    logger.warning("[Answers] LLM guidance suggestion failed for %s: %s", answer_id, exc)

        return GuidanceSuggestResponse(suggestions=suggestions[: payload.max_suggestions], mode=mode)

    @router.post("/{answer_id}/guidance", response_model=AnswerGuidance, dependencies=[Depends(combined_auth)])
    async def create_guidance(request: Request, answer_id: str, payload: GuidanceCreateRequest):
        workspace, db = await db_for_request(request)
        answer_row = await _get_answer_row(db, workspace, answer_id)
        guidance_id = f"agd-{uuid.uuid4().hex}"
        row = await db.query(
            """
            INSERT INTO LIGHTRAG_ANSWER_GUIDANCE
                (guidance_id, workspace, answer_id, guidance_type, text, weight, metadata)
            VALUES ($1, $2, $3, $4, $5, $6, $7::jsonb)
            RETURNING guidance_id, workspace, answer_id, guidance_type, text, weight, metadata, create_time
            """,
            [
                guidance_id,
                workspace,
                answer_id,
                payload.guidance_type,
                payload.text,
                payload.weight,
                _json(payload.metadata),
            ],
        )
        if not row:
            raise HTTPException(status_code=500, detail="Failed to create guidance")
        await _delete_answer_vectors(db, workspace, answer_id)
        asyncio.create_task(
            _refresh_answer_vector_safely(
                db=db,
                workspace=workspace,
                answer_id=answer_id,
            )
        )
        await _mark_answer_graph_stale(
            db, workspace, answer_id, int(answer_row.get("version") or 1)
        )
        asyncio.create_task(
            _refresh_answer_graph_safely(
                db=db,
                workspace=workspace,
                answer_id=answer_id,
                trigger_source="guidance_create",
            )
        )
        return _guidance_from_row(dict(row))

    @router.delete("/{answer_id}/guidance/{guidance_id}", dependencies=[Depends(combined_auth)])
    async def delete_guidance(request: Request, answer_id: str, guidance_id: str):
        workspace, db = await db_for_request(request)
        answer_row = await _get_answer_row(db, workspace, answer_id)
        await db.execute(
            """
            DELETE FROM LIGHTRAG_ANSWER_GUIDANCE
            WHERE workspace = $1 AND answer_id = $2 AND guidance_id = $3
            """,
            {"workspace": workspace, "answer_id": answer_id, "guidance_id": guidance_id},
        )
        await _delete_answer_vectors(db, workspace, answer_id)
        asyncio.create_task(
            _refresh_answer_vector_safely(
                db=db,
                workspace=workspace,
                answer_id=answer_id,
            )
        )
        await _mark_answer_graph_stale(
            db, workspace, answer_id, int(answer_row.get("version") or 1)
        )
        asyncio.create_task(
            _refresh_answer_graph_safely(
                db=db,
                workspace=workspace,
                answer_id=answer_id,
                trigger_source="guidance_delete",
            )
        )
        return {"message": "Guidance deleted", "guidance_id": guidance_id}

    @router.post("/{answer_id}/vectors/rebuild", dependencies=[Depends(combined_auth)])
    async def rebuild_answer_vector(request: Request, answer_id: str):
        workspace, db = await db_for_request(request)
        answer = _answer_from_row(await _get_answer_row(db, workspace, answer_id))
        guidance_rows = await db.query(
            """
            SELECT guidance_id, workspace, answer_id, guidance_type, text, weight, metadata, create_time
            FROM LIGHTRAG_ANSWER_GUIDANCE
            WHERE workspace = $1 AND answer_id = $2
            """,
            [workspace, answer_id],
            multirows=True,
        )
        workspace_rag = await rag_for_workspace(workspace)
        embedding = await _ensure_answer_vector(
            db,
            workspace,
            workspace_rag,
            answer,
            [_guidance_from_row(dict(row)) for row in guidance_rows or []],
        )
        if not embedding:
            raise HTTPException(status_code=409, detail="Answer vector could not be built. Check embedding configuration.")
        return {"message": "Answer vector rebuilt", "answer_id": answer_id, "dimensions": len(embedding)}

    @router.post("/vectors/rebuild", dependencies=[Depends(combined_auth)])
    async def rebuild_answer_vectors(
        request: Request,
        status: Optional[str] = Query(default=None),
        limit: int = Query(default=500, ge=1, le=1000),
        only_missing: bool = Query(
            default=True,
            description="Build the next answers without a current vector instead of rebuilding the same latest answers.",
        ),
    ):
        workspace, db = await db_for_request(request)
        status_filter = [status] if status and status != "all" else ["draft", "published"]
        rows = await db.query(
            """
            SELECT
                answers.workspace,
                answers.answer_id,
                answers.title,
                answers.body,
                answers.approved_summary,
                answers.content_format,
                answers.display_policy,
                answers.status,
                answers.version,
                answers.valid_from,
                answers.valid_until,
                answers.priority,
                answers.tags,
                answers.metadata,
                answers.publish_time,
                answers.create_time,
                answers.update_time
            FROM LIGHTRAG_ANSWER_ITEMS AS answers
            LEFT JOIN LIGHTRAG_ANSWER_VECTORS AS vectors
              ON vectors.workspace = answers.workspace
             AND vectors.answer_id = answers.answer_id
             AND vectors.answer_version = answers.version
             AND vectors.vector_kind = 'combined'
             AND COALESCE((vectors.metadata->>'projection_version')::integer, 0) = $5
            WHERE answers.workspace = $1
              AND answers.status = ANY($2::text[])
              AND ($3::boolean = FALSE OR vectors.answer_id IS NULL)
            ORDER BY answers.priority DESC, answers.update_time DESC
            LIMIT $4
            """,
            [
                workspace,
                status_filter,
                only_missing,
                limit,
                ANSWER_VECTOR_PROJECTION_VERSION,
            ],
            multirows=True,
        )
        answer_ids = [str(row["answer_id"]) for row in rows or []]
        guidance_rows = []
        if answer_ids:
            guidance_rows = await db.query(
                """
                SELECT guidance_id, workspace, answer_id, guidance_type, text, weight, metadata, create_time
                FROM LIGHTRAG_ANSWER_GUIDANCE
                WHERE workspace = $1
                  AND answer_id = ANY($2::text[])
                """,
                [workspace, answer_ids],
                multirows=True,
            )
        guidance_by_answer: dict[str, list[AnswerGuidance]] = {}
        for row in guidance_rows or []:
            guidance = _guidance_from_row(dict(row))
            guidance_by_answer.setdefault(guidance.answer_id, []).append(guidance)

        workspace_rag = await rag_for_workspace(workspace)
        rebuilt = 0
        failed: list[str] = []
        for row in rows or []:
            answer = _answer_from_row(dict(row))
            embedding = await _ensure_answer_vector(
                db,
                workspace,
                workspace_rag,
                answer,
                guidance_by_answer.get(answer.answer_id, []),
            )
            if embedding:
                rebuilt += 1
            else:
                failed.append(answer.answer_id)
        remaining = await db.query(
            """
            SELECT COUNT(*) AS count
            FROM LIGHTRAG_ANSWER_ITEMS AS answers
            LEFT JOIN LIGHTRAG_ANSWER_VECTORS AS vectors
              ON vectors.workspace = answers.workspace
             AND vectors.answer_id = answers.answer_id
             AND vectors.answer_version = answers.version
             AND vectors.vector_kind = 'combined'
             AND COALESCE((vectors.metadata->>'projection_version')::integer, 0) = $3
            WHERE answers.workspace = $1
              AND answers.status = ANY($2::text[])
              AND vectors.answer_id IS NULL
            """,
            [workspace, status_filter, ANSWER_VECTOR_PROJECTION_VERSION],
        )
        return {
            "message": "Answer vectors rebuild completed",
            "processed": len(answer_ids),
            "rebuilt": rebuilt,
            "failed": failed,
            "remaining": int(remaining.get("count") or 0) if remaining else 0,
        }

    async def resolve_answer_candidates(
        workspace: str,
        db,
        payload: ResolveRequest,
        *,
        event_type: str = "resolve",
        event_metadata: Optional[dict[str, Any]] = None,
    ) -> ResolveResponse:
        started = time.time()
        status_filter = ["published"] if not payload.include_drafts else ["draft", "published"]
        allowed_answer_ids = (
            [str(answer_id) for answer_id in payload.allowed_answer_ids if str(answer_id).strip()]
            if payload.allowed_answer_ids is not None
            else None
        )
        alias_expansions = await _expand_query_aliases(db, workspace, payload.query)
        keyword_candidate_ids, keyword_candidates_truncated = await _keyword_candidate_ids(
            db,
            workspace,
            payload.query,
            status_filter,
            allowed_answer_ids,
            alias_expansions,
        )
        vector_scores: dict[str, float] = {}
        vector_status = "not_requested"
        graph_scores: dict[str, float] = {}
        graph_evidence: dict[str, list[AnswerGraphEvidence]] = {}
        graph_status = "not_requested"
        graph_config = await _get_answer_graph_config(db, workspace)
        workspace_rag = await rag_for_workspace(workspace)

        if payload.retrieval_mode in {"hybrid", "graph_hybrid", "llm_rerank"}:
            vector_scores, vector_status = await _answer_vector_scores(
                db,
                workspace,
                workspace_rag,
                payload.query,
                payload.vector_top_k,
                status_filter,
                allowed_answer_ids,
            )
        graph_requested = payload.retrieval_mode == "graph_hybrid" or (
            payload.retrieval_mode == "llm_rerank" and graph_config.enabled
        )
        if graph_requested:
            graph_scores, graph_evidence, graph_status = await _answer_graph_scores(
                db,
                workspace,
                workspace_rag,
                payload.query,
                payload.vector_top_k,
                allowed_answer_ids,
                graph_config,
            )
        effective_retrieval_mode, retrieval_fallback_reason = (
            _effective_answer_retrieval(
                payload.retrieval_mode,
                graph_status,
            )
        )

        candidate_answer_ids = list(
            dict.fromkeys(
                [
                    *keyword_candidate_ids,
                    *vector_scores.keys(),
                    *graph_scores.keys(),
                ]
            )
        )
        rows = await db.query(
            """
            SELECT workspace, answer_id, title, body, approved_summary, content_format,
                   display_policy, status, version, valid_from, valid_until, priority,
                   tags, metadata, publish_time, create_time, update_time
            FROM LIGHTRAG_ANSWER_ITEMS
            WHERE workspace = $1
              AND answer_id = ANY($2::text[])
              AND status = ANY($3::text[])
              AND (valid_from IS NULL OR valid_from <= NOW())
              AND (valid_until IS NULL OR valid_until >= NOW())
            """,
            [workspace, candidate_answer_ids, status_filter],
            multirows=True,
        )
        guidance_rows = []
        if candidate_answer_ids:
            guidance_rows = await db.query(
                """
                SELECT guidance_id, workspace, answer_id, guidance_type, text, weight, metadata, create_time
                FROM LIGHTRAG_ANSWER_GUIDANCE
                WHERE workspace = $1
                  AND answer_id = ANY($2::text[])
                """,
                [workspace, candidate_answer_ids],
                multirows=True,
            )
        guidance_by_answer: dict[str, list[AnswerGuidance]] = {}
        for row in guidance_rows or []:
            guidance = _guidance_from_row(dict(row))
            guidance_by_answer.setdefault(guidance.answer_id, []).append(guidance)

        answers = [_answer_from_row(dict(row)) for row in rows or []]
        await _hydrate_answer_assets(db, workspace, answers)
        answer_by_id = {answer.answer_id: answer for answer in answers}
        keyword_scores: dict[str, tuple[float, list[str], str, dict[str, float]]] = {}

        for row in rows or []:
            answer = answer_by_id[str(row["answer_id"])]
            score, matched_guidance, reason, score_details = _score_candidate(
                answer,
                guidance_by_answer.get(answer.answer_id, []),
                payload.query,
                payload.strategy,
                alias_expansions,
            )
            keyword_scores[answer.answer_id] = (score, matched_guidance, reason, score_details)

        candidate_ids = {
            answer_id
            for answer_id, (keyword_score, _, _, _) in keyword_scores.items()
            if keyword_score > 0
        } | {
            answer_id
            for answer_id, vector_score in vector_scores.items()
            if vector_score > 0
        } | {
            answer_id
            for answer_id, graph_score in graph_scores.items()
            if graph_score > 0
        }

        candidates: list[ResolveCandidate] = []
        for answer_id in candidate_ids:
            answer = answer_by_id.get(answer_id)
            if answer is None:
                continue
            keyword_score, matched_guidance, reason, score_details = keyword_scores.get(
                answer_id,
                (0.0, [], "vector_match", {}),
            )
            if reason == "blocked_by_negative_guidance":
                continue
            vector_score = vector_scores.get(answer_id, 0.0)
            selected_by = "keyword"
            final_score = keyword_score
            detail_payload = dict(score_details)
            _add_score(detail_payload, "keyword_score", keyword_score)
            if (
                payload.retrieval_mode in {"hybrid", "graph_hybrid", "llm_rerank"}
                and vector_score > 0
            ):
                priority_boost = min(0.05, max(answer.priority, 0) * 0.005)
                final_score = min(1.0, keyword_score * 0.6 + vector_score * 0.35 + priority_boost)
                _add_score(detail_payload, "vector_score", vector_score)
                _add_score(detail_payload, "hybrid_priority", priority_boost)
                _add_score(detail_payload, "hybrid_score", final_score)
                if vector_score > keyword_score:
                    selected_by = "vector"
                elif vector_score > 0:
                    selected_by = "hybrid"
                reason = (
                    f"hybrid keyword:{keyword_score:.2f}, vector:{vector_score:.2f}"
                    if vector_score > 0
                    else reason
                )
            graph_score = graph_scores.get(answer_id, 0.0)
            candidate_graph_evidence = graph_evidence.get(answer_id, [])
            if graph_requested and graph_score > 0:
                graph_contribution = graph_score * graph_config.graph_weight
                final_score = min(
                    1.0,
                    1.0 - (1.0 - final_score) * (1.0 - graph_contribution),
                )
                _add_score(detail_payload, "graph_score", graph_score)
                _add_score(
                    detail_payload,
                    "graph_weight",
                    graph_config.graph_weight,
                )
                _add_score(detail_payload, "graph_hybrid_score", final_score)
                selected_by = (
                    "graph"
                    if keyword_score <= 0 and vector_score <= 0
                    else "graph_hybrid"
                )
                reason = (
                    f"{reason}; graph:{graph_score:.2f} "
                    f"(weight:{graph_config.graph_weight:.2f})"
                )
            candidates.append(
                ResolveCandidate(
                    answer=_search_answer_view(answer),
                    score=round(final_score, 4),
                    matched_guidance=matched_guidance,
                    reason=reason,
                    score_details=detail_payload,
                    selected_by=selected_by,
                    graph_evidence=candidate_graph_evidence,
                )
            )

        candidates.sort(key=lambda item: (item.score, item.answer.priority, item.answer.update_time or ""), reverse=True)
        llm_selection: dict[str, Any] = {"status": "not_requested"}
        if payload.retrieval_mode == "llm_rerank":
            llm_candidates = candidates[: payload.llm_candidate_count]
            llm_answer_id, llm_selection = await _llm_select_candidate(workspace_rag, payload.query, llm_candidates)
            if llm_answer_id:
                candidates.sort(
                    key=lambda item: (
                        item.answer.answer_id == llm_answer_id,
                        item.score,
                        item.answer.priority,
                        item.answer.update_time or "",
                    ),
                    reverse=True,
                )
                for item in candidates:
                    if item.answer.answer_id == llm_answer_id:
                        item.selected_by = "llm_id_selector"
                        item.reason = f"LLM selected candidate ID. {item.reason}"
                        item.score_details["llm_confidence"] = _safe_float(
                            llm_selection.get("confidence"),
                            item.score,
                            0.0,
                            1.0,
                        )
                        break

        effective_selection_policy: Literal["coverage", "precision"] = (
            "precision"
            if payload.selection_policy == "precision"
            or (
                payload.selection_policy == "workspace"
                and graph_config.precision_mode
            )
            else "coverage"
        )
        precision_diagnostics: dict[str, Any] = {}
        abstention_reason: Optional[str] = None
        selected = candidates[0] if candidates else None
        if effective_selection_policy == "precision":
            abstention_reason, precision_diagnostics = _precision_abstention_reason(
                selected,
                candidates,
                payload,
                graph_config,
                llm_selection,
            )
            if abstention_reason:
                selected = None
        elif selected and selected.score < payload.min_score:
            selected = None
            abstention_reason = "low_score"

        candidates = candidates[: payload.top_k]
        selected_by = selected.selected_by if selected else "none"
        matched_id = _matched_business_id(selected.answer if selected else None)
        clarification_question = (
            _clarification_question(
                candidates,
                abstention_reason,
                llm_selection,
            )
            if selected is None
            else None
        )
        rationale = (
            f"Selected {selected.answer.answer_id} by {selected_by}: {selected.reason}."
            if selected
            else (
                "No answer was selected because the precision policy rejected "
                f"the candidates: {abstention_reason}."
                if effective_selection_policy == "precision"
                else "No candidate reached the minimum answer matching score."
            )
        )
        if (
            selected
            and llm_selection.get("answer_id")
            and llm_selection.get("rationale")
        ):
            rationale = f"{rationale} LLM rationale: {llm_selection['rationale']}"
        trace_id = await _log_event(
            db,
            workspace,
            event_type,
            query=payload.query,
            selected_answer_id=selected.answer.answer_id if selected else None,
            candidate_ids=[item.answer.answer_id for item in candidates],
            scores={item.answer.answer_id: item.score for item in candidates},
            metadata={
                **(event_metadata or {}),
                "latency_ms": int((time.time() - started) * 1000),
                "mode": payload.retrieval_mode,
                "requested_retrieval_mode": payload.retrieval_mode,
                "effective_retrieval_mode": effective_retrieval_mode,
                "retrieval_fallback_reason": retrieval_fallback_reason,
                "strategy": payload.strategy,
                "selected_by": selected_by,
                "selection_policy": effective_selection_policy,
                "abstention_reason": abstention_reason,
                "clarification_question": clarification_question,
                "precision_diagnostics": precision_diagnostics,
                "matched_id": matched_id,
                "vector_status": vector_status,
                "graph_status": graph_status,
                "graph_schema_version": graph_config.schema_version,
                "graph_weight": graph_config.graph_weight,
                "graph_min_similarity": graph_config.min_similarity,
                "keyword_candidate_count": len(keyword_candidate_ids),
                "keyword_candidates_truncated": keyword_candidates_truncated,
                "keyword_candidate_limit": MAX_KEYWORD_SEARCH_CANDIDATES,
                "alias_expansions": [
                    expansion.model_dump()
                    for expansion in alias_expansions
                ],
                "llm_selection": llm_selection,
                "score_details": {item.answer.answer_id: item.score_details for item in candidates},
                "graph_evidence": {
                    item.answer.answer_id: [
                        evidence_item.model_dump()
                        for evidence_item in item.graph_evidence
                    ]
                    for item in candidates
                    if item.graph_evidence
                },
            },
        )
        return ResolveResponse(
            selected_answer=selected.answer if selected else None,
            matched_id=matched_id,
            confidence=selected.score if selected else 0.0,
            candidates=candidates,
            trace_id=trace_id,
            rationale=rationale,
            retrieval_mode=payload.retrieval_mode,
            requested_retrieval_mode=payload.retrieval_mode,
            effective_retrieval_mode=effective_retrieval_mode,
            graph_status=graph_status,
            retrieval_fallback_reason=retrieval_fallback_reason,
            selected_by=selected_by,
            alias_expansions=alias_expansions,
            selection_policy=effective_selection_policy,
            abstention_reason=abstention_reason,
            clarification_question=clarification_question,
        )

    @router.post("/resolve", response_model=ResolveResponse, dependencies=[Depends(combined_auth)])
    async def resolve_answer(request: Request, payload: ResolveRequest):
        workspace, db = await db_for_request(request)
        return await resolve_answer_candidates(workspace, db, payload, event_type="resolve")

    @router.post("/search", response_model=AnswerSearchResponse, dependencies=[Depends(combined_auth)])
    async def search_answer(request: Request, payload: AnswerSearchRequest):
        workspace, db = await db_for_request(request)
        result = await resolve_answer_candidates(
            workspace,
            db,
            ResolveRequest(
                query=payload.query,
                top_k=payload.top_k,
                min_score=payload.min_score,
                include_drafts=payload.include_drafts,
                strategy=payload.strategy,
                retrieval_mode=payload.retrieval_mode,
                vector_top_k=payload.vector_top_k,
                llm_candidate_count=payload.llm_candidate_count,
                allowed_answer_ids=payload.allowed_answer_ids,
                selection_policy=payload.selection_policy,
            ),
            event_type="search",
            event_metadata={
                "response_policy": payload.response_policy,
                "include_candidates": payload.include_candidates,
            },
        )
        answer = result.selected_answer
        if answer is None:
            return AnswerSearchResponse(
                matched=False,
                confidence=0.0,
                candidates=result.candidates if payload.include_candidates else [],
                trace_id=result.trace_id,
                rationale=result.rationale,
                retrieval_mode=result.retrieval_mode,
                requested_retrieval_mode=result.requested_retrieval_mode,
                effective_retrieval_mode=result.effective_retrieval_mode,
                graph_status=result.graph_status,
                retrieval_fallback_reason=result.retrieval_fallback_reason,
                selected_by=result.selected_by,
                alias_expansions=result.alias_expansions,
                selection_policy=result.selection_policy,
                abstention_reason=result.abstention_reason,
                clarification_question=result.clarification_question,
            )

        response_text, display_policy = _render_answer_response(answer, payload.response_policy)
        return AnswerSearchResponse(
            matched=True,
            answer_id=answer.answer_id,
            matched_id=result.matched_id,
            title=answer.title,
            response=response_text,
            summary=answer.approved_summary,
            full_content=answer.body,
            display_policy=display_policy,
            content_format=answer.content_format,
            status=answer.status,
            version=answer.version,
            valid_from=answer.valid_from,
            valid_until=answer.valid_until,
            confidence=result.confidence,
            tags=answer.tags,
            assets=answer.assets,
            source_type=answer.metadata.get("source_type") or answer.metadata.get("created_from"),
            source_uri=answer.metadata.get("source_uri"),
            candidates=result.candidates if payload.include_candidates else [],
            trace_id=result.trace_id,
            rationale=result.rationale,
            retrieval_mode=result.retrieval_mode,
            requested_retrieval_mode=result.requested_retrieval_mode,
            effective_retrieval_mode=result.effective_retrieval_mode,
            graph_status=result.graph_status,
            retrieval_fallback_reason=result.retrieval_fallback_reason,
            selected_by=result.selected_by,
            alias_expansions=result.alias_expansions,
            selection_policy=result.selection_policy,
            abstention_reason=result.abstention_reason,
            clarification_question=result.clarification_question,
        )

    @router.get("/stats/events", response_model=AnswerEventStatsResponse, dependencies=[Depends(combined_auth)])
    async def answer_event_stats(
        request: Request,
        event_type: Optional[str] = Query(default=None),
        match_status: Optional[Literal["all", "matched", "no_match"]] = Query(default=None),
        source: Optional[str] = Query(default=None),
        mode: Optional[str] = Query(default=None),
        date_from: Optional[datetime] = Query(default=None),
        date_to: Optional[datetime] = Query(default=None),
        min_confidence: Optional[float] = Query(default=None, ge=0.0, le=1.0),
        max_latency_ms: Optional[int] = Query(default=None, ge=0),
        search: Optional[str] = Query(default=None),
        timezone: str = Query(default="Asia/Seoul", max_length=64),
        limit: int = Query(default=20, ge=1, le=100),
    ):
        workspace, db = await db_for_request(request)
        timezone_name = _safe_timezone(timezone)
        params: list[Any] = [workspace]
        where = ["e.workspace = $1"]
        _add_answer_event_filters(
            params,
            where,
            event_type=event_type,
            match_status=None if match_status == "all" else match_status,
            source=source,
            mode=mode,
            date_from=date_from,
            date_to=date_to,
            min_confidence=min_confidence,
            max_latency_ms=max_latency_ms,
            search=search.strip() if search else None,
            timezone_name=timezone_name,
        )
        where_sql = " AND ".join(where)
        from_sql = """
            FROM LIGHTRAG_ANSWER_EVENTS e
            LEFT JOIN LIGHTRAG_ANSWER_ITEMS a
              ON a.workspace = e.workspace AND a.answer_id = e.selected_answer_id
        """
        base = await db.query(
            f"""
            SELECT
              COUNT(*)::INT AS total_events,
              COUNT(*) FILTER (WHERE e.selected_answer_id IS NULL)::INT AS no_match,
              COALESCE(ROUND(AVG(
                CASE
                  WHEN e.metadata->>'latency_ms' ~ '^[0-9]+(\\.[0-9]+)?$'
                  THEN (e.metadata->>'latency_ms')::DOUBLE PRECISION
                  ELSE NULL
                END
              ))::INT, 0) AS avg_latency_ms
            {from_sql}
            WHERE {where_sql}
            """,
            params,
        )

        async def group_rows(select_sql: str, extra_params: Optional[list[Any]] = None) -> list[AnswerAnalyticsGroupRow]:
            local_params = [*params, *(extra_params or []), limit]
            rows = await db.query(
                f"""
                SELECT key, label, COUNT(*)::INT AS count
                FROM (
                  {select_sql}
                  {from_sql}
                  WHERE {where_sql}
                ) grouped
                GROUP BY key, label
                ORDER BY count DESC, label ASC
                LIMIT ${len(local_params)}
                """,
                local_params,
                multirows=True,
            )
            return [_analytics_group_from_row(dict(row)) for row in rows or []]

        selected_answers = await group_rows(
            """
            SELECT
              COALESCE(e.selected_answer_id, '__none__') AS key,
              CASE
                WHEN e.selected_answer_id IS NULL THEN 'No Match'
                ELSE COALESCE(a.title, e.selected_answer_id)
              END AS label
            """
        )
        queries = await group_rows(
            """
            SELECT
              LOWER(COALESCE(NULLIF(BTRIM(e.query), ''), '-')) AS key,
              COALESCE(NULLIF(BTRIM(e.query), ''), '-') AS label
            """
        )
        sources = await group_rows(
            """
            SELECT
              CASE
                WHEN e.selected_answer_id IS NULL THEN '__none__'
                ELSE COALESCE(a.metadata->>'source_type', a.metadata->>'created_from', a.metadata->>'materialization_mode', 'manual')
              END AS key,
              CASE
                WHEN e.selected_answer_id IS NULL THEN 'No Match'
                ELSE COALESCE(a.metadata->>'source_type', a.metadata->>'created_from', a.metadata->>'materialization_mode', 'manual')
              END AS label
            """
        )
        modes = await group_rows(
            """
            SELECT
              COALESCE(e.metadata->>'mode', e.metadata->>'retrieval_mode', '-') AS key,
              COALESCE(e.metadata->>'mode', e.metadata->>'retrieval_mode', '-') AS label
            """
        )
        dates = await group_rows(
            f"""
            SELECT
              TO_CHAR(e.create_time AT TIME ZONE ${len(params) + 1}, 'YYYY-MM-DD') AS key,
              TO_CHAR(e.create_time AT TIME ZONE ${len(params) + 1}, 'YYYY-MM-DD') AS label
            """,
            [timezone_name],
        )
        hours = await group_rows(
            f"""
            SELECT
              TO_CHAR(e.create_time AT TIME ZONE ${len(params) + 1}, 'HH24:00') AS key,
              TO_CHAR(e.create_time AT TIME ZONE ${len(params) + 1}, 'HH24:00') AS label
            """,
            [timezone_name],
        )
        return AnswerEventStatsResponse(
            workspace=workspace,
            total_events=int((base or {}).get("total_events") or 0),
            no_match=int((base or {}).get("no_match") or 0),
            avg_latency_ms=int((base or {}).get("avg_latency_ms") or 0),
            timezone=timezone_name,
            selected_answers=selected_answers,
            queries=queries,
            sources=sources,
            modes=modes,
            dates=dates,
            hours=hours,
        )

    @router.get("/stats/summary", dependencies=[Depends(combined_auth)])
    async def answer_stats(request: Request):
        workspace, db = await db_for_request(request)
        counts = await db.query(
            """
            SELECT
              COUNT(*)::INT AS total,
              COUNT(*) FILTER (WHERE status = 'draft')::INT AS draft,
              COUNT(*) FILTER (WHERE status = 'published')::INT AS published,
              COUNT(*) FILTER (WHERE status = 'archived')::INT AS archived
            FROM LIGHTRAG_ANSWER_ITEMS
            WHERE workspace = $1
            """,
            [workspace],
        )
        events = await db.query(
            """
            SELECT
              COUNT(*) FILTER (WHERE event_type = 'resolve')::INT AS resolves,
              COUNT(*) FILTER (WHERE event_type = 'search')::INT AS searches,
              COUNT(*) FILTER (WHERE event_type = 'view')::INT AS views,
              COUNT(*) FILTER (WHERE event_type = 'feedback')::INT AS feedback
            FROM LIGHTRAG_ANSWER_EVENTS
            WHERE workspace = $1
            """,
            [workspace],
        )
        return {
            "workspace": workspace,
            "answers": counts or {},
            "events": events or {},
        }

    return router
